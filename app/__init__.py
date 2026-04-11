from datetime import datetime, timedelta
from collections import deque
import os
import json
import csv
import io
import re
import unicodedata
from urllib.parse import urlparse

from flask import Response, flash, has_request_context, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from sqlalchemy.exc import OperationalError, ProgrammingError
from sqlalchemy import inspect, text
from sqlalchemy.orm import selectinload
from werkzeug.utils import secure_filename

from models import Categoria, ClientePublico, EnderecoEstoque, Estoque, Fornecedor, Funcionario, FuncaoRH, PerfilAcesso, LancamentoFinanceiro, Mesa, Movimentacao, ProcessoEvento, Produto, PermissaoAcesso, Caixa, MovimentacaoCaixa, Pedido, ItemPedido, Garcom, EmpresaConfig, AuditoriaEvento, AssistenteLocalFeedback, EquipamentoMovimentacao, ManutencaoEquipamento, OrdemServico, ChamadoInterno, FundoSolicitacao, RecebimentoFornecedor, AlmoxarifadoAtribuicao, funcionario_estoques, db
from routes.public_routes import obter_resumo_carrinho_site
from security import csrf_input_tag, csrf_protect_request, ensure_csrf_token, is_json_request, json_response
from app import extensions
from app.cli import register_cli
from app.decorators import _limit, login_required, require_role
from app.exceptions import AppError, BusinessRuleError, NotFound, PermissionDenied, ValidationError
from app.factory import create_app as create_base_app
from app.bootstrap import bootstrap_app
from app.blueprints_registry import register_api_module, register_auth_module, register_domain_modules
from app.core.context_processors import register_context_processors
from app.core.menu import (
    menu_agrupado_para_paginas as core_menu_agrupado_para_paginas,
    menu_navegacao_principal as core_menu_navegacao_principal,
    paginas_ordenadas_menu as core_paginas_ordenadas_menu,
    registrar_debug_menu as core_registrar_debug_menu,
)
from app.core.permissions import (
    carregar_paginas_json as core_carregar_paginas_json,
    expandir_paginas_relacionadas as core_expandir_paginas_relacionadas,
    mapa_permissoes_personalizadas_funcionario as core_mapa_permissoes_personalizadas_funcionario,
    paginas_efetivas_funcionario as core_paginas_efetivas_funcionario,
    paginas_perfil_acesso as core_paginas_perfil_acesso,
    paginas_permitidas_para_funcionario as core_paginas_permitidas_para_funcionario,
    salvar_permissoes_funcionario as core_salvar_permissoes_funcionario,
    serializar_paginas_json as core_serializar_paginas_json,
)
from app.core.runtime_patches import apply_runtime_schema_patches, runtime_schema_patches_enabled
from app.core.startup import (
    bootstrap_admin_configurado as startup_bootstrap_admin_configurado,
    garantir_admin_primeiro_acesso as startup_garantir_admin_primeiro_acesso,
    garantir_cargos_permanentes as startup_garantir_cargos_permanentes,
    migrar_funcoes_legadas_para_perfis as startup_migrar_funcoes_legadas_para_perfis,
)
from app.helpers import (
    _client_ip,
    _coletar_dashboard_analytics,
    _is_login_rate_limited,
    _normalizar_texto,
    _parse_date_range,
    _register_login_attempt,
    get_funcionario_logado,
)
from app.user_messages import build_flash_message, flash_category_for_status, resolve_action
from app.services.assistente_service import LocalAIAssistant
from app.services.permissao_service import PermissaoService
from app.services.estoque_service import aplicar_movimentacao_estoque
from app.services.financeiro_operacional import aplicar_acao_fundo, criar_lancamento_financeiro, criar_solicitacao_fundo
from app.services.master_data import (
    normalize_cost_center,
    validate_employee_payload,
    validate_payment_options_configuration,
)
from app.services.rh_service import sincronizar_garcom_funcionario
from app.services.traceability import build_timeline
from app.constants import (
    API_FALLBACK_ACCESS_PAGES,
    ENDPOINT_TO_PAGINA,
    NIVEIS_ORGANOGRAMA,
    PAGINA_ENDPOINTS,
    PAGINAS_SISTEMA,
    PAGINAS_SISTEMA_MENU_ORDEM,
    ROLES_PERMITIDOS,
    TIPOS_MOVIMENTACAO_VALIDOS,
)
from app.utils.payment_config import (
    api_integrations_text_to_json,
    api_integrations_to_text,
    payment_options_to_text,
    payment_text_to_json,
)
from app.utils.validators import (
    normalizar_matricula as shared_normalizar_matricula,
    validar_cpf as shared_validar_cpf,
    validar_data as shared_validar_data,
)

# Informacoes do SystemLR
APP_NAME = 'SystemLR'
APP_VERSION = '1.0.0'
APP_DOMAIN = 'systemlr.com'

PRIMEIRO_ACESSO_EMAIL = ((os.environ.get('SYSTEMLR_BOOTSTRAP_ADMIN_EMAIL') or 'admin@systemlr.com').strip().lower())
PRIMEIRO_ACESSO_SENHA = os.environ.get('SYSTEMLR_BOOTSTRAP_ADMIN_PASSWORD') or ''
PRIMEIRO_ACESSO_NOME = 'Administrador SystemLR'

TIPO_NEGOCIO_PRESETS = {
    EmpresaConfig.TIPO_NEGOCIO_CONVENIENCIA: {
        EmpresaConfig.CANAL_OPERACAO_FISICO: {'atendimento_mesas_ativo': True, 'separacao_entrega_ativa': False},
        EmpresaConfig.CANAL_OPERACAO_ECOMMERCE: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
        EmpresaConfig.CANAL_OPERACAO_HIBRIDO: {'atendimento_mesas_ativo': True, 'separacao_entrega_ativa': True},
    },
    EmpresaConfig.TIPO_NEGOCIO_SUPERMERCADO: {
        EmpresaConfig.CANAL_OPERACAO_FISICO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': False},
        EmpresaConfig.CANAL_OPERACAO_ECOMMERCE: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
        EmpresaConfig.CANAL_OPERACAO_HIBRIDO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
    },
    EmpresaConfig.TIPO_NEGOCIO_FARMACIA: {
        EmpresaConfig.CANAL_OPERACAO_FISICO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': False},
        EmpresaConfig.CANAL_OPERACAO_ECOMMERCE: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
        EmpresaConfig.CANAL_OPERACAO_HIBRIDO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
    },
    EmpresaConfig.TIPO_NEGOCIO_MODA: {
        EmpresaConfig.CANAL_OPERACAO_FISICO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': False},
        EmpresaConfig.CANAL_OPERACAO_ECOMMERCE: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
        EmpresaConfig.CANAL_OPERACAO_HIBRIDO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
    },
    EmpresaConfig.TIPO_NEGOCIO_HOME_CENTER: {
        EmpresaConfig.CANAL_OPERACAO_FISICO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': False},
        EmpresaConfig.CANAL_OPERACAO_ECOMMERCE: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
        EmpresaConfig.CANAL_OPERACAO_HIBRIDO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
    },
    EmpresaConfig.TIPO_NEGOCIO_OUTRO: {
        EmpresaConfig.CANAL_OPERACAO_FISICO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': False},
        EmpresaConfig.CANAL_OPERACAO_ECOMMERCE: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
        EmpresaConfig.CANAL_OPERACAO_HIBRIDO: {'atendimento_mesas_ativo': False, 'separacao_entrega_ativa': True},
    },
}

app = create_base_app()
local_ai_assistant = None

# Inicializar banco de dados
bootstrap_app(app, db=db, extensions_module=extensions, register_cli_fn=register_cli)


def _garantir_cargos_permanentes():
    # Extraido para app/core/startup.py
    startup_garantir_cargos_permanentes()


def _carregar_paginas_json(valor_json):
    # Extraido para app/core/permissions.py
    return core_carregar_paginas_json(valor_json)


def _serializar_paginas_json(paginas):
    # Extraido para app/core/permissions.py
    return core_serializar_paginas_json(paginas)


def _expandir_paginas_relacionadas(paginas, bloqueadas=None):
    # Extraido para app/core/permissions.py
    return core_expandir_paginas_relacionadas(paginas, bloqueadas=bloqueadas)


def _paginas_perfil_acesso(perfil_acesso):
    # Extraido para app/core/permissions.py
    return core_paginas_perfil_acesso(perfil_acesso)


def _mapa_permissoes_personalizadas_funcionario(funcionario):
    # Extraido para app/core/permissions.py
    return core_mapa_permissoes_personalizadas_funcionario(funcionario)


def _paginas_efetivas_funcionario(funcionario):
    # Extraido para app/core/permissions.py
    return core_paginas_efetivas_funcionario(funcionario)


def _salvar_permissoes_funcionario(funcionario, paginas_selecionadas):
    # Extraido para app/core/permissions.py
    return core_salvar_permissoes_funcionario(funcionario, paginas_selecionadas)


def _migrar_funcoes_legadas_para_perfis():
    # Extraido para app/core/startup.py
    startup_migrar_funcoes_legadas_para_perfis(
        carregar_paginas_json=_carregar_paginas_json,
        serializar_paginas_json=_serializar_paginas_json,
    )


def _bootstrap_admin_configurado():
    # Extraido para app/core/startup.py
    return startup_bootstrap_admin_configurado(
        email=PRIMEIRO_ACESSO_EMAIL,
        senha=PRIMEIRO_ACESSO_SENHA,
    )


def _garantir_admin_primeiro_acesso():
    # Extraido para app/core/startup.py
    startup_garantir_admin_primeiro_acesso(
        app=app,
        email=PRIMEIRO_ACESSO_EMAIL,
        senha=PRIMEIRO_ACESSO_SENHA,
        nome=PRIMEIRO_ACESSO_NOME,
        gerar_numero_cadastro_unico=_gerar_numero_cadastro_unico,
        gerar_matricula_unica=_gerar_matricula_unica,
    )


def _runtime_schema_patches_enabled():
    # Extraido para app/core/runtime_patches.py
    return runtime_schema_patches_enabled(app=app)


# ============ STARTUP / RUNTIME PATCHES ============
apply_runtime_schema_patches(
    app=app,
    db=db,
    perfil_acesso_model=PerfilAcesso,
    processo_evento_model=ProcessoEvento,
    cliente_publico_model=ClientePublico,
    lancamento_financeiro_model=LancamentoFinanceiro,
    fundo_solicitacao_model=FundoSolicitacao,
    equipamento_movimentacao_model=EquipamentoMovimentacao,
    manutencao_equipamento_model=ManutencaoEquipamento,
    ordem_servico_model=OrdemServico,
    chamado_interno_model=ChamadoInterno,
    almoxarifado_atribuicao_model=AlmoxarifadoAtribuicao,
    assistente_local_feedback_model=AssistenteLocalFeedback,
    funcionario_estoques_table=funcionario_estoques,
)

with app.app_context():
    _garantir_cargos_permanentes()
    _migrar_funcoes_legadas_para_perfis()
    try:
        db.session.query(Produto).filter(
            Produto.imagem_path == 'img/placeholders/produto-sem-foto.svg'
        ).update(
            {Produto.imagem_path: 'img/placeholders/imgindisponivel.png'},
            synchronize_session=False
        )
        db.session.commit()
    except Exception:
        db.session.rollback()


def _destino_interno_seguro(destino):
    texto = (destino or '').strip()
    if not texto:
        return None
    parsed = urlparse(texto)
    if parsed.scheme or parsed.netloc:
        host_atual = urlparse(request.host_url).netloc
        if parsed.netloc != host_atual:
            return None
        caminho = parsed.path or '/'
        if not caminho.startswith('/'):
            caminho = f'/{caminho}'
        return f'{caminho}?{parsed.query}' if parsed.query else caminho
    if not texto.startswith('/') or texto.startswith('//'):
        return None
    return texto


def _redirect_interno_seguro(destino, fallback):
    return redirect(_destino_interno_seguro(destino) or fallback)


def _normalizar_matricula(valor):
    return shared_normalizar_matricula(valor)


def _normalizar_numero_cadastro(valor):
    try:
        numero = int(valor)
    except (TypeError, ValueError):
        return None
    return numero if numero > 0 else None


def _normalizar_cpf(valor):
    return shared_validar_cpf(valor)


def _normalizar_estado(valor):
    estado = (valor or '').strip().upper()
    return estado[:2] if estado else None


def _normalizar_codigo_identificacao(valor, fallback=None, maxlen=12):
    codigo = re.sub(r'[^A-Z0-9]+', '', (valor or '').strip().upper())
    if not codigo:
        codigo = (fallback or '').strip().upper()
    return codigo[:maxlen] if codigo else None


def _estoques_permitidos_ids_funcionario(funcionario):
    if not funcionario:
        return None
    ids = set()
    if getattr(funcionario, 'estoque_principal_id', None):
        ids.add(funcionario.estoque_principal_id)
    for estoque in getattr(funcionario, 'estoques_permitidos', []) or []:
        if getattr(estoque, 'id', None):
            ids.add(estoque.id)
    return ids


def _codigo_filial_matricula(funcionario):
    estoque_principal = getattr(funcionario, 'estoque_principal', None)
    if not estoque_principal and funcionario:
        ids = _estoques_permitidos_ids_funcionario(funcionario) or set()
        if ids:
            estoque_principal = Estoque.query.filter(Estoque.id.in_(ids)).order_by(Estoque.nome.asc()).first()
    codigo_base = getattr(estoque_principal, 'codigo_filial', None) if estoque_principal else None
    return _normalizar_codigo_identificacao(codigo_base, fallback='GERAL', maxlen=10) or 'GERAL'


def _gerar_numero_cadastro_unico(funcionario):
    funcionario_id = getattr(funcionario, 'id', funcionario)
    numero_atual = _normalizar_numero_cadastro(getattr(funcionario, 'numero_cadastro', None))
    if numero_atual and not Funcionario.query.filter(
        Funcionario.numero_cadastro == numero_atual,
        Funcionario.id != funcionario_id,
    ).first():
        return numero_atual

    numero_base = _normalizar_numero_cadastro(getattr(funcionario, 'id', None))
    if numero_base and not Funcionario.query.filter(
        Funcionario.numero_cadastro == numero_base,
        Funcionario.id != funcionario_id,
    ).first():
        return numero_base

    maior_numero = db.session.query(db.func.max(Funcionario.numero_cadastro)).scalar() or 0
    candidata = max(maior_numero + 1, numero_base or 1)
    while Funcionario.query.filter(
        Funcionario.numero_cadastro == candidata,
        Funcionario.id != funcionario_id,
    ).first():
        candidata += 1
    return candidata


def _gerar_matricula_padrao(funcionario):
    numero_cadastro = _normalizar_numero_cadastro(getattr(funcionario, 'numero_cadastro', None))
    if not numero_cadastro:
        numero_cadastro = _gerar_numero_cadastro_unico(funcionario)
    return f'{_codigo_filial_matricula(funcionario)}-{numero_cadastro:06d}'


def _gerar_matricula_unica(funcionario):
    funcionario_id = getattr(funcionario, 'id', funcionario)
    matricula_base = _gerar_matricula_padrao(funcionario)
    if not Funcionario.query.filter(
        db.func.lower(Funcionario.matricula) == matricula_base.lower(),
        Funcionario.id != funcionario_id,
    ).first():
        return matricula_base
    sufixo = 1
    while True:
        candidata = f'{matricula_base}-{sufixo:02d}'
        existe = Funcionario.query.filter(
            db.func.lower(Funcionario.matricula) == candidata.lower(),
            Funcionario.id != funcionario_id,
        ).first()
        if not existe:
            return candidata
        sufixo += 1


def _garantir_matriculas_funcionarios():
    houve_mudanca = False
    for funcionario in Funcionario.query.order_by(Funcionario.id.asc()).all():
        numero_cadastro_atual = _normalizar_numero_cadastro(getattr(funcionario, 'numero_cadastro', None))
        numero_cadastro_esperado = _gerar_numero_cadastro_unico(funcionario)
        if numero_cadastro_atual != numero_cadastro_esperado:
            funcionario.numero_cadastro = numero_cadastro_esperado
            houve_mudanca = True
        matricula_atual = _normalizar_matricula(funcionario.matricula)
        matricula_esperada = _gerar_matricula_unica(funcionario)
        if matricula_atual != matricula_esperada:
            funcionario.matricula = matricula_esperada
            houve_mudanca = True
    if houve_mudanca:
        db.session.commit()


with app.app_context():
    try:
        _garantir_admin_primeiro_acesso()
        _garantir_matriculas_funcionarios()
    except Exception:
        db.session.rollback()


def _role_para_cargo_padrao(role):
    mapa = {
        'admin': 'Administrador',
        'gerente': 'Gerente',
        'caixa': 'Caixa',
        'operador': 'Operador',
        'garcom': 'Garcom',
    }
    return mapa.get((role or '').strip().lower(), 'Operador')


def _role_para_nivel_organograma(role):
    mapa = {
        'admin': 'Diretoria',
        'gerente': 'Gerencia',
        'caixa': 'Operacao',
        'operador': 'Operacao',
        'garcom': 'Operacao',
    }
    return mapa.get((role or '').strip().lower(), 'Operacao')


def _deve_exigir_superior_hierarquico(role, nivel_organograma):
    role_normalizado = (role or '').strip().lower()
    nivel_normalizado = (nivel_organograma or '').strip().lower()
    if role_normalizado == 'admin':
        return False
    if nivel_normalizado == 'diretoria':
        return False
    return True


def _normalizar_campo_organograma(valor):
    texto = (valor or '').strip()
    return texto or None


def _listar_cadastros_organograma():
    departamentos = sorted({
        (f.departamento or '').strip()
        for f in Funcionario.query.filter(Funcionario.departamento.isnot(None)).all()
        if (f.departamento or '').strip()
    }, key=str.lower)
    times = sorted({
        (f.time_nome or '').strip()
        for f in Funcionario.query.filter(Funcionario.time_nome.isnot(None)).all()
        if (f.time_nome or '').strip()
    }, key=str.lower)
    return departamentos, times


def _primeiro_funcionario_id():
    primeiro = Funcionario.query.order_by(Funcionario.id.asc()).first()
    return primeiro.id if primeiro else None


def _listar_estoques_para_vinculo_funcionario():
    return Estoque.query.order_by(Estoque.nome.asc()).all()


def _resolver_vinculos_estoque_funcionario(request_obj):
    restricao_estoques_ativa = (request_obj.form.get('restricao_estoques_ativa') == 'on')
    estoque_principal_id = request_obj.form.get('estoque_principal_id', type=int)
    estoques_ids = set()
    for valor in request_obj.form.getlist('estoques_permitidos_ids'):
        try:
            estoques_ids.add(int(valor))
        except (TypeError, ValueError):
            continue

    if estoque_principal_id:
        estoques_ids.add(estoque_principal_id)

    estoques_cadastrados = _listar_estoques_para_vinculo_funcionario()
    estoques_por_id = {estoque.id: estoque for estoque in estoques_cadastrados}

    if estoque_principal_id and estoque_principal_id not in estoques_por_id:
        return None, 'Estoque principal inválido.'

    ids_invalidos = [estoque_id for estoque_id in estoques_ids if estoque_id not in estoques_por_id]
    if ids_invalidos:
        return None, 'Há estoques/lojas inválidos no cadastro do colaborador.'

    if restricao_estoques_ativa and not estoques_ids:
        return None, 'Selecione ao menos um estoque/loja quando a restrição de acesso estiver ativa.'

    estoques_permitidos = [
        estoques_por_id[estoque_id]
        for estoque_id in sorted(estoques_ids, key=lambda item: (estoques_por_id[item].nome or '').lower())
    ]
    return {
        'restricao_estoques_ativa': restricao_estoques_ativa,
        'estoque_principal': estoques_por_id.get(estoque_principal_id),
        'estoques_permitidos': estoques_permitidos,
    }, None


def _estoques_contexto_base_ids_funcionario(funcionario):
    if not funcionario or funcionario.role == 'admin' or not getattr(funcionario, 'restricao_estoques_ativa', False):
        return None
    return _estoques_permitidos_ids_funcionario(funcionario) or set()


def _estoques_contexto_disponiveis(funcionario, apenas_ativos=True):
    query = Estoque.query
    ids = _estoques_contexto_base_ids_funcionario(funcionario)
    if ids is not None:
        if not ids:
            return []
        query = query.filter(Estoque.id.in_(ids))
    if apenas_ativos:
        query = query.filter(Estoque.ativo.is_(True))
    return query.order_by(Estoque.nome.asc()).all()


def _estoque_contexto_selecionado_id(funcionario):
    valor = session.get('estoque_contexto_id')
    if valor in (None, '', 'all'):
        return None
    try:
        estoque_id = int(valor)
    except (TypeError, ValueError):
        session.pop('estoque_contexto_id', None)
        return None

    disponiveis_ids = {estoque.id for estoque in _estoques_contexto_disponiveis(funcionario)}
    if estoque_id not in disponiveis_ids:
        session.pop('estoque_contexto_id', None)
        return None
    return estoque_id


def _paginas_permitidas_para_funcionario(funcionario):
    # Extraido para app/core/permissions.py
    return core_paginas_permitidas_para_funcionario(funcionario)


def _paginas_ordenadas_menu():
    # Extraido para app/core/menu.py
    return core_paginas_ordenadas_menu()


def _menu_agrupado_para_paginas(paginas_permitidas):
    # Extraido para app/core/menu.py
    return core_menu_agrupado_para_paginas(paginas_permitidas)


def _url_for_safe(endpoint, **values):
    # Mantido por compatibilidade; delega para app/core/menu.py
    from app.core.menu import url_for_safe as core_url_for_safe
    return core_url_for_safe(endpoint, **values)


def _item_menu_interno(*, label, endpoint, page_keys, current_page_key=None, visible=True):
    # Mantido por compatibilidade; delega para app/core/menu.py
    from app.core.menu import item_menu_interno as core_item_menu_interno
    return core_item_menu_interno(
        label=label,
        endpoint=endpoint,
        page_keys=page_keys,
        current_page_key=current_page_key,
        visible=visible,
    )


def _menu_navegacao_principal(funcionario, empresa_config=None, atendimento_mesas_ativo=True, current_page_key=None):
    # Extraido para app/core/menu.py
    return core_menu_navegacao_principal(
        funcionario,
        resolver_paginas_permitidas=_paginas_permitidas_para_funcionario,
        empresa_config=empresa_config,
        atendimento_mesas_ativo=atendimento_mesas_ativo,
        current_page_key=current_page_key,
    )


def _registrar_debug_menu(funcionario, paginas_permitidas, menu_agrupado=None):
    # Extraido para app/core/menu.py
    return core_registrar_debug_menu(
        funcionario,
        paginas_permitidas,
        menu_agrupado,
        resolver_menu_agrupado=_menu_agrupado_para_paginas,
    )


def _tabela_existe(nome_tabela):
    try:
        return inspect(db.engine).has_table(nome_tabela)
    except Exception:
        return False


def _extrair_permissoes_padrao_form():
    paginas_enviadas = set(request.form.getlist('paginas'))
    paginas_validas = set(PAGINAS_SISTEMA.keys())
    permissoes_padrao = set(paginas_enviadas.intersection(paginas_validas))

    setores_totais = set(request.form.getlist('setores_totais'))
    for secao_nome, itens in _paginas_ordenadas_menu():
        if secao_nome in setores_totais:
            permissoes_padrao.update(pagina_key for pagina_key, _ in itens)

    return sorted(permissoes_padrao)


def funcionario_tem_acesso(funcionario, endpoint):
    permissao_service = PermissaoService(resolver_paginas=_paginas_efetivas_funcionario)
    return permissao_service.tem_acesso(
        funcionario,
        endpoint,
        is_api_request=request.path.startswith('/api/'),
    )


def _resumir_payload_requisicao():
    chaves_sensiveis = {
        'senha',
        'confirmacao_senha',
        'password',
        'cpf',
        'rg',
        'celular',
        'telefone',
        'cep',
        'endereco',
        'token',
        'csrf_token',
    }

    def _is_sensitive_key(chave):
        chave_norm = (chave or '').strip().lower()
        return any(item in chave_norm for item in chaves_sensiveis)

    def _flatten_payload(payload, prefix=''):
        if isinstance(payload, dict):
            for chave, valor in payload.items():
                chave_composta = f'{prefix}.{chave}' if prefix else str(chave)
                if _is_sensitive_key(chave_composta) or _is_sensitive_key(chave):
                    continue
                yield from _flatten_payload(valor, chave_composta)
            return
        if isinstance(payload, list):
            for idx, item in enumerate(payload):
                chave_composta = f'{prefix}[{idx}]' if prefix else f'[{idx}]'
                yield from _flatten_payload(item, chave_composta)
            return
        yield prefix, payload

    dados = {}
    try:
        if request.is_json:
            payload = request.get_json(silent=True) or {}
            if isinstance(payload, dict):
                dados = payload
        else:
            dados = request.form.to_dict(flat=True)
    except Exception:
        dados = {}

    resumo = []
    for chave, valor in _flatten_payload(dados):
        if not chave or _is_sensitive_key(chave):
            continue
        texto_valor = str(valor)
        if len(texto_valor) > 80:
            texto_valor = f'{texto_valor[:77]}...'
        resumo.append(f'{chave}={texto_valor}')
        if len(resumo) >= 10:
            break

    return '; '.join(resumo)


def registrar_evento_auditoria(*, funcionario=None, acao='acao', entidade=None, detalhes='', status_code=None):
    try:
        funcionario_id = None
        funcionario_nome = None
        funcionario_email = None
        funcionario_role = None
        if funcionario:
            funcionario_id = funcionario.id
            funcionario_nome = funcionario.nome
            funcionario_email = funcionario.email
            funcionario_role = funcionario.role

        evento = AuditoriaEvento(
            funcionario_id=funcionario_id,
            funcionario_nome=funcionario_nome,
            funcionario_email=funcionario_email,
            funcionario_role=funcionario_role,
            metodo=request.method,
            endpoint=request.endpoint,
            rota=request.path,
            acao=acao,
            entidade=entidade,
            detalhes=detalhes,
            status_code=status_code,
            ip=(request.headers.get('X-Forwarded-For') or request.remote_addr),
        )
        db.session.add(evento)
        db.session.commit()
    except Exception:
        db.session.rollback()


def _normalizar_tipo_negocio(valor):
    tipo_negocio = (valor or EmpresaConfig.TIPO_NEGOCIO_CONVENIENCIA).strip().lower()
    if tipo_negocio not in EmpresaConfig.TIPOS_NEGOCIO_VALIDOS:
        return EmpresaConfig.TIPO_NEGOCIO_CONVENIENCIA
    return tipo_negocio


def _normalizar_canal_operacao(valor):
    canal_operacao = (valor or EmpresaConfig.CANAL_OPERACAO_HIBRIDO).strip().lower()
    if canal_operacao not in EmpresaConfig.CANAIS_OPERACAO_VALIDOS:
        return EmpresaConfig.CANAL_OPERACAO_HIBRIDO
    return canal_operacao


def _normalizar_cor_hex(valor, fallback):
    valor = (valor or '').strip()
    if not valor:
        return fallback
    if len(valor) == 7 and valor.startswith('#'):
        validos = '0123456789abcdefABCDEF'
        if all(c in validos for c in valor[1:]):
            return valor.lower()
    return fallback


def _normalizar_horario_hhmm(valor):
    texto = (valor or '').strip()
    if not texto:
        return None
    if not re.fullmatch(r'\d{2}:\d{2}', texto):
        return None
    try:
        hora_txt, minuto_txt = texto.split(':', 1)
        hora = int(hora_txt)
        minuto = int(minuto_txt)
    except (TypeError, ValueError):
        return None
    if hora < 0 or hora > 23 or minuto < 0 or minuto > 59:
        return None
    return f'{hora:02d}:{minuto:02d}'


def _humanizar_rotulo_endpoint(valor):
    partes = [parte for parte in str(valor or '').split('_') if parte]
    if not partes:
        return 'Tela Atual'
    rotulo = ' '.join(partes).title()
    substituicoes = {
        'Pdv': 'PDV',
        'Rh': 'RH',
        'Cd': 'CD',
        'Ids': 'IDs',
        'Api': 'API',
        'Qr': 'QR',
    }
    for origem, destino in substituicoes.items():
        rotulo = rotulo.replace(origem, destino)
    return rotulo


def _titulo_tela_atual():
    endpoint = (request.endpoint or '').split('.')[-1]
    if not endpoint:
        return None, None

    pagina = ENDPOINT_TO_PAGINA.get(endpoint)
    secao_menu = None
    if pagina:
        for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
            if pagina in secao_paginas:
                secao_menu = secao_nome
                break

    titulos_especiais = {
        'dashboard': 'Meu Perfil',
        'boas_vindas': 'Home Operacional',
        'gestao_negocio': 'Central do Negócio',
        'financeiro': 'Visão Financeira',
        'financeiro_lancamentos': 'Lançamentos Financeiros',
        'financeiro_fundos': 'Gestão Monetária e Fundos',
        'pdv': 'PDV',
        'configurar_ecommerce': 'Tema e Loja Online',
        'configurar_marketing_ecommerce': 'Promocoes e Campanhas',
        'configurar_ativacao_ecommerce': 'Ativação da Loja',
        'central_expedicao': 'Central de Expedição',
        'frota_expedicao': 'Frota Própria e Terceiros',
        'listar_separacao_entrega': 'Separação e Entrega',
        'listar_roteirizacao_entrega': 'Roteirização',
        'painel_expedicao': 'Painel em Tempo Real',
        'listar_transferencias_estoque': 'Histórico de Transferências',
        'transferir_armazenamento': 'Nova Transferência entre Lojas e CDs',
        'listar_movimentacoes': 'Entradas e Saídas Internas',
        'nova_movimentacao': 'Nova Entrada ou Saída Interna',
        'movimentacao_rapida': 'Movimentação Rápida Interna',
        'listar_almoxarifado': 'Central de Almoxarifado',
        'nova_atribuicao_almoxarifado': 'Nova Atribuição de Almoxarifado',
        'listar_recebimentos_fornecedor': 'Central de Recebimentos',
        'novo_recebimento_fornecedor': 'Novo Recebimento',
        'conferir_recebimento_fornecedor': 'Conferir Recebimento',
        'armazenar_recebimento_fornecedor': 'Armazenar Recebimento',
        'listar_funcoes_rh': 'Cargos',
        'listar_perfis_rh': 'Perfis de Acesso',
        'central_ajuda': 'Guia do Sistema',
        'detalhe_ajuda': 'Ajuda',
        'login': 'Login',
        'registro': 'Registro',
    }
    if endpoint in titulos_especiais:
        return secao_menu, titulos_especiais[endpoint]

    prefixos = (
        ('listar_', ''),
        ('novo_', 'Novo '),
        ('nova_', 'Nova '),
        ('criar_', 'Criar '),
        ('editar_', 'Editar '),
        ('deletar_', 'Excluir '),
        ('detalhes_', 'Detalhes de '),
        ('detalhe_', 'Detalhe de '),
        ('visualizar_', 'Visualizar '),
        ('configurar_', 'Configurar '),
        ('organograma_', 'Organograma '),
        ('indicadores_', 'Indicadores '),
        ('historico_', 'Histórico de '),
        ('abrir_', 'Abrir '),
        ('fechar_', 'Fechar '),
        ('imprimir_', 'Imprimir '),
        ('preview_', 'Preview '),
    )
    for prefixo, inicio in prefixos:
        if endpoint.startswith(prefixo):
            return secao_menu, f'{inicio}{_humanizar_rotulo_endpoint(endpoint[len(prefixo):])}'.strip()

    if pagina and pagina in PAGINAS_SISTEMA:
        return secao_menu, PAGINAS_SISTEMA[pagina]
    return secao_menu, _humanizar_rotulo_endpoint(endpoint)


def _carregar_json_lista(valor_json):
    if not valor_json:
        return []
    try:
        dados = json.loads(valor_json)
        if isinstance(dados, list):
            return dados
    except Exception:
        pass
    return []


def _normalizar_json_lista_para_texto(valor_json):
    linhas = []
    for item in _carregar_json_lista(valor_json):
        texto = ''
        if isinstance(item, str):
            texto = item.strip()
        elif isinstance(item, dict):
            for chave in ('nome', 'veiculo', 'empresa', 'titulo', 'descricao', 'texto', 'label'):
                valor = item.get(chave)
                if isinstance(valor, str) and valor.strip():
                    texto = valor.strip()
                    break
            if not texto:
                for valor in item.values():
                    if isinstance(valor, str) and valor.strip():
                        texto = valor.strip()
                        break
            if not texto:
                texto = json.dumps(item, ensure_ascii=False)
        elif item is not None:
            texto = str(item).strip()

        if texto:
            linhas.append(texto)
    return linhas


def _normalizar_linhas_configuracao(valor_texto, tamanho_max=200):
    linhas = []
    for linha in (valor_texto or '').splitlines():
        item = linha.strip()
        if not item:
            continue
        if len(item) > tamanho_max:
            item = item[:tamanho_max]
        linhas.append(item)
    return linhas


def _count_recebimentos_by_status_safe(status):
    try:
        inspector = inspect(db.engine)
        if not inspector.has_table('recebimentos_fornecedor'):
            return 0
        colunas = {col['name'] for col in inspector.get_columns('recebimentos_fornecedor')}
        if 'status' not in colunas:
            return 0
        resultado = db.session.execute(
            text("SELECT COUNT(*) FROM recebimentos_fornecedor WHERE status = :status"),
            {'status': status},
        ).scalar()
        return int(resultado or 0)
    except Exception:
        return 0


def _oldest_pending_storage_recebimento_safe():
    try:
        inspector = inspect(db.engine)
        if not inspector.has_table('recebimentos_fornecedor'):
            return None
        colunas = {col['name'] for col in inspector.get_columns('recebimentos_fornecedor')}
        if 'status' not in colunas:
            return None

        order_parts = []
        if 'conferido_em' in colunas:
            order_parts.append('conferido_em ASC')
        if 'criado_em' in colunas:
            order_parts.append('criado_em ASC')
        if not order_parts:
            order_parts.append('id ASC')

        select_cols = ['id']
        if 'conferido_em' in colunas:
            select_cols.append('conferido_em')
        if 'criado_em' in colunas:
            select_cols.append('criado_em')

        sql = (
            f"SELECT {', '.join(select_cols)} "
            "FROM recebimentos_fornecedor "
            "WHERE status = :status "
            f"ORDER BY {', '.join(order_parts)} "
            "LIMIT 1"
        )
        return db.session.execute(
            text(sql),
            {'status': RecebimentoFornecedor.STATUS_AGUARDANDO_ARMAZENAGEM},
        ).mappings().first()
    except Exception:
        return None


def _parse_datetime_local(valor):
    valor = (valor or '').strip()
    if not valor:
        return None
    try:
        return datetime.strptime(valor, '%Y-%m-%dT%H:%M')
    except ValueError:
        return None


def _parse_date_iso(valor):
    return shared_validar_data(valor)


def _periodo_datetime_ativo(inicio, fim, agora=None):
    agora = agora or datetime.utcnow()
    if inicio and agora < inicio:
        return False
    if fim and agora > fim:
        return False
    return True


def _periodo_data_ativo(inicio, fim, hoje=None):
    hoje = hoje or datetime.utcnow().date()
    if inicio and hoje < inicio:
        return False
    if fim and hoje > fim:
        return False
    return True


def _aplicar_preset_negocio(empresa):
    if not empresa:
        return
    preset_tipo = TIPO_NEGOCIO_PRESETS.get(empresa.tipo_negocio or EmpresaConfig.TIPO_NEGOCIO_CONVENIENCIA, {})
    preset = preset_tipo.get(empresa.canal_operacao or EmpresaConfig.CANAL_OPERACAO_HIBRIDO, {})
    for chave, valor in preset.items():
        setattr(empresa, chave, valor)


# ============ ROTAS - AUTENTICACAO ============

register_auth_module(app, {
    'login_required': login_required,
    '_limit': _limit,
    '_client_ip': _client_ip,
    '_is_login_rate_limited': _is_login_rate_limited,
    '_register_login_attempt': _register_login_attempt,
    'get_funcionario_logado': get_funcionario_logado,
    '_normalizar_texto': _normalizar_texto,
    '_normalizar_matricula': _normalizar_matricula,
    '_normalizar_cpf': _normalizar_cpf,
    '_normalizar_estado': _normalizar_estado,
    '_normalizar_campo_organograma': _normalizar_campo_organograma,
    '_parse_date_iso': _parse_date_iso,
    '_role_para_cargo_padrao': _role_para_cargo_padrao,
    '_role_para_nivel_organograma': _role_para_nivel_organograma,
    '_gerar_numero_cadastro_unico': _gerar_numero_cadastro_unico,
    '_gerar_matricula_unica': _gerar_matricula_unica,
    '_listar_cadastros_organograma': _listar_cadastros_organograma,
    '_paginas_permitidas_para_funcionario': _paginas_permitidas_para_funcionario,
    'sincronizar_garcom_funcionario': sincronizar_garcom_funcionario,
    'registrar_evento_auditoria': registrar_evento_auditoria,
    '_bootstrap_admin_configurado': _bootstrap_admin_configurado,
    'PRIMEIRO_ACESSO_EMAIL': PRIMEIRO_ACESSO_EMAIL,
    'ROLES_PERMITIDOS': ROLES_PERMITIDOS,
    'NIVEIS_ORGANOGRAMA': NIVEIS_ORGANOGRAMA,
    'extensions': extensions,
})


@app.route('/operacao/contexto-filial', methods=['POST'])
@login_required
def atualizar_contexto_filial():
    funcionario = get_funcionario_logado()
    valor = (request.form.get('estoque_contexto_id') or '').strip().lower()
    destino = request.form.get('next') or request.referrer

    if valor in ('', 'all', 'tudo'):
        session.pop('estoque_contexto_id', None)
        return _redirect_interno_seguro(destino, url_for('boas_vindas'))

    try:
        estoque_id = int(valor)
    except (TypeError, ValueError):
        session.pop('estoque_contexto_id', None)
        flash('Selecione uma filial/estoque valido.', 'warning')
        return _redirect_interno_seguro(destino, url_for('boas_vindas'))

    estoques_disponiveis = {estoque.id for estoque in _estoques_contexto_disponiveis(funcionario)}
    if estoque_id not in estoques_disponiveis:
        session.pop('estoque_contexto_id', None)
        flash('A filial/estoque selecionado nao esta disponivel para este colaborador.', 'warning')
        return _redirect_interno_seguro(destino, url_for('boas_vindas'))

    session['estoque_contexto_id'] = estoque_id
    return _redirect_interno_seguro(destino, url_for('boas_vindas'))


# ============ ROTAS - SISTEMA ============

@app.route('/')
def index():
    empresa = EmpresaConfig.query.first()
    if empresa and empresa.ecommerce_ativo is False:
        if session.get('funcionario_id'):
            return redirect(url_for('boas_vindas'))
        return redirect(url_for('login'))
    cliente_publico = None
    cliente_publico_id = session.get('site_cliente_id')
    if cliente_publico_id:
        query_cliente_publico = ClientePublico.query
        inspector = inspect(db.engine)
        tabelas = set(inspector.get_table_names())
        if 'clientes_enderecos' in tabelas:
            query_cliente_publico = query_cliente_publico.options(selectinload(ClientePublico.enderecos))
        if 'clientes_favoritos' in tabelas:
            query_cliente_publico = query_cliente_publico.options(selectinload(ClientePublico.favoritos))
        cliente_publico = query_cliente_publico.get(cliente_publico_id)
    resumo_carrinho = obter_resumo_carrinho_site(cliente_publico)
    produtos_disponiveis = Produto.query.filter(
        Produto.ativo.is_(True),
        Produto.status_disponibilidade.in_(Produto.STATUS_DISPONIBILIDADE_ONLINE_EQUIVALENTES)
    )

    produtos_destaque = produtos_disponiveis.order_by(
        Produto.atualizado_em.desc(),
        Produto.criado_em.desc()
    ).limit(12).all()

    if not produtos_destaque:
        produtos_destaque = Produto.query.filter_by(ativo=True).order_by(
            Produto.atualizado_em.desc(),
            Produto.criado_em.desc()
        ).limit(12).all()

    agora = datetime.utcnow()
    hoje = agora.date()
    banners_ativos = []
    for banner in _carregar_json_lista(empresa.ecom_banners_json if empresa else None):
        inicio = _parse_datetime_local(banner.get('inicio_em'))
        fim = _parse_datetime_local(banner.get('fim_em'))
        if not _periodo_datetime_ativo(inicio, fim, agora=agora):
            continue
        if not banner.get('ativo', True):
            continue
        image_path = (banner.get('image_path') or '').strip()
        if not image_path:
            continue
        banners_ativos.append({
            'titulo': (banner.get('titulo') or '').strip(),
            'subtitulo': (banner.get('subtitulo') or '').strip(),
            'link_url': (banner.get('link_url') or '').strip(),
            'link_label': (banner.get('link_label') or '').strip(),
            'image_path': image_path,
        })

    if not banners_ativos and empresa and empresa.ecom_banner_path:
        banners_ativos.append({
            'titulo': empresa.ecom_titulo_banner or 'Destaque da loja',
            'subtitulo': empresa.ecom_subtitulo_banner or '',
            'link_url': '#ofertas',
            'link_label': 'Ver ofertas',
            'image_path': empresa.ecom_banner_path,
        })

    campanhas_ativas = []
    for campanha in _carregar_json_lista(empresa.ecom_campanhas_json if empresa else None):
        inicio = _parse_date_iso(campanha.get('inicio_em'))
        fim = _parse_date_iso(campanha.get('fim_em'))
        if not _periodo_data_ativo(inicio, fim, hoje=hoje):
            continue
        if not campanha.get('ativo', True):
            continue
        nome = (campanha.get('nome') or '').strip()
        texto = (campanha.get('texto') or '').strip()
        if not nome and not texto:
            continue
        campanhas_ativas.append({'nome': nome, 'texto': texto})

    campanha_principal = None
    if campanhas_ativas:
        campanha_principal = campanhas_ativas[0].get('texto') or campanhas_ativas[0].get('nome')
    if not campanha_principal and empresa:
        campanha_principal = empresa.ecom_texto_promocao

    def _normalizar_chave_categoria_imagem(valor):
        texto = unicodedata.normalize('NFKD', (valor or '').strip())
        texto = ''.join(ch for ch in texto if not unicodedata.combining(ch))
        texto = re.sub(r'[^a-zA-Z0-9]+', ' ', texto).strip().lower()
        return texto

    def _mapear_imagens_categoria_existentes():
        relative_dir = os.path.join('uploads', 'categorias')
        absolute_dir = os.path.join(app.static_folder, relative_dir)
        imagens = {}
        if not os.path.isdir(absolute_dir):
            return imagens

        for nome_arquivo in sorted(os.listdir(absolute_dir)):
            caminho_absoluto = os.path.join(absolute_dir, nome_arquivo)
            if not os.path.isfile(caminho_absoluto):
                continue
            _, ext = os.path.splitext(nome_arquivo.lower())
            if ext not in {'.png', '.jpg', '.jpeg', '.webp', '.gif'}:
                continue
            chave = _normalizar_chave_categoria_imagem(
                os.path.splitext(nome_arquivo)[0].replace('_', ' ').replace('-', ' ')
            )
            if not chave:
                continue
            imagens.setdefault(chave, os.path.join(relative_dir, nome_arquivo).replace('\\', '/'))
        return imagens

    def _resolver_imagem_categoria_vitrine(categoria, imagens_existentes):
        caminho_salvo = (getattr(categoria, 'imagem_path', None) or '').strip().replace('\\', '/')
        if caminho_salvo:
            return caminho_salvo
        return imagens_existentes.get(_normalizar_chave_categoria_imagem(getattr(categoria, 'nome', '')))

    categorias_vitrine = []
    imagens_categoria_existentes = _mapear_imagens_categoria_existentes()
    categorias_home = Categoria.query.order_by(Categoria.nome.asc()).all()
    for categoria in categorias_home:
        produtos_categoria = Produto.query.filter(
            Produto.categoria_id == categoria.id,
            Produto.ativo.is_(True),
            Produto.status_disponibilidade.in_(Produto.STATUS_DISPONIBILIDADE_ONLINE_EQUIVALENTES)
        ).order_by(
            Produto.atualizado_em.desc(),
            Produto.criado_em.desc()
        ).limit(8).all()
        if not produtos_categoria:
            continue
        categorias_vitrine.append({
            'categoria': categoria,
            'categoria_imagem_path': _resolver_imagem_categoria_vitrine(categoria, imagens_categoria_existentes),
            'produtos': produtos_categoria,
            'ancora': f'categoria-{categoria.id}',
        })

    return render_template(
        'public/home_varejo.html',
        app_name=APP_NAME,
        produtos_destaque=produtos_destaque,
        carrinho_site=resumo_carrinho,
        total_produtos_ativos=Produto.query.filter_by(ativo=True).count(),
        banners_ativos=banners_ativos,
        campanhas_ativas=campanhas_ativas,
        campanha_principal=campanha_principal,
        categorias_vitrine=categorias_vitrine,
        cliente_publico=cliente_publico,
    )


@app.route('/dashboard', methods=['GET', 'POST'])
@login_required
def dashboard():
    funcionario = get_funcionario_logado()
    if not funcionario:
        flash('Sessao invalida. Faca login novamente.', 'danger')
        return redirect(url_for('login'))

    if request.method == 'POST':
        acao = (request.form.get('acao') or '').strip().lower()
        if acao == 'atualizar_perfil':
            remover_imagem_perfil = (request.form.get('remover_imagem_perfil') == 'on')
            arquivo_imagem = request.files.get('imagem_perfil')
            novo_imagem_perfil_path = None

            if (arquivo_imagem and arquivo_imagem.filename) or remover_imagem_perfil:
                if not funcionario.permitir_editar_imagem_perfil:
                    flash('Seu perfil nao possui permissao para alterar a imagem.', 'danger')
                    return redirect(url_for('dashboard'))

            imagem_anterior = funcionario.imagem_perfil_path
            allowed_ext = {'.png', '.jpg', '.jpeg', '.webp', '.gif'}
            if arquivo_imagem and arquivo_imagem.filename:
                _, ext = os.path.splitext(arquivo_imagem.filename.lower())
                if ext not in allowed_ext:
                    flash('Formato de imagem invalido. Use PNG, JPG, JPEG, WEBP ou GIF.', 'danger')
                    return redirect(url_for('dashboard'))
                nome_base = secure_filename((funcionario.nome or f'usuario_{funcionario.id}').strip()) or f'usuario_{funcionario.id}'
                relative_dir = os.path.join('uploads', 'usuarios')
                absolute_dir = os.path.join(app.static_folder, relative_dir)
                os.makedirs(absolute_dir, exist_ok=True)
                image_name = f'{nome_base}_{funcionario.id}_perfil{ext}'
                novo_imagem_perfil_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_imagem_perfil_path)
                arquivo_imagem.save(absolute_path)
                funcionario.imagem_perfil_path = novo_imagem_perfil_path
            elif remover_imagem_perfil:
                funcionario.imagem_perfil_path = None

            funcionario.numero_cadastro = funcionario.numero_cadastro or _gerar_numero_cadastro_unico(funcionario)
            funcionario.matricula = funcionario.matricula or _gerar_matricula_unica(funcionario)
            db.session.commit()
            if novo_imagem_perfil_path and imagem_anterior and imagem_anterior != novo_imagem_perfil_path:
                caminho_anterior = os.path.join(app.static_folder, imagem_anterior)
                if os.path.exists(caminho_anterior):
                    os.remove(caminho_anterior)
            if remover_imagem_perfil and imagem_anterior:
                caminho_anterior = os.path.join(app.static_folder, imagem_anterior)
                if os.path.exists(caminho_anterior):
                    os.remove(caminho_anterior)
            flash('Perfil atualizado com sucesso.', 'success')
            return redirect(url_for('dashboard'))

        if acao == 'atualizar_senha':
            senha_atual = request.form.get('senha_atual', '')
            nova_senha = request.form.get('nova_senha', '')
            confirmacao_senha = request.form.get('confirmacao_senha', '')

            if not senha_atual or not nova_senha:
                flash('Informe senha atual e nova senha.', 'danger')
                return redirect(url_for('dashboard'))
            if not funcionario.check_password(senha_atual):
                flash('Senha atual invalida.', 'danger')
                return redirect(url_for('dashboard'))
            if len(nova_senha) < 6:
                flash('A nova senha deve ter no minimo 6 caracteres.', 'danger')
                return redirect(url_for('dashboard'))
            if nova_senha != confirmacao_senha:
                flash('Confirmacao de senha nao confere.', 'danger')
                return redirect(url_for('dashboard'))

            funcionario.set_password(nova_senha)
            funcionario.senha_provisoria = False
            db.session.commit()
            flash('Senha alterada com sucesso.', 'success')
            return redirect(url_for('dashboard'))

        if acao == 'atualizar_preferencias':
            receber_alertas = request.form.get('receber_alertas') == 'on'
            funcionario.receber_alertas = receber_alertas
            db.session.commit()
            flash('Preferencias salvas com sucesso.', 'success')
            return redirect(url_for('dashboard'))

        flash('Acao de configuracao invalida.', 'danger')
        return redirect(url_for('dashboard'))

    agora = datetime.utcnow()
    hora_atual = agora.hour
    if hora_atual < 12:
        saudacao = 'Bom dia'
    elif hora_atual < 18:
        saudacao = 'Boa tarde'
    else:
        saudacao = 'Boa noite'

    atividades_recentes = AuditoriaEvento.query.filter(
        AuditoriaEvento.funcionario_id == funcionario.id
    ).order_by(AuditoriaEvento.criado_em.desc()).limit(8).all()

    paginas_liberadas = len(_paginas_efetivas_funcionario(funcionario))

    alertas_usuario = []
    if funcionario.receber_alertas:
        if not (funcionario.matricula or '').strip():
            alertas_usuario.append('Defina sua matricula para rastreabilidade no RH.')
        if not (funcionario.cpf or '').strip():
            alertas_usuario.append('Cadastre seu CPF para completar os dados pessoais.')
        if not (funcionario.celular or '').strip():
            alertas_usuario.append('Cadastre seu celular para facilitar contato operacional.')
        if not (funcionario.endereco or '').strip():
            alertas_usuario.append('Complete seu endereco para manter o cadastro atualizado.')
        if not (funcionario.departamento or '').strip():
            alertas_usuario.append('Informe seu departamento para manter o organograma atualizado.')
        if not (funcionario.time_nome or '').strip():
            alertas_usuario.append('Informe seu time/squad para relatórios de RH e produtividade.')
    if not alertas_usuario:
        alertas_usuario.append('Nenhuma pendencia pessoal critica no momento.')

    secoes_acesso = []
    paginas_permitidas = _paginas_permitidas_para_funcionario(funcionario)
    for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        paginas = [
            PAGINAS_SISTEMA.get(pagina_key, pagina_key)
            for pagina_key in secao_paginas
            if pagina_key in paginas_permitidas and pagina_key in PAGINAS_SISTEMA
        ]
        if paginas:
            secoes_acesso.append({'secao': secao_nome, 'paginas': paginas})

    return render_template(
        'dashboard/index.html',
        funcionario=funcionario,
        saudacao=saudacao,
        data_hora_atual=agora,
        paginas_liberadas=paginas_liberadas,
        total_alertas=len(alertas_usuario),
        alertas_usuario=alertas_usuario,
        atividades_recentes=atividades_recentes,
        secoes_acesso=secoes_acesso,
    )


@app.route('/gestao-negocio')
@require_role('admin', 'gerente')
def gestao_negocio():
    empresa = EmpresaConfig.query.first()
    total_funcionarios = Funcionario.query.count()
    funcionarios_ativos = Funcionario.query.filter_by(ativo=True).count()
    produtos_ativos = Produto.query.filter_by(ativo=True).count()
    categorias_total = Categoria.query.count()
    caixas_abertas = Caixa.query.filter_by(aberto=True).count()
    cargos_rh_ativos = FuncaoRH.query.filter_by(ativo=True).count()
    perfis_acesso_ativos = PerfilAcesso.query.filter_by(ativo=True).count()
    pedidos_abertos = Pedido.query.filter(Pedido.status.in_(['aberto', 'em_preparo', 'entregue'])).count()
    pendencias_armazenagem = _count_recebimentos_by_status_safe(
        RecebimentoFornecedor.STATUS_AGUARDANDO_ARMAZENAGEM
    )
    recebimento_armazenagem_mais_antigo = _oldest_pending_storage_recebimento_safe()
    if recebimento_armazenagem_mais_antigo:
        referencia_pendencia_armazenagem = (
            recebimento_armazenagem_mais_antigo.get('conferido_em')
            or recebimento_armazenagem_mais_antigo.get('criado_em')
        )
        detalhe_pendencia_armazenagem = (
            f'{pendencias_armazenagem} recebimento(s) aguardando armazenagem. '
            f'Mais antigo desde {referencia_pendencia_armazenagem.strftime("%d/%m/%Y %H:%M") if referencia_pendencia_armazenagem else "data indisponivel"}.'
        )
    else:
        detalhe_pendencia_armazenagem = 'Nenhum recebimento aguardando armazenagem.'

    checklist_config = [
        {
            'titulo': 'Dados principais da empresa',
            'ok': bool(empresa and (empresa.nome_fantasia or '').strip()),
            'detalhe_ok': 'Cadastro principal preenchido.',
            'detalhe_pendente': 'Defina nome fantasia e dados da empresa.',
            'url': url_for('editar_empresa'),
            'acao': 'Configurar empresa',
        },
        {
            'titulo': 'Canal de operacao e entrega',
            'ok': bool(empresa and empresa.canal_operacao and empresa.separacao_entrega_ativa is not False),
            'detalhe_ok': 'Canal e fluxo de entrega ativos.',
            'detalhe_pendente': 'Revise canal de operacao e habilite separacao/entrega.',
            'url': url_for('editar_empresa') + '#config-entrega',
            'acao': 'Revisar operacao',
        },
        {
            'titulo': 'Ativacao do e-commerce',
            'ok': bool(empresa and empresa.ecommerce_ativo is not False),
            'detalhe_ok': 'Loja online liberada para operacao.',
            'detalhe_pendente': 'Habilite a loja online antes de divulgar o canal para clientes.',
            'url': url_for('configurar_ativacao_ecommerce'),
            'acao': 'Ativar loja online',
        },
        {
            'titulo': 'Tema e identidade visual da loja',
            'ok': bool(empresa and (getattr(empresa, 'ecom_titulo_banner', None) or '').strip()),
            'detalhe_ok': 'Tema, banners e comunicacao principal definidos.',
            'detalhe_pendente': 'Ajuste tema, banner, promocoes e identidade da vitrine.',
            'url': url_for('configurar_ecommerce'),
            'acao': 'Configurar visual da loja',
        },
        {
            'titulo': 'Cargos e perfis de acesso',
            'ok': cargos_rh_ativos > 0 and perfis_acesso_ativos > 0,
            'detalhe_ok': f'{cargos_rh_ativos} cargo(s) e {perfis_acesso_ativos} perfil(is) ativo(s).',
            'detalhe_pendente': 'Cadastre cargos da empresa e perfis de acesso padrao para a equipe.',
            'url': url_for('listar_funcoes_rh'),
            'acao': 'Configurar RH',
        },
        {
            'titulo': 'Equipe cadastrada',
            'ok': funcionarios_ativos > 0,
            'detalhe_ok': f'{funcionarios_ativos} colaborador(es) ativo(s).',
            'detalhe_pendente': 'Cadastre colaboradores e niveis da hierarquia.',
            'url': url_for('listar_funcionarios'),
            'acao': 'Gerir funcionarios',
        },
        {
            'titulo': 'Pendencias de armazenagem',
            'ok': pendencias_armazenagem == 0,
            'detalhe_ok': 'Fila de armazenagem em dia.',
            'detalhe_pendente': detalhe_pendencia_armazenagem,
            'url': url_for(
                'listar_recebimentos_fornecedor',
                status=RecebimentoFornecedor.STATUS_AGUARDANDO_ARMAZENAGEM,
            ),
            'acao': 'Abrir fila de armazenagem',
        },
    ]
    pendencias_criticas = [item for item in checklist_config if not item['ok']]

    atalhos_dono = [
        {'titulo': 'Configuracoes da Empresa', 'descricao': 'Dados, operacao, entrega e servicos.', 'url': url_for('editar_empresa')},
        {'titulo': 'Ativacao da Loja Online', 'descricao': 'Liga ou desliga a vitrine publica do e-commerce.', 'url': url_for('configurar_ativacao_ecommerce')},
        {'titulo': 'Tema e Loja Online', 'descricao': 'Cores, banners, checkout e rodape da vitrine.', 'url': url_for('configurar_ecommerce')},
        {'titulo': 'Promocoes e Campanhas', 'descricao': 'Campanhas, cupons e temas sazonais para datas especiais.', 'url': url_for('configurar_marketing_ecommerce')},
        {'titulo': 'Cargos da Equipe', 'descricao': 'Estrutura organizacional usada nos cadastros e no organograma.', 'url': url_for('listar_funcoes_rh')},
        {'titulo': 'Perfis de Acesso', 'descricao': 'Conjuntos padrao de paginas para aplicar por colaborador.', 'url': url_for('listar_perfis_rh')},
        {'titulo': 'Central de Expedicao', 'descricao': 'Separacao, rotas, etiquetas e frota.', 'url': url_for('central_expedicao')},
        {'titulo': 'Lancamentos Financeiros', 'descricao': 'Controle monetario e exportacao contabil.', 'url': url_for('financeiro_lancamentos')},
        {'titulo': 'Guia do Sistema', 'descricao': 'Passo a passo para equipe e lideres.', 'url': url_for('central_ajuda')},
    ]

    return render_template(
        'sistema/gestao_negocio.html',
        empresa=empresa,
        total_funcionarios=total_funcionarios,
        funcionarios_ativos=funcionarios_ativos,
        produtos_ativos=produtos_ativos,
        categorias_total=categorias_total,
        caixas_abertas=caixas_abertas,
        pedidos_abertos=pedidos_abertos,
        pendencias_armazenagem=pendencias_armazenagem,
        detalhe_pendencia_armazenagem=detalhe_pendencia_armazenagem,
        checklist_config=checklist_config,
        pendencias_criticas=pendencias_criticas,
        atalhos_dono=atalhos_dono,
    )


@app.route('/financeiro')
@login_required
def financeiro():
    inicio_periodo, fim_periodo, data_inicial_str, data_final_str = _parse_date_range(
        request.args.get('data_inicial'),
        request.args.get('data_final'),
        default_days=7
    )
    analytics = _coletar_dashboard_analytics(inicio_periodo, fim_periodo)
    cache = extensions.cache
    cache_key = f'view:financeiro:{inicio_periodo.strftime("%Y%m%d")}:{fim_periodo.strftime("%Y%m%d")}'
    if cache is not None:
        html = cache.get(cache_key)
        if html is not None:
            return html

    html = render_template(
        'financeiro/index.html',
        periodo_dias=analytics['periodo_dias'],
        data_inicial=data_inicial_str,
        data_final=data_final_str,
        pedidos_periodo_total=analytics['pedidos_periodo_total'],
        faturamento_periodo=analytics['faturamento_periodo'],
        faturamento_periodo_anterior=analytics['faturamento_periodo_anterior'],
        crescimento_receita_pct=analytics['crescimento_receita_pct'],
        faturamento_hoje=analytics['faturamento_hoje'],
        receita_media_dia=analytics['receita_media_dia'],
        pedidos_media_dia=analytics['pedidos_media_dia'],
        ticket_medio_periodo=analytics['ticket_medio_periodo'],
        cmv_periodo=analytics['cmv_periodo'],
        lucro_bruto_periodo=analytics['lucro_bruto_periodo'],
        margem_bruta_pct=analytics['margem_bruta_pct'],
        despesas_operacionais_periodo=analytics['despesas_operacionais_periodo'],
        movimentacoes_financeiras_excluidas_periodo=analytics['movimentacoes_financeiras_excluidas_periodo'],
        categorias_excluidas_resultado=analytics['categorias_excluidas_resultado'],
        ajustes_financeiros_periodo=analytics['ajustes_financeiros_periodo'],
        resultado_operacional_periodo=analytics['resultado_operacional_periodo'],
        margem_operacional_pct=analytics['margem_operacional_pct'],
        pedidos_abertos=analytics['pedidos_abertos'],
        pedidos_cancelados_periodo=analytics['pedidos_cancelados_periodo'],
        valor_cancelado_periodo=analytics['valor_cancelado_periodo'],
        taxa_cancelamento_pct=analytics['taxa_cancelamento_pct'],
        metodo_mais_usado=analytics['metodo_mais_usado'],
        concentracao_top_pagamento_pct=analytics['concentracao_top_pagamento_pct'],
        vendas_periodo=analytics['vendas_periodo'],
        top_produtos_vendidos=analytics['top_produtos_vendidos'],
        pedidos_por_status=analytics['pedidos_por_status'],
        top_clientes=analytics['top_clientes'],
        desempenho_garcons=analytics['desempenho_garcons'],
        desempenho_caixas=analytics['desempenho_caixas'],
        metodos_pagamento=analytics['metodos_pagamento']
    )
    if cache is not None:
        cache.set(cache_key, html, timeout=300)
    return html


@app.route('/financeiro/fundos', methods=['GET', 'POST'])
@require_role('admin', 'gerente', 'caixa')
def financeiro_fundos():
    funcionario = get_funcionario_logado()
    if request.method == 'POST':
        acao = (request.form.get('acao') or 'solicitar').strip().lower()
        fundo_id = request.form.get('fundo_id', type=int)

        if acao == 'solicitar':
            valor_texto = (request.form.get('valor') or '').strip().replace('.', '').replace(',', '.')
            try:
                valor = float(valor_texto)
            except Exception:
                valor = 0.0

            try:
                fundo = criar_solicitacao_fundo(
                    tipo=request.form.get('tipo'),
                    descricao=request.form.get('descricao'),
                    categoria=request.form.get('categoria'),
                    centro_custo=normalize_cost_center(request.form.get('centro_custo')),
                    referencia_documento=request.form.get('referencia_documento'),
                    valor=valor,
                    solicitado_por_id=(funcionario.id if funcionario else None),
                    actor=funcionario,
                )
                db.session.commit()
                flash(f'Solicitacao de fundo #{fundo.id} registrada.', 'success')
            except Exception as exc:
                db.session.rollback()
                flash(f'Erro ao registrar solicitacao: {str(exc)}', 'danger')
            return redirect(url_for('financeiro_fundos'))

        fundo = FundoSolicitacao.query.get_or_404(fundo_id)
        if acao == 'aprovar':
            if funcionario.role not in {'admin', 'gerente'}:
                flash('Somente admin/gerente pode aprovar solicitacoes.', 'danger')
                return redirect(url_for('financeiro_fundos'))
            try:
                aplicar_acao_fundo(
                    fundo,
                    acao='aprovar',
                    actor=funcionario,
                )
                db.session.commit()
                flash(f'Solicitacao #{fundo.id} aprovada.', 'success')
            except AppError as exc:
                db.session.rollback()
                flash(str(exc), 'warning')
            return redirect(url_for('financeiro_fundos'))

        if acao == 'rejeitar':
            if funcionario.role not in {'admin', 'gerente'}:
                flash('Somente admin/gerente pode rejeitar solicitacoes.', 'danger')
                return redirect(url_for('financeiro_fundos'))
            try:
                aplicar_acao_fundo(
                    fundo,
                    acao='rejeitar',
                    actor=funcionario,
                    motivo_rejeicao=request.form.get('motivo_rejeicao'),
                )
                db.session.commit()
                flash(f'Solicitacao #{fundo.id} rejeitada.', 'warning')
            except AppError as exc:
                db.session.rollback()
                flash(str(exc), 'warning')
            return redirect(url_for('financeiro_fundos'))

        if acao == 'liberar':
            if funcionario.role not in {'admin', 'gerente'}:
                flash('Somente admin/gerente pode liberar fundos.', 'danger')
                return redirect(url_for('financeiro_fundos'))
            try:
                aplicar_acao_fundo(
                    fundo,
                    acao='liberar',
                    actor=funcionario,
                )
                db.session.commit()
                flash(f'Fundo #{fundo.id} liberado e lancamento financeiro gerado.', 'success')
            except AppError as exc:
                db.session.rollback()
                flash(str(exc), 'warning')
            except Exception as exc:
                db.session.rollback()
                flash(f'Erro ao liberar fundo: {str(exc)}', 'danger')
            return redirect(url_for('financeiro_fundos'))

    status_filtro = (request.args.get('status') or '').strip().lower()
    tipo_filtro = (request.args.get('tipo') or '').strip().lower()
    query = FundoSolicitacao.query.order_by(FundoSolicitacao.solicitado_em.desc())
    if status_filtro in FundoSolicitacao.STATUS_VALIDOS:
        query = query.filter(FundoSolicitacao.status == status_filtro)
    if tipo_filtro in FundoSolicitacao.TIPOS_VALIDOS:
        query = query.filter(FundoSolicitacao.tipo == tipo_filtro)
    fundos = query.limit(300).all()

    total_solicitado = sum(float(item.valor or 0.0) for item in fundos if item.status == FundoSolicitacao.STATUS_SOLICITADA)
    total_liberado = sum(float(item.valor or 0.0) for item in fundos if item.status == FundoSolicitacao.STATUS_LIBERADA)
    total_aprovado = sum(float(item.valor or 0.0) for item in fundos if item.status == FundoSolicitacao.STATUS_APROVADA)

    return render_template(
        'financeiro/fundos.html',
        fundos=fundos,
        status_validos=FundoSolicitacao.STATUS_VALIDOS,
        tipos_validos=FundoSolicitacao.TIPOS_VALIDOS,
        status_filtro=status_filtro,
        tipo_filtro=tipo_filtro,
        total_solicitado=total_solicitado,
        total_aprovado=total_aprovado,
        total_liberado=total_liberado,
    )


@app.route('/financeiro/lancamentos', methods=['GET', 'POST'])
@require_role('admin', 'gerente', 'caixa')
def financeiro_lancamentos():
    funcionario = get_funcionario_logado()
    if request.method == 'POST':
        tipo = (request.form.get('tipo') or '').strip().lower()
        descricao = (request.form.get('descricao') or '').strip()
        categoria = (request.form.get('categoria') or '').strip() or None
        referencia_documento = (request.form.get('referencia_documento') or '').strip() or None
        try:
            centro_custo = normalize_cost_center(request.form.get('centro_custo'))
        except ValidationError as exc:
            flash(str(exc), 'danger')
            return redirect(url_for('financeiro_lancamentos'))
        data_competencia = _parse_date_iso(request.form.get('data_competencia')) or datetime.utcnow().date()
        incluir_contabilidade = (request.form.get('incluir_contabilidade') == 'on')

        if tipo not in LancamentoFinanceiro.TIPOS:
            flash('Tipo de lancamento invalido.', 'danger')
            return redirect(url_for('financeiro_lancamentos'))
        if not descricao:
            flash('Descricao e obrigatoria.', 'danger')
            return redirect(url_for('financeiro_lancamentos'))

        valor_texto = (request.form.get('valor') or '').strip().replace('.', '').replace(',', '.')
        valor = 0.0
        if valor_texto:
            try:
                valor = float(valor_texto)
            except ValueError:
                flash('Valor invalido.', 'danger')
                return redirect(url_for('financeiro_lancamentos'))

        produto = None
        produto_id = request.form.get('produto_id', type=int)
        quantidade = request.form.get('quantidade', type=float)
        if tipo == LancamentoFinanceiro.TIPO_CONSUMO_PROPRIO:
            if not produto_id:
                flash('Selecione um produto para consumo proprio.', 'danger')
                return redirect(url_for('financeiro_lancamentos'))
            produto = Produto.query.get(produto_id)
            if not produto:
                flash('Produto nao encontrado para consumo proprio.', 'danger')
                return redirect(url_for('financeiro_lancamentos'))
            if not quantidade or quantidade <= 0:
                flash('Informe quantidade valida para consumo proprio.', 'danger')
                return redirect(url_for('financeiro_lancamentos'))
            valor = round((produto.preco_custo or 0) * quantidade, 2)
        elif tipo in {LancamentoFinanceiro.TIPO_DESPESA, LancamentoFinanceiro.TIPO_RECEITA}:
            if valor <= 0:
                flash('Valor deve ser maior que zero.', 'danger')
                return redirect(url_for('financeiro_lancamentos'))
        elif tipo == LancamentoFinanceiro.TIPO_AJUSTE and valor == 0:
            flash('Ajuste deve ter valor diferente de zero.', 'danger')
            return redirect(url_for('financeiro_lancamentos'))

        try:
            criar_lancamento_financeiro(
                tipo=tipo,
                categoria=categoria,
                descricao=descricao,
                valor=valor,
                data_competencia=data_competencia,
                incluir_contabilidade=incluir_contabilidade,
                referencia_documento=referencia_documento,
                centro_custo=centro_custo,
                produto=produto,
                produto_id=produto_id,
                quantidade=quantidade,
                criado_por_id=(funcionario.id if funcionario else None),
                actor=funcionario,
            )
            db.session.commit()
            flash('Lancamento financeiro registrado com sucesso.', 'success')
        except Exception as exc:
            db.session.rollback()
            flash(f'Erro ao registrar lancamento: {str(exc)}', 'danger')
        return redirect(url_for('financeiro_lancamentos'))

    data_inicial_txt = (request.args.get('data_inicial') or '').strip()
    data_final_txt = (request.args.get('data_final') or '').strip()
    tipo_filtro = (request.args.get('tipo') or '').strip().lower()
    somente_contabilidade = (request.args.get('somente_contabilidade') == '1')

    data_inicial = _parse_date_iso(data_inicial_txt)
    data_final = _parse_date_iso(data_final_txt)

    query = LancamentoFinanceiro.query.order_by(
        LancamentoFinanceiro.data_competencia.desc(),
        LancamentoFinanceiro.criado_em.desc()
    )
    if data_inicial:
        query = query.filter(LancamentoFinanceiro.data_competencia >= data_inicial)
    if data_final:
        query = query.filter(LancamentoFinanceiro.data_competencia <= data_final)
    if tipo_filtro in LancamentoFinanceiro.TIPOS:
        query = query.filter(LancamentoFinanceiro.tipo == tipo_filtro)
    if somente_contabilidade:
        query = query.filter(LancamentoFinanceiro.incluir_contabilidade.is_(True))

    lancamentos = query.all()
    total_receitas = sum((item.valor or 0.0) for item in lancamentos if item.tipo == LancamentoFinanceiro.TIPO_RECEITA)
    total_despesas = sum(
        (item.valor or 0.0)
        for item in lancamentos
        if item.tipo in {LancamentoFinanceiro.TIPO_DESPESA, LancamentoFinanceiro.TIPO_CONSUMO_PROPRIO}
    )
    total_ajustes = sum((item.valor or 0.0) for item in lancamentos if item.tipo == LancamentoFinanceiro.TIPO_AJUSTE)
    saldo_periodo = total_receitas - total_despesas + total_ajustes

    produtos = Produto.query.filter_by(ativo=True).order_by(Produto.nome.asc()).all()
    return render_template(
        'financeiro/lancamentos.html',
        lancamentos=lancamentos,
        produtos=produtos,
        tipos_lancamento=LancamentoFinanceiro.TIPOS,
        data_inicial=data_inicial_txt,
        data_final=data_final_txt,
        tipo_filtro=tipo_filtro,
        somente_contabilidade=somente_contabilidade,
        total_receitas=total_receitas,
        total_despesas=total_despesas,
        total_ajustes=total_ajustes,
        saldo_periodo=saldo_periodo,
    )


@app.route('/api/rastreabilidade/pedidos/<int:pedido_id>/timeline')
@require_role('admin', 'gerente', 'caixa', 'operador')
def api_timeline_pedido(pedido_id):
    pedido = Pedido.query.get_or_404(pedido_id)
    timeline = build_timeline(pedido=pedido)
    return jsonify({
        'pedido_id': pedido.id,
        'timeline': [
            {**item, 'quando': (item['quando'].isoformat() if item.get('quando') else None)}
            for item in timeline
        ],
    })


@app.route('/api/rastreabilidade/recebimentos/<int:recebimento_id>/timeline')
@require_role('admin', 'gerente', 'caixa', 'operador')
def api_timeline_recebimento(recebimento_id):
    recebimento = RecebimentoFornecedor.query.get_or_404(recebimento_id)
    timeline = build_timeline(recebimento=recebimento)
    return jsonify({
        'recebimento_id': recebimento.id,
        'timeline': [
            {**item, 'quando': (item['quando'].isoformat() if item.get('quando') else None)}
            for item in timeline
        ],
    })


@app.route('/financeiro/lancamentos/<int:lancamento_id>/marcar-enviado', methods=['POST'])
@require_role('admin', 'gerente')
def marcar_lancamento_enviado_contador(lancamento_id):
    lancamento = LancamentoFinanceiro.query.get_or_404(lancamento_id)
    try:
        lancamento.enviado_contador = True
        lancamento.enviado_em = datetime.utcnow()
        db.session.commit()
        flash('Lancamento marcado como enviado ao contador.', 'success')
    except Exception as exc:
        db.session.rollback()
        flash(f'Erro ao atualizar lancamento: {str(exc)}', 'danger')
    return redirect(url_for('financeiro_lancamentos'))


@app.route('/financeiro/lancamentos/exportar')
@require_role('admin', 'gerente', 'caixa')
def exportar_lancamentos_financeiros():
    data_inicial = _parse_date_iso(request.args.get('data_inicial'))
    data_final = _parse_date_iso(request.args.get('data_final'))
    tipo_filtro = (request.args.get('tipo') or '').strip().lower()
    somente_contabilidade = (request.args.get('somente_contabilidade', '1') == '1')

    query = LancamentoFinanceiro.query.order_by(
        LancamentoFinanceiro.data_competencia.asc(),
        LancamentoFinanceiro.id.asc()
    )
    if data_inicial:
        query = query.filter(LancamentoFinanceiro.data_competencia >= data_inicial)
    if data_final:
        query = query.filter(LancamentoFinanceiro.data_competencia <= data_final)
    if tipo_filtro in LancamentoFinanceiro.TIPOS:
        query = query.filter(LancamentoFinanceiro.tipo == tipo_filtro)
    if somente_contabilidade:
        query = query.filter(LancamentoFinanceiro.incluir_contabilidade.is_(True))

    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow([
        'id',
        'data_competencia',
        'tipo',
        'categoria',
        'descricao',
        'valor',
        'produto_codigo',
        'produto_nome',
        'quantidade',
        'referencia_documento',
        'centro_custo',
        'incluir_contabilidade',
        'enviado_contador',
        'enviado_em',
        'criado_em',
    ])
    for item in query.all():
        writer.writerow([
            item.id,
            item.data_competencia.isoformat() if item.data_competencia else '',
            item.tipo,
            item.categoria or '',
            item.descricao or '',
            f'{(item.valor or 0.0):.2f}',
            item.produto.codigo if item.produto else '',
            item.produto.nome if item.produto else '',
            f'{item.quantidade:.3f}' if item.quantidade is not None else '',
            item.referencia_documento or '',
            item.centro_custo or '',
            'sim' if item.incluir_contabilidade else 'nao',
            'sim' if item.enviado_contador else 'nao',
            item.enviado_em.strftime('%Y-%m-%d %H:%M:%S') if item.enviado_em else '',
            item.criado_em.strftime('%Y-%m-%d %H:%M:%S') if item.criado_em else '',
        ])

    csv_content = output.getvalue()
    output.close()
    nome_arquivo = f'lancamentos_financeiros_{datetime.utcnow().strftime("%Y%m%d_%H%M%S")}.csv'
    return Response(
        csv_content,
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename={nome_arquivo}'}
    )


@app.route('/financeiro/lancamentos/exportar-xlsx')
@require_role('admin', 'gerente', 'caixa')
def exportar_lancamentos_financeiros_xlsx():
    try:
        from openpyxl import Workbook
    except Exception:
        flash('Dependencia openpyxl nao instalada. Execute: pip install openpyxl', 'danger')
        return redirect(url_for('financeiro_lancamentos'))

    data_inicial = _parse_date_iso(request.args.get('data_inicial'))
    data_final = _parse_date_iso(request.args.get('data_final'))
    tipo_filtro = (request.args.get('tipo') or '').strip().lower()
    somente_contabilidade = (request.args.get('somente_contabilidade', '1') == '1')

    query = LancamentoFinanceiro.query.order_by(
        LancamentoFinanceiro.data_competencia.asc(),
        LancamentoFinanceiro.id.asc()
    )
    if data_inicial:
        query = query.filter(LancamentoFinanceiro.data_competencia >= data_inicial)
    if data_final:
        query = query.filter(LancamentoFinanceiro.data_competencia <= data_final)
    if tipo_filtro in LancamentoFinanceiro.TIPOS:
        query = query.filter(LancamentoFinanceiro.tipo == tipo_filtro)
    if somente_contabilidade:
        query = query.filter(LancamentoFinanceiro.incluir_contabilidade.is_(True))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lancamentos'
    ws.append([
        'id',
        'data_competencia',
        'tipo',
        'categoria',
        'descricao',
        'valor',
        'produto_codigo',
        'produto_nome',
        'quantidade',
        'referencia_documento',
        'centro_custo',
        'incluir_contabilidade',
        'enviado_contador',
        'enviado_em',
        'criado_em',
    ])
    for item in query.all():
        ws.append([
            item.id,
            item.data_competencia.isoformat() if item.data_competencia else '',
            item.tipo,
            item.categoria or '',
            item.descricao or '',
            float(item.valor or 0.0),
            item.produto.codigo if item.produto else '',
            item.produto.nome if item.produto else '',
            float(item.quantidade) if item.quantidade is not None else None,
            item.referencia_documento or '',
            item.centro_custo or '',
            'sim' if item.incluir_contabilidade else 'nao',
            'sim' if item.enviado_contador else 'nao',
            item.enviado_em.strftime('%Y-%m-%d %H:%M:%S') if item.enviado_em else '',
            item.criado_em.strftime('%Y-%m-%d %H:%M:%S') if item.criado_em else '',
        ])

    arquivo = io.BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    nome_arquivo = f'lancamentos_financeiros_{datetime.utcnow().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return Response(
        arquivo.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={nome_arquivo}'}
    )


@app.route('/empresa', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def editar_empresa():
    empresa = EmpresaConfig.query.first()

    if request.method == 'POST':
        novo_logo_path = None
        novo_favicon_empresa_path = None
        novo_app_icon_path = None
        logo_anterior = (empresa.logo_path if empresa else None)
        favicon_empresa_anterior = (empresa.favicon_path if empresa else None)
        app_icon_anterior = (empresa.app_icon_path if empresa else None)
        try:
            if not empresa:
                empresa = EmpresaConfig()
                db.session.add(empresa)

            empresa.razao_social = request.form.get('razao_social', '').strip() or None
            empresa.nome_fantasia = request.form.get('nome_fantasia', '').strip() or None
            empresa.codigo_empresa = _normalizar_codigo_identificacao(
                request.form.get('codigo_empresa'),
                maxlen=10,
            ) or None
            empresa.cnpj = request.form.get('cnpj', '').strip() or None
            empresa.inscricao_estadual = request.form.get('inscricao_estadual', '').strip() or None
            empresa.telefone = request.form.get('telefone', '').strip() or None
            empresa.email = request.form.get('email', '').strip() or None
            empresa.endereco = request.form.get('endereco', '').strip() or None
            empresa.cidade = request.form.get('cidade', '').strip() or None
            empresa.estado = request.form.get('estado', '').strip().upper() or None
            empresa.cep = request.form.get('cep', '').strip() or None
            empresa.mensagem_comprovante = request.form.get('mensagem_comprovante', '').strip() or None
            empresa.cardapio_titulo = request.form.get('cardapio_titulo', '').strip() or None
            empresa.cardapio_subtitulo = request.form.get('cardapio_subtitulo', '').strip() or None
            empresa.cardapio_mensagem = request.form.get('cardapio_mensagem', '').strip() or None
            empresa.cardapio_mostrar_imagem = (request.form.get('cardapio_mostrar_imagem') == 'on')
            empresa.cardapio_mostrar_descricao = (request.form.get('cardapio_mostrar_descricao') == 'on')
            empresa.tipo_negocio = _normalizar_tipo_negocio(request.form.get('tipo_negocio'))
            empresa.canal_operacao = _normalizar_canal_operacao(request.form.get('canal_operacao'))
            empresa.atendimento_mesas_ativo = (request.form.get('atendimento_mesas_ativo') == 'on')
            empresa.separacao_entrega_ativa = (request.form.get('separacao_entrega_ativa') == 'on')
            empresa.emissao_etiqueta_entrega_ativa = (request.form.get('emissao_etiqueta_entrega_ativa') == 'on')
            empresa.separacao_entrega_unir_vendas_off = (request.form.get('separacao_entrega_unir_vendas_off') == 'on')
            empresa.roteirizacao_entrega_ativa = (request.form.get('roteirizacao_entrega_ativa') == 'on')
            empresa.emissao_nota_entrega_ativa = (request.form.get('emissao_nota_entrega_ativa') == 'on')
            empresa.entrega_local_saida_padrao = request.form.get('entrega_local_saida_padrao', '').strip() or None
            empresa.entrega_veiculo_padrao = request.form.get('entrega_veiculo_padrao', '').strip() or None
            empresa.entrega_motorista_padrao = request.form.get('entrega_motorista_padrao', '').strip() or None
            horario_fechamento_roteirizacao_raw = request.form.get('entrega_horario_fechamento_roteirizacao', '')
            horario_fechamento_roteirizacao = _normalizar_horario_hhmm(horario_fechamento_roteirizacao_raw)
            if horario_fechamento_roteirizacao_raw.strip() and not horario_fechamento_roteirizacao:
                flash('Horario de fechamento para roteirizacao invalido. Use o formato HH:MM.', 'error')
                return redirect(url_for('editar_empresa') + '#config-entrega')
            empresa.entrega_horario_fechamento_roteirizacao = horario_fechamento_roteirizacao
            veiculos_linhas = _normalizar_linhas_configuracao(
                request.form.get('entrega_veiculos_cadastro', ''),
                tamanho_max=160
            )
            terceirizadas_linhas = _normalizar_linhas_configuracao(
                request.form.get('entrega_terceirizadas_cadastro', ''),
                tamanho_max=180
            )
            empresa.entrega_veiculos_json = json.dumps(veiculos_linhas, ensure_ascii=False) if veiculos_linhas else None
            empresa.entrega_terceirizadas_json = json.dumps(terceirizadas_linhas, ensure_ascii=False) if terceirizadas_linhas else None
            empresa.pagamentos_pdv_json = validate_payment_options_configuration(
                request.form.get('pagamentos_pdv_config', ''),
                channel='pdv',
            )
            empresa.integracoes_pdv_json = api_integrations_text_to_json(
                request.form.get('integracoes_pdv_config', '')
            )
            empresa.reposicao_loja_fisica_ativa = (request.form.get('reposicao_loja_fisica_ativa') == 'on')
            empresa.emissao_etiqueta_loja_ativa = (request.form.get('emissao_etiqueta_loja_ativa') == 'on')
            empresa.emissao_etiqueta_endereco_ativa = (request.form.get('emissao_etiqueta_endereco_ativa') == 'on')
            empresa.servicos_tecnicos_ativos = (request.form.get('servicos_tecnicos_ativos') == 'on')
            empresa.servico_montagem_instalacao_ativo = (request.form.get('servico_montagem_instalacao_ativo') == 'on')
            if request.form.get('aplicar_preset_operacao') == 'on':
                _aplicar_preset_negocio(empresa)

            remover_logo = (request.form.get('remover_logo') == 'on')
            remover_favicon_empresa = (request.form.get('remover_favicon') == 'on')
            remover_app_icon = (request.form.get('remover_app_icon') == 'on')
            arquivo_logo = request.files.get('logo')
            arquivo_favicon_empresa = request.files.get('favicon')
            arquivo_app_icon = request.files.get('app_icon')
            allowed_ext = {'.png', '.jpg', '.jpeg', '.webp', '.gif'}
            allowed_favicon_ext = {'.ico', '.png', '.svg', '.jpg', '.jpeg', '.webp'}
            allowed_app_icon_ext = {'.png', '.jpg', '.jpeg', '.webp'}
            if arquivo_logo and arquivo_logo.filename:
                _, ext = os.path.splitext(arquivo_logo.filename.lower())
                if ext not in allowed_ext:
                    flash('Formato de logo inválido. Use PNG, JPG, JPEG, WEBP ou GIF.', 'error')
                    return redirect(url_for('editar_empresa'))
                nome_empresa_base = secure_filename((empresa.nome_fantasia or empresa.razao_social or 'empresa').strip())
                if not nome_empresa_base:
                    nome_empresa_base = 'empresa'
                relative_dir = os.path.join('uploads', 'empresa')
                absolute_dir = os.path.join(app.static_folder, relative_dir)
                os.makedirs(absolute_dir, exist_ok=True)
                image_name = f'{nome_empresa_base}_logo{ext}'
                novo_logo_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_logo_path)
                arquivo_logo.save(absolute_path)
                empresa.logo_path = novo_logo_path
            elif remover_logo:
                empresa.logo_path = None

            if arquivo_favicon_empresa and arquivo_favicon_empresa.filename:
                _, ext = os.path.splitext(arquivo_favicon_empresa.filename.lower())
                if ext not in allowed_favicon_ext:
                    flash('Formato de favicon inválido. Use ICO, PNG, SVG, JPG, JPEG ou WEBP.', 'error')
                    return redirect(url_for('editar_empresa'))
                nome_empresa_base = secure_filename((empresa.nome_fantasia or empresa.razao_social or 'empresa').strip())
                if not nome_empresa_base:
                    nome_empresa_base = 'empresa'
                relative_dir = os.path.join('uploads', 'empresa')
                absolute_dir = os.path.join(app.static_folder, relative_dir)
                os.makedirs(absolute_dir, exist_ok=True)
                image_name = f'{nome_empresa_base}_favicon{ext}'
                novo_favicon_empresa_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_favicon_empresa_path)
                arquivo_favicon_empresa.save(absolute_path)
                empresa.favicon_path = novo_favicon_empresa_path
            elif remover_favicon_empresa:
                empresa.favicon_path = None

            if arquivo_app_icon and arquivo_app_icon.filename:
                _, ext = os.path.splitext(arquivo_app_icon.filename.lower())
                if ext not in allowed_app_icon_ext:
                    flash('Formato do ícone do app inválido. Use PNG, JPG, JPEG ou WEBP.', 'error')
                    return redirect(url_for('editar_empresa'))
                nome_empresa_base = secure_filename((empresa.nome_fantasia or empresa.razao_social or 'empresa').strip())
                if not nome_empresa_base:
                    nome_empresa_base = 'empresa'
                relative_dir = os.path.join('uploads', 'empresa')
                absolute_dir = os.path.join(app.static_folder, relative_dir)
                os.makedirs(absolute_dir, exist_ok=True)
                image_name = f'{nome_empresa_base}_app_icon{ext}'
                novo_app_icon_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_app_icon_path)
                arquivo_app_icon.save(absolute_path)
                empresa.app_icon_path = novo_app_icon_path
            elif remover_app_icon:
                empresa.app_icon_path = None

            qtd_maxima = request.form.get('cardapio_qtd_maxima', type=int)
            if qtd_maxima is None or qtd_maxima <= 0:
                qtd_maxima = 20
            empresa.cardapio_qtd_maxima = qtd_maxima
            if 'distribuicao_ativa' in request.form or 'modo_distribuicao_pedidos' in request.form:
                empresa.distribuicao_ativa = (request.form.get('distribuicao_ativa') == 'on')
                modo_distribuicao = (request.form.get('modo_distribuicao_pedidos') or 'round_robin').strip().lower()
                if modo_distribuicao not in {'round_robin', 'menos_pedidos', 'manual'}:
                    modo_distribuicao = 'round_robin'
                empresa.modo_distribuicao_pedidos = modo_distribuicao

            db.session.commit()
            _garantir_matriculas_funcionarios()
            if novo_logo_path and logo_anterior and logo_anterior != novo_logo_path:
                caminho_logo_anterior = os.path.join(app.static_folder, logo_anterior)
                if os.path.exists(caminho_logo_anterior):
                    os.remove(caminho_logo_anterior)
            if remover_logo and logo_anterior:
                caminho_logo_anterior = os.path.join(app.static_folder, logo_anterior)
                if os.path.exists(caminho_logo_anterior):
                    os.remove(caminho_logo_anterior)
            if novo_favicon_empresa_path and favicon_empresa_anterior and favicon_empresa_anterior != novo_favicon_empresa_path:
                caminho_favicon_empresa_anterior = os.path.join(app.static_folder, favicon_empresa_anterior)
                if os.path.exists(caminho_favicon_empresa_anterior):
                    os.remove(caminho_favicon_empresa_anterior)
            if remover_favicon_empresa and favicon_empresa_anterior:
                caminho_favicon_empresa_anterior = os.path.join(app.static_folder, favicon_empresa_anterior)
                if os.path.exists(caminho_favicon_empresa_anterior):
                    os.remove(caminho_favicon_empresa_anterior)
            if novo_app_icon_path and app_icon_anterior and app_icon_anterior != novo_app_icon_path:
                caminho_app_icon_anterior = os.path.join(app.static_folder, app_icon_anterior)
                if os.path.exists(caminho_app_icon_anterior):
                    os.remove(caminho_app_icon_anterior)
            if remover_app_icon and app_icon_anterior:
                caminho_app_icon_anterior = os.path.join(app.static_folder, app_icon_anterior)
                if os.path.exists(caminho_app_icon_anterior):
                    os.remove(caminho_app_icon_anterior)
            flash('Dados da empresa salvos com sucesso.', 'success')
            return redirect(url_for('editar_empresa'))
        except Exception as e:
            db.session.rollback()
            if novo_logo_path:
                caminho_novo_logo = os.path.join(app.static_folder, novo_logo_path)
                if os.path.exists(caminho_novo_logo):
                    os.remove(caminho_novo_logo)
            if novo_favicon_empresa_path:
                caminho_novo_favicon = os.path.join(app.static_folder, novo_favicon_empresa_path)
                if os.path.exists(caminho_novo_favicon):
                    os.remove(caminho_novo_favicon)
            if novo_app_icon_path:
                caminho_novo_app_icon = os.path.join(app.static_folder, novo_app_icon_path)
                if os.path.exists(caminho_novo_app_icon):
                    os.remove(caminho_novo_app_icon)
            flash(f'Erro ao salvar dados da empresa: {str(e)}', 'error')

    tipos_negocio = [
        (EmpresaConfig.TIPO_NEGOCIO_CONVENIENCIA, 'Loja de conveniência'),
        (EmpresaConfig.TIPO_NEGOCIO_SUPERMERCADO, 'Supermercado/mercearia'),
        (EmpresaConfig.TIPO_NEGOCIO_FARMACIA, 'Farmacia/perfumaria'),
        (EmpresaConfig.TIPO_NEGOCIO_MODA, 'Moda/acessorios'),
        (EmpresaConfig.TIPO_NEGOCIO_HOME_CENTER, 'Casa/construcao'),
        (EmpresaConfig.TIPO_NEGOCIO_OUTRO, 'Outro varejo'),
    ]
    canais_operacao = [
        (EmpresaConfig.CANAL_OPERACAO_FISICO, 'Somente loja física'),
        (EmpresaConfig.CANAL_OPERACAO_ECOMMERCE, 'Somente e-commerce'),
        (EmpresaConfig.CANAL_OPERACAO_HIBRIDO, 'Hibrido (fisico + e-commerce)'),
    ]
    veiculos_texto = '\n'.join(_normalizar_json_lista_para_texto(empresa.entrega_veiculos_json if empresa else None))
    terceirizadas_texto = '\n'.join(
        _normalizar_json_lista_para_texto(empresa.entrega_terceirizadas_json if empresa else None)
    )
    pagamentos_pdv_texto = payment_options_to_text(empresa.pagamentos_pdv_json if empresa else None, 'pdv')
    integracoes_pdv_texto = api_integrations_to_text(empresa.integracoes_pdv_json if empresa else None)
    return render_template(
        'sistema/empresa.html',
        empresa=empresa,
        tipos_negocio=tipos_negocio,
        canais_operacao=canais_operacao,
        veiculos_texto=veiculos_texto,
        terceirizadas_texto=terceirizadas_texto,
        pagamentos_pdv_texto=pagamentos_pdv_texto,
        integracoes_pdv_texto=integracoes_pdv_texto,
    )


@app.route('/ecommerce-config', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def configurar_ecommerce():
    empresa = EmpresaConfig.query.first()
    if not empresa:
        empresa = EmpresaConfig()
        db.session.add(empresa)
        db.session.commit()

    def _montar_slots_banners():
        slots = []
        origem = _carregar_json_lista(empresa.ecom_banners_json)
        for i in range(5):
            item = origem[i] if i < len(origem) and isinstance(origem[i], dict) else {}
            slots.append({
                'titulo': item.get('titulo') or '',
                'subtitulo': item.get('subtitulo') or '',
                'link_url': item.get('link_url') or '',
                'link_label': item.get('link_label') or '',
                'inicio_em': item.get('inicio_em') or '',
                'fim_em': item.get('fim_em') or '',
                'image_path': item.get('image_path') or '',
                'ativo': bool(item.get('ativo', False)),
            })
        return slots

    if request.method == 'POST':
        novo_banner_path = None
        novo_favicon_path = None
        novo_produto_placeholder_path = None
        banner_anterior = empresa.ecom_banner_path
        favicon_anterior = empresa.ecom_favicon_path
        produto_placeholder_anterior = empresa.ecom_produto_placeholder_path
        novos_arquivos = []
        arquivos_para_remover = []
        try:
            empresa.ecom_cor_primaria = _normalizar_cor_hex(
                request.form.get('ecom_cor_primaria'),
                empresa.ecom_cor_primaria or '#ff7848'
            )
            empresa.ecom_cor_secundaria = _normalizar_cor_hex(
                request.form.get('ecom_cor_secundaria'),
                empresa.ecom_cor_secundaria or '#ff5a2a'
            )
            empresa.ecom_card_bg = _normalizar_cor_hex(
                request.form.get('ecom_card_bg'),
                empresa.ecom_card_bg or '#ffffff'
            )
            empresa.ecom_titulo_banner = request.form.get('ecom_titulo_banner', '').strip() or None
            empresa.ecom_subtitulo_banner = request.form.get('ecom_subtitulo_banner', '').strip() or None
            empresa.ecom_footer_bg = _normalizar_cor_hex(
                request.form.get('ecom_footer_bg'),
                empresa.ecom_footer_bg or '#1f2b38'
            )
            empresa.ecom_footer_texto = request.form.get('ecom_footer_texto', '').strip() or None
            empresa.ecom_footer_contato = request.form.get('ecom_footer_contato', '').strip() or None
            empresa.ecom_footer_creditos = request.form.get('ecom_footer_creditos', '').strip() or None
            empresa.pagamentos_ecommerce_json = validate_payment_options_configuration(
                request.form.get('pagamentos_ecommerce_config', ''),
                channel='ecommerce',
            )
            empresa.integracoes_ecommerce_json = api_integrations_text_to_json(
                request.form.get('integracoes_ecommerce_config', '')
            )

            remover_banner = (request.form.get('remover_ecom_banner') == 'on')
            remover_favicon = (request.form.get('remover_ecom_favicon') == 'on')
            remover_produto_placeholder = (request.form.get('remover_ecom_produto_placeholder') == 'on')
            arquivo_banner = request.files.get('ecom_banner')
            arquivo_favicon = request.files.get('ecom_favicon')
            arquivo_produto_placeholder = request.files.get('ecom_produto_placeholder')
            allowed_ext = {'.png', '.jpg', '.jpeg', '.webp', '.gif'}
            allowed_favicon_ext = {'.ico', '.png', '.svg', '.jpg', '.jpeg', '.webp'}
            nome_empresa_base = secure_filename((empresa.nome_fantasia or empresa.razao_social or 'empresa').strip())
            if not nome_empresa_base:
                nome_empresa_base = 'empresa'
            relative_dir = os.path.join('uploads', 'ecommerce')
            absolute_dir = os.path.join(app.static_folder, relative_dir)
            os.makedirs(absolute_dir, exist_ok=True)
            if arquivo_banner and arquivo_banner.filename:
                _, ext = os.path.splitext(arquivo_banner.filename.lower())
                if ext not in allowed_ext:
                    flash('Formato de banner inválido. Use PNG, JPG, JPEG, WEBP ou GIF.', 'error')
                    return redirect(url_for('configurar_ecommerce'))
                image_name = f'{nome_empresa_base}_banner{ext}'
                novo_banner_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_banner_path)
                arquivo_banner.save(absolute_path)
                novos_arquivos.append(novo_banner_path)
                empresa.ecom_banner_path = novo_banner_path
            elif remover_banner:
                empresa.ecom_banner_path = None

            if arquivo_favicon and arquivo_favicon.filename:
                _, ext = os.path.splitext(arquivo_favicon.filename.lower())
                if ext not in allowed_favicon_ext:
                    flash('Formato de favicon inválido. Use ICO, PNG, SVG, JPG, JPEG ou WEBP.', 'error')
                    return redirect(url_for('configurar_ecommerce'))
                image_name = f'{nome_empresa_base}_favicon{ext}'
                novo_favicon_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_favicon_path)
                arquivo_favicon.save(absolute_path)
                novos_arquivos.append(novo_favicon_path)
                empresa.ecom_favicon_path = novo_favicon_path
            elif remover_favicon:
                empresa.ecom_favicon_path = None

            if arquivo_produto_placeholder and arquivo_produto_placeholder.filename:
                _, ext = os.path.splitext(arquivo_produto_placeholder.filename.lower())
                if ext not in allowed_ext:
                    flash('Formato da imagem padrão inválido. Use PNG, JPG, JPEG, WEBP ou GIF.', 'error')
                    return redirect(url_for('configurar_ecommerce'))
                image_name = f'{nome_empresa_base}_produto_padrao{ext}'
                novo_produto_placeholder_path = os.path.join(relative_dir, image_name).replace('\\', '/')
                absolute_path = os.path.join(app.static_folder, novo_produto_placeholder_path)
                arquivo_produto_placeholder.save(absolute_path)
                novos_arquivos.append(novo_produto_placeholder_path)
                empresa.ecom_produto_placeholder_path = novo_produto_placeholder_path
            elif remover_produto_placeholder:
                empresa.ecom_produto_placeholder_path = None

            banners_slots = []
            for idx in range(5):
                titulo = (request.form.get(f'banner_titulo_{idx}') or '').strip()
                subtitulo = (request.form.get(f'banner_subtitulo_{idx}') or '').strip()
                link_url = (request.form.get(f'banner_link_url_{idx}') or '').strip()
                link_label = (request.form.get(f'banner_link_label_{idx}') or '').strip()
                inicio_em = (request.form.get(f'banner_inicio_{idx}') or '').strip()
                fim_em = (request.form.get(f'banner_fim_{idx}') or '').strip()
                image_path_atual = (request.form.get(f'banner_path_{idx}') or '').strip() or None
                ativo = (request.form.get(f'banner_ativo_{idx}') == 'on')
                remover_slot = (request.form.get(f'remover_banner_{idx}') == 'on')
                arquivo_slot = request.files.get(f'banner_imagem_{idx}')
                image_path_slot = image_path_atual

                if arquivo_slot and arquivo_slot.filename:
                    _, ext = os.path.splitext(arquivo_slot.filename.lower())
                    if ext not in allowed_ext:
                        flash(f'Formato de imagem inválido no banner {idx + 1}.', 'error')
                        return redirect(url_for('configurar_ecommerce'))
                    nome_slot = f'{nome_empresa_base}_carrossel_{idx}_{int(datetime.utcnow().timestamp())}{ext}'
                    novo_slot_path = os.path.join(relative_dir, nome_slot).replace('\\', '/')
                    caminho_slot = os.path.join(app.static_folder, novo_slot_path)
                    arquivo_slot.save(caminho_slot)
                    novos_arquivos.append(novo_slot_path)
                    image_path_slot = novo_slot_path
                    if image_path_atual and image_path_atual != novo_slot_path:
                        arquivos_para_remover.append(image_path_atual)
                elif remover_slot:
                    if image_path_atual:
                        arquivos_para_remover.append(image_path_atual)
                    image_path_slot = None

                inicio_dt = _parse_datetime_local(inicio_em)
                fim_dt = _parse_datetime_local(fim_em)
                if fim_dt and inicio_dt and fim_dt < inicio_dt:
                    flash(f'Banner {idx + 1}: o fim da vigência não pode ser menor que o início.', 'error')
                    return redirect(url_for('configurar_ecommerce'))

                if not any([titulo, subtitulo, link_url, link_label, image_path_slot, inicio_em, fim_em]):
                    continue

                banners_slots.append({
                    'titulo': titulo,
                    'subtitulo': subtitulo,
                    'link_url': link_url,
                    'link_label': link_label,
                    'inicio_em': inicio_em,
                    'fim_em': fim_em,
                    'image_path': image_path_slot,
                    'ativo': ativo,
                })

            empresa.ecom_banners_json = json.dumps(banners_slots, ensure_ascii=False)

            placeholder_alvos = [None, '', 'img/placeholders/imgindisponivel.png']
            if produto_placeholder_anterior:
                placeholder_alvos.append(produto_placeholder_anterior)
            placeholder_alvos = list({item for item in placeholder_alvos})

            if novo_produto_placeholder_path:
                Produto.query.filter(
                    db.or_(
                        Produto.imagem_path.is_(None),
                        Produto.imagem_path == '',
                        Produto.imagem_path.in_(placeholder_alvos),
                    )
                ).update(
                    {Produto.imagem_path: novo_produto_placeholder_path},
                    synchronize_session=False
                )
            elif remover_produto_placeholder and produto_placeholder_anterior:
                Produto.query.filter(
                    db.or_(
                        Produto.imagem_path.is_(None),
                        Produto.imagem_path == '',
                        Produto.imagem_path == produto_placeholder_anterior,
                    )
                ).update(
                    {Produto.imagem_path: 'img/placeholders/imgindisponivel.png'},
                    synchronize_session=False
                )

            db.session.commit()
            if novo_banner_path and banner_anterior and banner_anterior != novo_banner_path:
                caminho_banner_anterior = os.path.join(app.static_folder, banner_anterior)
                if os.path.exists(caminho_banner_anterior):
                    os.remove(caminho_banner_anterior)
            if remover_banner and banner_anterior:
                caminho_banner_anterior = os.path.join(app.static_folder, banner_anterior)
                if os.path.exists(caminho_banner_anterior):
                    os.remove(caminho_banner_anterior)
            if novo_favicon_path and favicon_anterior and favicon_anterior != novo_favicon_path:
                caminho_favicon_anterior = os.path.join(app.static_folder, favicon_anterior)
                if os.path.exists(caminho_favicon_anterior):
                    os.remove(caminho_favicon_anterior)
            if remover_favicon and favicon_anterior:
                caminho_favicon_anterior = os.path.join(app.static_folder, favicon_anterior)
                if os.path.exists(caminho_favicon_anterior):
                    os.remove(caminho_favicon_anterior)
            if novo_produto_placeholder_path and produto_placeholder_anterior and produto_placeholder_anterior != novo_produto_placeholder_path:
                caminho_placeholder_anterior = os.path.join(app.static_folder, produto_placeholder_anterior)
                if os.path.exists(caminho_placeholder_anterior):
                    os.remove(caminho_placeholder_anterior)
            if remover_produto_placeholder and produto_placeholder_anterior:
                caminho_placeholder_anterior = os.path.join(app.static_folder, produto_placeholder_anterior)
                if os.path.exists(caminho_placeholder_anterior):
                    os.remove(caminho_placeholder_anterior)
            for caminho_rel in arquivos_para_remover:
                caminho_abs = os.path.join(app.static_folder, caminho_rel)
                if os.path.exists(caminho_abs):
                    os.remove(caminho_abs)
            flash('Configurações de e-commerce salvas com sucesso.', 'success')
            return redirect(url_for('configurar_ecommerce'))
        except Exception as e:
            db.session.rollback()
            for caminho_rel in novos_arquivos:
                caminho_abs = os.path.join(app.static_folder, caminho_rel)
                if os.path.exists(caminho_abs):
                    os.remove(caminho_abs)
            flash(f'Erro ao salvar configurações de e-commerce: {str(e)}', 'error')

    return render_template(
        'sistema/ecommerce_config.html',
        empresa=empresa,
        banners_config=_montar_slots_banners(),
        pagamentos_ecommerce_texto=payment_options_to_text(empresa.pagamentos_ecommerce_json, 'ecommerce'),
        integracoes_ecommerce_texto=api_integrations_to_text(empresa.integracoes_ecommerce_json),
    )


@app.route('/ecommerce-marketing', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def configurar_marketing_ecommerce():
    empresa = EmpresaConfig.query.first()
    if not empresa:
        empresa = EmpresaConfig()
        db.session.add(empresa)
        db.session.commit()

    temas_sazonais = [
        ('neutro', 'Neutro'),
        ('natal', 'Natal'),
        ('pascoa', 'Pascoa'),
        ('ano_novo', 'Ano Novo'),
        ('dia_das_maes', 'Dia das Maes'),
        ('dia_dos_pais', 'Dia dos Pais'),
        ('black_friday', 'Black Friday'),
        ('volta_as_aulas', 'Volta as aulas'),
        ('festa_junina', 'Festa junina'),
    ]

    tipos_desconto_cupom = [
        ('percentual', 'Percentual (%)'),
        ('fixo', 'Valor fixo (R$)'),
        ('frete', 'Frete promocional'),
    ]

    def _montar_slots_campanhas():
        slots = []
        origem = _carregar_json_lista(empresa.ecom_campanhas_json)
        for i in range(5):
            item = origem[i] if i < len(origem) and isinstance(origem[i], dict) else {}
            slots.append({
                'nome': item.get('nome') or '',
                'texto': item.get('texto') or '',
                'tema_sazonal': item.get('tema_sazonal') or 'neutro',
                'inicio_em': item.get('inicio_em') or '',
                'fim_em': item.get('fim_em') or '',
                'ativo': bool(item.get('ativo', False)),
            })
        return slots

    def _montar_slots_cupons():
        slots = []
        origem = _carregar_json_lista(empresa.ecom_cupons_json)
        for i in range(5):
            item = origem[i] if i < len(origem) and isinstance(origem[i], dict) else {}
            slots.append({
                'codigo': item.get('codigo') or '',
                'descricao': item.get('descricao') or '',
                'tipo_desconto': item.get('tipo_desconto') or 'percentual',
                'valor': item.get('valor') or '',
                'inicio_em': item.get('inicio_em') or '',
                'fim_em': item.get('fim_em') or '',
                'ativo': bool(item.get('ativo', False)),
            })
        return slots

    if request.method == 'POST':
        try:
            empresa.ecom_texto_promocao = request.form.get('ecom_texto_promocao', '').strip() or None

            campanhas_slots = []
            for idx in range(5):
                nome = (request.form.get(f'campanha_nome_{idx}') or '').strip()
                texto = (request.form.get(f'campanha_texto_{idx}') or '').strip()
                tema_sazonal = (request.form.get(f'campanha_tema_{idx}') or 'neutro').strip() or 'neutro'
                inicio_em = (request.form.get(f'campanha_inicio_{idx}') or '').strip()
                fim_em = (request.form.get(f'campanha_fim_{idx}') or '').strip()
                ativo = (request.form.get(f'campanha_ativa_{idx}') == 'on')

                if not any([nome, texto, inicio_em, fim_em]):
                    continue

                inicio_data = _parse_date_iso(inicio_em)
                fim_data = _parse_date_iso(fim_em)
                if fim_data and inicio_data and fim_data < inicio_data:
                    flash(f'Campanha {idx + 1}: o fim da vigencia nao pode ser menor que o inicio.', 'error')
                    return redirect(url_for('configurar_marketing_ecommerce'))

                campanhas_slots.append({
                    'nome': nome,
                    'texto': texto,
                    'tema_sazonal': tema_sazonal,
                    'inicio_em': inicio_em,
                    'fim_em': fim_em,
                    'ativo': ativo,
                })

            cupons_slots = []
            for idx in range(5):
                codigo = (request.form.get(f'cupom_codigo_{idx}') or '').strip().upper()
                descricao = (request.form.get(f'cupom_descricao_{idx}') or '').strip()
                tipo_desconto = (request.form.get(f'cupom_tipo_{idx}') or 'percentual').strip() or 'percentual'
                valor = (request.form.get(f'cupom_valor_{idx}') or '').strip()
                inicio_em = (request.form.get(f'cupom_inicio_{idx}') or '').strip()
                fim_em = (request.form.get(f'cupom_fim_{idx}') or '').strip()
                ativo = (request.form.get(f'cupom_ativo_{idx}') == 'on')

                if not any([codigo, descricao, valor, inicio_em, fim_em]):
                    continue

                if tipo_desconto not in {item[0] for item in tipos_desconto_cupom}:
                    flash(f'Cupom {idx + 1}: tipo de desconto invalido.', 'error')
                    return redirect(url_for('configurar_marketing_ecommerce'))

                try:
                    valor_numerico = float(valor.replace(',', '.')) if valor else 0.0
                except ValueError:
                    flash(f'Cupom {idx + 1}: informe um valor numerico valido.', 'error')
                    return redirect(url_for('configurar_marketing_ecommerce'))

                if valor_numerico < 0:
                    flash(f'Cupom {idx + 1}: o valor nao pode ser negativo.', 'error')
                    return redirect(url_for('configurar_marketing_ecommerce'))

                inicio_data = _parse_date_iso(inicio_em)
                fim_data = _parse_date_iso(fim_em)
                if fim_data and inicio_data and fim_data < inicio_data:
                    flash(f'Cupom {idx + 1}: o fim da vigencia nao pode ser menor que o inicio.', 'error')
                    return redirect(url_for('configurar_marketing_ecommerce'))

                cupons_slots.append({
                    'codigo': codigo,
                    'descricao': descricao,
                    'tipo_desconto': tipo_desconto,
                    'valor': f'{valor_numerico:.2f}',
                    'inicio_em': inicio_em,
                    'fim_em': fim_em,
                    'ativo': ativo,
                })

            empresa.ecom_campanhas_json = json.dumps(campanhas_slots, ensure_ascii=False)
            empresa.ecom_cupons_json = json.dumps(cupons_slots, ensure_ascii=False)
            db.session.commit()
            flash('Promocoes, cupons e temas sazonais salvos com sucesso.', 'success')
            return redirect(url_for('configurar_marketing_ecommerce'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao salvar configuracoes de marketing do e-commerce: {str(e)}', 'error')

    return render_template(
        'sistema/ecommerce_marketing.html',
        empresa=empresa,
        campanhas_config=_montar_slots_campanhas(),
        cupons_config=_montar_slots_cupons(),
        temas_sazonais=temas_sazonais,
        tipos_desconto_cupom=tipos_desconto_cupom,
    )


@app.route('/ecommerce-ativacao', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def configurar_ativacao_ecommerce():
    empresa = EmpresaConfig.query.first()
    if not empresa:
        empresa = EmpresaConfig()
        db.session.add(empresa)
        db.session.commit()

    if request.method == 'POST':
        try:
            empresa.ecommerce_ativo = (request.form.get('ecommerce_ativo') == 'on')
            db.session.commit()
            flash('Ativacao do e-commerce salva com sucesso.', 'success')
            return redirect(url_for('configurar_ativacao_ecommerce'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao salvar ativacao do e-commerce: {str(e)}', 'error')

    return render_template('sistema/ecommerce_ativacao.html', empresa=empresa)


@app.route('/empresa/config-cardapio/preview')
@require_role('admin', 'gerente')
def preview_cardapio_empresa():
    empresa = EmpresaConfig.query.first()
    if empresa and empresa.atendimento_mesas_ativo is False:
        flash('Atendimento por mesas e garcons esta desativado na empresa.', 'warning')
        return redirect(url_for('editar_empresa'))
    qtd_max = (empresa.cardapio_qtd_maxima if empresa and empresa.cardapio_qtd_maxima else 20)
    if qtd_max <= 0:
        qtd_max = 20

    categorias = Categoria.query.order_by(Categoria.nome.asc()).all()
    categorias_cardapio = []
    for categoria in categorias:
        produtos = Produto.query.filter_by(categoria_id=categoria.id, ativo=True).order_by(Produto.nome.asc()).all()
        if not produtos:
            continue
        categorias_cardapio.append({
            'categoria': categoria,
            'produtos': produtos
        })

    class MesaPreview:
        numero = 'Preview'

    return render_template(
        'public/cardapio.html',
        mesa=MesaPreview(),
        empresa=empresa,
        qtd_max=qtd_max,
        cliente_nome='Cliente Exemplo',
        cliente_celular='(00) 00000-0000',
        cliente_slug='preview',
        categorias_cardapio=categorias_cardapio,
        pedidos_cliente=[],
        preview_mode=True
    )


def _query_chamados_home(funcionario):
    if not _tabela_existe('chamados_internos'):
        return None
    query = ChamadoInterno.query.filter(
        ChamadoInterno.status.notin_([ChamadoInterno.STATUS_CONCLUIDO, ChamadoInterno.STATUS_CANCELADO])
    )
    if not _pode_gerir_chamados(funcionario):
        query = query.filter(
            db.or_(
                ChamadoInterno.solicitante_id == funcionario.id,
                ChamadoInterno.responsavel_id == funcionario.id,
            )
        )
    return query


def _query_ordens_home(funcionario):
    if not _tabela_existe('ordens_servico'):
        return None
    query = OrdemServico.query.filter(
        OrdemServico.status.in_([
            OrdemServico.STATUS_ABERTA,
            OrdemServico.STATUS_ENVIADA,
            OrdemServico.STATUS_EM_EXECUCAO,
        ])
    )
    if funcionario.role not in {'admin', 'gerente'}:
        query = query.filter(OrdemServico.funcionario_destino_id == funcionario.id)
    return query


def _montar_prioridades_home(funcionario, paginas_permitidas, empresa):
    prioridades = []
    origens_entrega = ['site']
    if empresa and empresa.separacao_entrega_unir_vendas_off:
        origens_entrega.append('interno')

    if 'pedidos' in paginas_permitidas:
        pedidos_abertos = Pedido.query.filter(
            Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO, Pedido.STATUS_ENTREGUE])
        ).count()
        pedidos_preparo = Pedido.query.filter(
            Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO])
        ).count()
        prioridades.append({
            'titulo': 'Pedidos na fila',
            'quantidade': pedidos_abertos,
            'descricao': 'Pedidos aguardando acao comercial ou operacional.',
            'acao': 'Abrir pedidos',
            'url': url_for('listar_pedidos'),
            'tom': 'primary',
            'detalhe': f'{pedidos_preparo} ainda em registro ou preparo.',
        })

    if 'expedicao' in paginas_permitidas and empresa and empresa.separacao_entrega_ativa is not False:
        pedidos_separacao = Pedido.query.filter(
            Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO, Pedido.STATUS_ENTREGUE]),
            Pedido.origem.in_(origens_entrega),
            db.or_(
                Pedido.separacao_entrega_concluida.is_(False),
                Pedido.separacao_entrega_concluida.is_(None),
            ),
        ).count()
        pedidos_despacho = Pedido.query.filter(
            Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO, Pedido.STATUS_ENTREGUE]),
            Pedido.origem.in_(origens_entrega),
            Pedido.separacao_entrega_concluida.is_(True),
            Pedido.entrega_concluida_em.is_(None),
        ).count()
        prioridades.append({
            'titulo': 'Separacao de entrega',
            'quantidade': pedidos_separacao,
            'descricao': 'Pedidos prontos comercialmente e aguardando separacao fisica.',
            'acao': 'Abrir separacao',
            'url': url_for('listar_separacao_entrega'),
            'tom': 'warning',
            'detalhe': f'{pedidos_despacho} ja separados aguardando roteirizacao ou despacho.',
        })

    if 'recebimentos' in paginas_permitidas:
        recebimentos_conferencia = _count_recebimentos_by_status_safe(
            RecebimentoFornecedor.STATUS_CRIADO
        )
        pendencias_armazenagem = _count_recebimentos_by_status_safe(
            RecebimentoFornecedor.STATUS_AGUARDANDO_ARMAZENAGEM
        )
        prioridades.append({
            'titulo': 'Recebimentos operacionais',
            'quantidade': recebimentos_conferencia,
            'descricao': 'Entradas aguardando conferencia inicial.',
            'acao': 'Abrir recebimentos',
            'url': url_for('listar_recebimentos_fornecedor'),
            'tom': 'info',
            'detalhe': f'{pendencias_armazenagem} pendencia(s) de armazenagem.',
        })

    if 'produtos' in paginas_permitidas:
        produtos_alerta = Produto.query.filter(
            Produto.quantidade_estoque < Produto.quantidade_minima,
            Produto.ativo.is_(True),
        ).count()
        prioridades.append({
            'titulo': 'Risco de ruptura',
            'quantidade': produtos_alerta,
            'descricao': 'Produtos abaixo do minimo configurado.',
            'acao': 'Revisar produtos',
            'url': url_for('listar_produtos', ruptura='sim'),
            'tom': 'danger',
            'detalhe': 'Use filtros e picking para priorizar reposicao.',
        })

    if 'financeiro' in paginas_permitidas:
        lancamentos_pendentes = LancamentoFinanceiro.query.filter(
            LancamentoFinanceiro.incluir_contabilidade.is_(True),
            db.or_(
                LancamentoFinanceiro.enviado_contador.is_(False),
                LancamentoFinanceiro.enviado_contador.is_(None),
            ),
        ).count()
        prioridades.append({
            'titulo': 'Financeiro e contador',
            'quantidade': lancamentos_pendentes,
            'descricao': 'Lancamentos ainda nao enviados para rotina contabil.',
            'acao': 'Abrir lancamentos',
            'url': url_for('financeiro_lancamentos'),
            'tom': 'secondary',
            'detalhe': 'Revise perdas, consumo interno e despesas operacionais.',
        })

    if 'caixas' in paginas_permitidas:
        caixas_abertos = Caixa.query.filter_by(aberto=True).count()
        prioridades.append({
            'titulo': 'Caixas abertos',
            'quantidade': caixas_abertos,
            'descricao': 'Caixas com operacao em andamento neste momento.',
            'acao': 'Abrir caixas',
            'url': url_for('listar_caixas'),
            'tom': 'success',
            'detalhe': 'Conferir aberturas, responsaveis e fechamento do turno.',
        })

    if 'chamados_internos' in paginas_permitidas:
        chamados_query = _query_chamados_home(funcionario)
        chamados_abertos = chamados_query.count() if chamados_query is not None else 0
        prioridades.append({
            'titulo': 'Chamados internos',
            'quantidade': chamados_abertos,
            'descricao': 'Pendencias de sistema, infraestrutura ou suporte.',
            'acao': 'Abrir chamados',
            'url': url_for('listar_chamados_internos'),
            'tom': 'dark',
            'detalhe': 'Priorize itens criticos e em triagem.',
        })

    if 'servicos_tecnicos' in paginas_permitidas:
        ordens_query = _query_ordens_home(funcionario)
        ordens_abertas = ordens_query.count() if ordens_query is not None else 0
        prioridades.append({
            'titulo': 'Ordens de servico',
            'quantidade': ordens_abertas,
            'descricao': 'Ordens tecnicas aguardando envio, execucao ou conclusao.',
            'acao': 'Abrir ordens',
            'url': url_for('listar_ordens_servico'),
            'tom': 'primary',
            'detalhe': 'Acompanhe agenda tecnica e retornos em aberto.',
        })

    prioridades.sort(key=lambda item: item['quantidade'], reverse=True)
    return prioridades


def _montar_atalhos_home(paginas_permitidas, empresa):
    candidatos = [
        ('pdv', 'PDV', 'Abrir venda no caixa ou pedido interno.', 'Abrir PDV', 'pdv'),
        ('pedidos', 'Pedidos', 'Ver filas de registro, preparo e fechamento.', 'Abrir pedidos', 'listar_pedidos'),
        ('recebimentos', 'Recebimentos', 'Conferir entradas e armazenagem.', 'Abrir recebimentos', 'listar_recebimentos_fornecedor'),
        ('produtos', 'Produtos', 'Consultar ruptura, picking e cadastro.', 'Abrir produtos', 'listar_produtos'),
        ('financeiro', 'Lancamentos', 'Registrar perdas, consumo e despesas.', 'Abrir lancamentos', 'financeiro_lancamentos'),
        ('caixas', 'Caixas', 'Abrir, fechar e revisar historico.', 'Abrir caixas', 'listar_caixas'),
        ('funcionarios', 'Funcionarios', 'Revisar equipe, acessos e estrutura.', 'Abrir funcionarios', 'listar_funcionarios'),
        ('rh_funcoes', 'Perfis e cargos', 'Ajustar cargos e perfis de acesso.', 'Abrir RH', 'listar_perfis_rh'),
        ('expedicao', 'Expedicao', 'Separar, roteirizar e despachar entregas.', 'Abrir expedicao', 'central_expedicao'),
        ('ajuda', 'Ajuda', 'Consultar fluxos, FAQ e treinamento.', 'Abrir ajuda', 'central_ajuda'),
    ]
    atalhos = []
    for pagina, titulo, descricao, acao, endpoint in candidatos:
        if pagina not in paginas_permitidas:
            continue
        if endpoint == 'central_expedicao' and empresa and empresa.separacao_entrega_ativa is False:
            continue
        atalhos.append({
            'titulo': titulo,
            'descricao': descricao,
            'acao': acao,
            'url': url_for(endpoint),
        })
    return atalhos[:8]


def _montar_fluxos_home(paginas_permitidas, empresa):
    passos = []
    if 'pedidos' in paginas_permitidas:
        passos.append({
            'titulo': 'Validar pedidos novos',
            'descricao': 'Comece por registro e preparo para manter a fila comercial organizada.',
            'url': url_for('listar_pedidos'),
            'acao': 'Ir para pedidos',
        })
    if 'expedicao' in paginas_permitidas and empresa and empresa.separacao_entrega_ativa is not False:
        passos.append({
            'titulo': 'Separar e despachar entregas',
            'descricao': 'Conclua separacao, roteirizacao e saida de pedidos de entrega.',
            'url': url_for('listar_separacao_entrega'),
            'acao': 'Ir para expedicao',
        })
    if 'recebimentos' in paginas_permitidas:
        passos.append({
            'titulo': 'Conferir entradas e armazenagem',
            'descricao': 'Receba, confira avarias e finalize o put-away antes da reposicao.',
            'url': url_for('listar_recebimentos_fornecedor'),
            'acao': 'Ir para recebimentos',
        })
    if 'financeiro' in paginas_permitidas:
        passos.append({
            'titulo': 'Atualizar perdas e despesas',
            'descricao': 'Lance consumo interno, despesas operacionais e pendencias contabeis.',
            'url': url_for('financeiro_lancamentos'),
            'acao': 'Ir para lancamentos',
        })
    if 'funcionarios' in paginas_permitidas:
        passos.append({
            'titulo': 'Revisar equipe e acessos',
            'descricao': 'Confirme perfil, cargo, responsavel direto e filiais liberadas para cada colaborador.',
            'url': url_for('listar_funcionarios'),
            'acao': 'Ir para funcionarios',
        })
    if 'ajuda' in paginas_permitidas:
        passos.append({
            'titulo': 'Treinar e alinhar operacao',
            'descricao': 'Use a ajuda para reforcar o fluxo correto antes de executar processos novos.',
            'url': url_for('central_ajuda'),
            'acao': 'Abrir ajuda',
        })
    return passos


def _perfil_estoquista(funcionario, paginas_permitidas):
    cargo = (getattr(funcionario, 'cargo', '') or '').strip().lower()
    return 'estoqu' in cargo or 'almox' in cargo or (
        {'recebimentos', 'relatorios', 'movimentacoes'}.intersection(paginas_permitidas)
        and not {'financeiro', 'pdv'}.intersection(paginas_permitidas)
    )


def _coletar_metricas_dashboard_tempo_real():
    agora = datetime.utcnow()
    inicio_dia = agora.replace(hour=0, minute=0, second=0, microsecond=0)
    try:
        pedidos_ultimos_15_min = Pedido.query.filter(Pedido.criado_em >= (agora - timedelta(minutes=15))).count()
        faturamento_dia = db.session.query(db.func.sum(Pedido.total)).filter(
            Pedido.status == Pedido.STATUS_FECHADO,
            Pedido.fechado_em >= inicio_dia,
            Pedido.fechado_em < inicio_dia + timedelta(days=1),
        ).scalar() or 0.0
        produtos_vendidos_turno = db.session.query(db.func.sum(ItemPedido.quantidade)).join(
            Pedido, Pedido.id == ItemPedido.pedido_id
        ).filter(
            Pedido.status == Pedido.STATUS_FECHADO,
            Pedido.fechado_em >= inicio_dia,
            Pedido.fechado_em < inicio_dia + timedelta(days=1),
        ).scalar() or 0
    except (OperationalError, ProgrammingError):
        app.logger.warning('metricas_dashboard_tempo_real indisponiveis por schema/banco inconsistente', exc_info=True)
        db.session.rollback()
        pedidos_ultimos_15_min = 0
        faturamento_dia = 0.0
        produtos_vendidos_turno = 0
    return {
        'pedidos_ultimos_15_min': int(pedidos_ultimos_15_min or 0),
        'faturamento_dia': float(faturamento_dia or 0.0),
        'produtos_vendidos_turno': int(produtos_vendidos_turno or 0),
    }


def _montar_indicadores_contexto_usuario(funcionario, paginas_permitidas):
    indicadores = []
    metricas_tempo_real = _coletar_metricas_dashboard_tempo_real()
    analytics = _coletar_dashboard_analytics(datetime.utcnow() - timedelta(days=7), datetime.utcnow() + timedelta(days=1))

    if funcionario.role in {'admin', 'gerente'}:
        indicadores.extend([
            {
                'titulo': 'Pedidos em aberto',
                'valor': Pedido.query.filter(Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO])).count(),
                'detalhe': 'Fila operacional atual.',
                'metric_key': 'pedidos_ultimos_15_min',
            },
            {
                'titulo': 'Faturamento do dia',
                'valor': f"R$ {metricas_tempo_real['faturamento_dia']:.2f}",
                'detalhe': 'Fechamentos confirmados hoje.',
                'metric_key': 'faturamento_dia',
            },
            {
                'titulo': 'Tempo medio de preparo',
                'valor': f"{analytics['tempo_medio_preparo_minutos']:.1f} min",
                'detalhe': 'Media de criado ate fechado.',
                'metric_key': 'tempo_medio_preparo_minutos',
            },
        ])
    elif funcionario.role in {'operador', 'caixa'}:
        top_produto = analytics['top_produtos_vendidos'][0]['nome'] if analytics.get('top_produtos_vendidos') else '-'
        indicadores.extend([
            {
                'titulo': 'Pedidos em aberto',
                'valor': Pedido.query.filter(Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO])).count(),
                'detalhe': 'Fila para atendimento imediato.',
                'metric_key': 'pedidos_ultimos_15_min',
            },
            {
                'titulo': 'Faturamento do turno',
                'valor': f"R$ {metricas_tempo_real['faturamento_dia']:.2f}",
                'detalhe': 'Acumulado de hoje.',
                'metric_key': 'faturamento_dia',
            },
            {
                'titulo': 'Top vendido',
                'valor': top_produto,
                'detalhe': 'Produto lider no periodo.',
                'metric_key': 'produtos_vendidos_turno',
            },
        ])
    elif _perfil_estoquista(funcionario, paginas_permitidas):
        produtos_em_falta = Produto.query.filter(
            Produto.ativo.is_(True),
            Produto.quantidade_estoque < Produto.quantidade_minima,
        ).count()
        recebimentos_pendentes = RecebimentoFornecedor.query.filter(
            RecebimentoFornecedor.status == RecebimentoFornecedor.STATUS_AGUARDANDO_ARMAZENAGEM
        ).count()
        total_enderecos_ativos = EnderecoEstoque.query.filter_by(ativo=True).count()
        enderecos_ocupados = db.session.query(EnderecoEstoque.id).join(
            Produto, Produto.endereco_id == EnderecoEstoque.id
        ).filter(
            EnderecoEstoque.ativo.is_(True),
            Produto.ativo.is_(True),
        ).distinct().count()
        taxa_ocupacao = ((enderecos_ocupados / total_enderecos_ativos) * 100.0) if total_enderecos_ativos > 0 else 0.0
        indicadores.extend([
            {
                'titulo': 'Produtos em falta',
                'valor': produtos_em_falta,
                'detalhe': 'Abaixo do estoque minimo.',
                'metric_key': 'produtos_em_falta',
            },
            {
                'titulo': 'Recebimentos pendentes',
                'valor': recebimentos_pendentes,
                'detalhe': 'Aguardando armazenagem.',
                'metric_key': 'recebimentos_pendentes',
            },
            {
                'titulo': 'Ocupacao de enderecos',
                'valor': f'{taxa_ocupacao:.1f}%',
                'detalhe': 'Endereco ocupado x ativo.',
                'metric_key': 'ocupacao_enderecos',
            },
        ])
    return indicadores


def _montar_alertas_acionaveis_dashboard(paginas_permitidas):
    alertas = []
    analytics = _coletar_dashboard_analytics(datetime.utcnow() - timedelta(days=7), datetime.utcnow() + timedelta(days=1))
    produto_ruptura = Produto.query.filter(
        Produto.ativo.is_(True),
        Produto.quantidade_estoque < Produto.quantidade_minima,
    ).order_by((Produto.quantidade_estoque - Produto.quantidade_minima).asc()).first()
    if produto_ruptura and 'recebimentos' in paginas_permitidas:
        alertas.append({
            'tipo': 'warning',
            'mensagem': f'Produto {produto_ruptura.nome} em ruptura ou abaixo do minimo.',
            'acao_label': 'Reabastecer',
            'acao_url': url_for('novo_recebimento_fornecedor'),
        })

    if analytics['margem_bruta_pct'] < 25 and 'produtos' in paginas_permitidas:
        alertas.append({
            'tipo': 'danger',
            'mensagem': 'Margem bruta abaixo da meta de 25% no periodo.',
            'acao_label': 'Revisar precos',
            'acao_url': url_for('listar_produtos'),
        })

    total_mesas = Mesa.query.count()
    mesas_ocupadas = Mesa.query.filter(Mesa.status == 'ocupada').count()
    ocupacao_mesas = ((mesas_ocupadas / total_mesas) * 100.0) if total_mesas > 0 else 0.0
    if ocupacao_mesas > 90 and 'mesas' in paginas_permitidas:
        alertas.append({
            'tipo': 'info',
            'mensagem': 'Alta demanda no salao. Mais de 90% das mesas estao ocupadas.',
            'acao_label': 'Orientar garcons',
            'acao_url': url_for('listar_mesas'),
        })

    pedidos_em_preparo = Pedido.query.filter(Pedido.status == Pedido.STATUS_EM_PREPARO).count()
    if analytics['tempo_medio_preparo_minutos'] > 30 and pedidos_em_preparo > 0 and 'pedidos' in paginas_permitidas:
        alertas.append({
            'tipo': 'warning',
            'mensagem': 'Tempo medio de preparo acima de 30 minutos.',
            'acao_label': 'Abrir fila de producao',
            'acao_url': url_for('listar_pedidos', status='em_preparo'),
        })

    return alertas


@app.route('/boas-vindas')
@login_required
def boas_vindas():
    funcionario = get_funcionario_logado()
    if not funcionario:
        flash('Sessao invalida. Faca login novamente.', 'danger')
        return redirect(url_for('login'))

    empresa = EmpresaConfig.query.first()
    paginas_permitidas = _paginas_permitidas_para_funcionario(funcionario)
    saudacao_hora = datetime.now().hour
    if saudacao_hora < 12:
        saudacao = 'Bom dia'
    elif saudacao_hora < 18:
        saudacao = 'Boa tarde'
    else:
        saudacao = 'Boa noite'

    estoques_disponiveis = _estoques_contexto_disponiveis(funcionario)
    estoque_contexto_id = _estoque_contexto_selecionado_id(funcionario)
    estoque_contexto_atual = next((item for item in estoques_disponiveis if item.id == estoque_contexto_id), None)
    perfil_nome = funcionario.perfil_acesso.nome if getattr(funcionario, 'perfil_acesso', None) else 'Sem perfil padrao'
    prioridades = _montar_prioridades_home(funcionario, paginas_permitidas, empresa)
    atalhos = _montar_atalhos_home(paginas_permitidas, empresa)
    fluxos = _montar_fluxos_home(paginas_permitidas, empresa)
    indicadores_contexto_usuario = _montar_indicadores_contexto_usuario(funcionario, paginas_permitidas)
    alertas_dashboard = _montar_alertas_acionaveis_dashboard(paginas_permitidas)
    metricas_tempo_real = _coletar_metricas_dashboard_tempo_real()

    return render_template(
        'sistema/boas_vindas.html',
        app_name=APP_NAME,
        app_version=APP_VERSION,
        app_domain=APP_DOMAIN,
        funcionario=funcionario,
        saudacao=saudacao,
        agora=datetime.now(),
        prioridades=prioridades,
        atalhos=atalhos,
        fluxos=fluxos,
        perfil_nome=perfil_nome,
        estoque_contexto_atual=estoque_contexto_atual,
        estoques_disponiveis=estoques_disponiveis,
        paginas_liberadas=len(paginas_permitidas),
        indicadores_contexto_usuario=indicadores_contexto_usuario,
        alertas_dashboard=alertas_dashboard,
        metricas_tempo_real=metricas_tempo_real,
    )


AJUDA_TOPICOS = {
    'primeiros-passos': {
        'slug': 'primeiros-passos',
        'titulo': 'Primeiros passos no sistema',
        'resumo': 'Configura o ambiente inicial para operar com seguranca e consistencia.',
        'objetivo': 'Preparar a empresa para operar do cadastro inicial ao primeiro teste controlado.',
        'paginas': [],
        'checklist': [
            'Validar quem sera responsavel por acessos, empresa e operacao.',
            'Confirmar dados da empresa, canal de operacao e recursos ativos.',
            'Separar alguns produtos e usuarios de teste antes de iniciar vendas reais.',
        ],
        'passos': [
            'Acesse Meu RH > Funcionarios e valide os perfis de acesso.',
            'Cadastre os dados da empresa em Meu RH > Empresa.',
            'Defina o tipo de negocio e o canal de operacao.',
            'Crie ao menos uma categoria e alguns produtos para teste.',
            'Abra um caixa e valide o fluxo completo no PDV.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-primeiros-passos.svg',
            'alt': 'Fluxograma dos primeiros passos para implantacao do sistema.',
            'legenda': 'Fluxo recomendado para iniciar a operacao sem pular configuracoes basicas.',
        },
        'alertas': [
            'Nao use vendas reais antes de testar caixa, pedidos e estoque com dados de exemplo.',
            'Ative o controle de acesso somente depois de revisar cargo e paginas permitidas.',
            'Se houver e-commerce ou expedicao, valide esses recursos antes de divulgar a operacao.',
        ],
        'duvidas': [
            {
                'pergunta': 'Por que alguns menus nao aparecem para mim?',
                'resposta': 'O sistema mostra apenas paginas liberadas para o seu perfil. Peça revisao de acesso em RH > Funcionarios > Acessos.',
            },
            {
                'pergunta': 'Qual modulo devo configurar primeiro?',
                'resposta': 'Comece por Empresa, depois Estoque basico (categoria/produto), em seguida PDV/Pedidos e por fim Expedicao, Financeiro e RH.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Nao consigo salvar cadastro inicial.',
                'acao': 'Revise campos obrigatorios e se ha mensagens em vermelho no topo da pagina.',
            },
            {
                'situacao': 'Nao encontro um colaborador recem-criado.',
                'acao': 'Confirme se ele esta ativo e se o filtro da tela de Funcionarios nao esta limitado.',
            },
        ],
    },
    'duvidas-acesso': {
        'slug': 'duvidas-acesso',
        'titulo': 'Duvidas frequentes de acesso e navegacao',
        'resumo': 'Responde bloqueios comuns de menu, permissoes e acoes indisponiveis.',
        'objetivo': 'Ajudar o usuario a identificar se o bloqueio e de perfil, permissao, configuracao ou sessao.',
        'paginas': [],
        'checklist': [
            'Conferir se o usuario esta ativo e com senha correta.',
            'Verificar cargo, perfil e paginas liberadas em Funcionarios > Acessos.',
            'Fazer novo login depois de qualquer ajuste de permissao.',
        ],
        'passos': [
            'Valide se o colaborador esta ativo em Meu RH > Funcionarios.',
            'Confira o cargo/perfil do colaborador e a matricula utilizada no login.',
            'Se o controle de acesso estiver ativo, revise permissoes por pagina em Funcionarios > Acessos.',
            'Confirme se a funcionalidade esta habilitada em configuracoes da empresa.',
            'Peça novo login apos alteracoes de permissao para atualizar a sessao.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-acesso.svg',
            'alt': 'Fluxograma para diagnostico de bloqueios de acesso e navegacao.',
            'legenda': 'Use esta sequencia para identificar rapidamente o motivo de um menu ou acao nao aparecer.',
        },
        'alertas': [
            'Trocar apenas o cargo do colaborador pode nao ser suficiente se o controle de acesso por pagina estiver ativo.',
            'Permissoes novas normalmente exigem novo login para atualizar a sessao.',
            'Modulos desativados na empresa continuam ocultos mesmo para usuarios com cargo alto.',
        ],
        'duvidas': [
            {
                'pergunta': 'Por que o botao aparece para um usuario e para outro nao?',
                'resposta': 'As acoes obedecem permissao de pagina e perfil (admin/gerente/operacional).',
            },
            {
                'pergunta': 'Ajuda deve ficar visivel para todos os perfis?',
                'resposta': 'Sim. A ajuda fica disponivel para todos, mas o conteudo listado respeita o que cada perfil pode acessar.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Recebo mensagem de acesso negado.',
                'acao': 'Solicite ao gestor revisao do perfil de acesso e dos ajustes individuais da pagina.',
            },
            {
                'situacao': 'Menu sumiu apos troca de cargo.',
                'acao': 'Revalide o perfil de acesso padrao e os ajustes individuais do colaborador, depois entre novamente no sistema.',
            },
        ],
    },
    'estoque-operacao': {
        'slug': 'estoque-operacao',
        'titulo': 'Operacao de estoque',
        'resumo': 'Passo a passo para cadastro, entrada, saida e acompanhamento de estoque.',
        'paginas': ['estoques', 'produtos', 'categorias', 'fornecedores', 'movimentacoes', 'relatorios'],
        'objetivo': 'Manter produtos, saldos e reposicao sob controle para evitar ruptura e divergencia fisica.',
        'checklist': [
            'Ter fornecedores, categorias e estoques cadastrados.',
            'Definir estoque minimo e preco nos produtos principais.',
            'Escolher o processo certo: recebimento, movimentacao, transferencia ou ajuste rapido.',
        ],
        'passos': [
            'Cadastre fornecedores com os dados principais.',
            'Crie categorias e depois cadastre os produtos.',
            'Informe preco de custo, preco de venda e estoque minimo.',
            'Registre entradas, saidas e recebimentos por fornecedor.',
            'Acompanhe relatorios para ruptura, giro e valor em estoque.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-estoque.svg',
            'alt': 'Fluxograma da operacao de estoque desde cadastro ate acompanhamento.',
            'legenda': 'Ciclo basico do estoque: cadastrar, movimentar, armazenar e acompanhar indicadores.',
        },
        'alertas': [
            'Nunca corrija saldo no produto sem registrar a movimentacao correspondente.',
            'Recebimentos e transferencias devem ser fechados no mesmo fluxo para nao gerar saldo solto.',
            'Produtos sem endereco ou estoque minimo ficam mais sujeitos a ruptura e retrabalho.',
        ],
        'duvidas': [
            {
                'pergunta': 'Quando usar movimentacao rapida?',
                'resposta': 'Use para ajustes operacionais de entrada/saida em lote e abastecimento rapido da operacao.',
            },
            {
                'pergunta': 'Como evitar ruptura em itens de alto giro?',
                'resposta': 'Defina estoque minimo coerente, acompanhe alertas de reposicao e use enderecos inteligentes.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Produto sem endereco no coletor.',
                'acao': 'Acesse Enderecos Inteligentes e vincule o produto para reduzir retrabalho de separacao.',
            },
            {
                'situacao': 'Estoque ficou negativo ou incoerente.',
                'acao': 'Revise movimentacoes recentes, recebimentos e possiveis saidas duplicadas.',
            },
        ],
    },
    'vendas-pdv': {
        'slug': 'vendas-pdv',
        'titulo': 'Vendas e PDV',
        'resumo': 'Fluxo recomendado para abertura de caixa, venda e fechamento do turno.',
        'paginas': ['pdv', 'pedidos', 'caixas', 'mesas', 'garcons'],
        'objetivo': 'Padronizar a venda do caixa ao acompanhamento do pedido e do fechamento do turno.',
        'checklist': [
            'Confirmar que existe caixa aberto para a operacao.',
            'Verificar itens, precos e meios de pagamento liberados.',
            'Definir se a venda sera balcão, pedido, mesa ou comanda conforme o processo da empresa.',
        ],
        'passos': [
            'Abra o caixa com o valor inicial do turno.',
            'Lance os pedidos no PDV e confirme os itens.',
            'Selecione o metodo de pagamento e finalize a venda.',
            'Acompanhe os pedidos na tela de pedidos e status.',
            'Feche o caixa ao final do expediente e confira o historico.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-vendas.svg',
            'alt': 'Fluxograma do processo de vendas e PDV.',
            'legenda': 'Fluxo recomendado para venda no PDV com acompanhamento do pedido e fechamento do caixa.',
        },
        'alertas': [
            'Nao feche o caixa antes de conferir pedidos em aberto, pagamentos pendentes e diferencas.',
            'Pedidos finalizados devem ser tratados como historico operacional, nao como rascunho editavel.',
            'Se a venda estiver ligada a mesas ou garcons, valide a configuracao da empresa antes do turno.',
        ],
        'duvidas': [
            {
                'pergunta': 'Qual a diferenca entre Pedido e PDV?',
                'resposta': 'PDV e o ponto de registro/finalizacao da venda; Pedidos e a fila consolidada para acompanhamento operacional.',
            },
            {
                'pergunta': 'Posso editar pedido fechado?',
                'resposta': 'Normalmente nao. Pedidos fechados sao tratados como imutaveis para seguranca operacional e fiscal.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Nao consigo fechar pedido.',
                'acao': 'Confirme metodo de pagamento, valor recebido e se o caixa informado esta aberto.',
            },
            {
                'situacao': 'Pedido nao aparece no monitor.',
                'acao': 'Verifique filtros de status/busca e atualize a tela para recarregar a listagem.',
            },
        ],
    },
    'expedicao-entregas': {
        'slug': 'expedicao-entregas',
        'titulo': 'Expedicao, transferencias e abastecimento',
        'resumo': 'Organiza separacao, rota, etiquetas, notas e fluxo de abastecimento entre estoques/lojas.',
        'paginas': ['expedicao', 'pedidos', 'movimentacoes', 'enderecos_estoque'],
        'objetivo': 'Garantir que pedidos saiam separados, roteirizados e despachados com rastreio operacional.',
        'checklist': [
            'Confirmar se a separacao de entrega esta ativa na empresa.',
            'Validar frota, etiquetas e regras de roteirizacao antes do despacho.',
            'Checar se o pedido esta separado e elegivel para entrar na fila de entrega.',
        ],
        'passos': [
            'Acesse Expedicao > Central de Expedicao para iniciar o fluxo diario.',
            'Faça separacao e embalagem dos pedidos na fila operacional.',
            'Preencha rota, ordem, local de saida e dados de motorista/veiculo.',
            'Emita etiqueta de entrega e registre nota fiscal quando aplicavel.',
            'Use transferencias e entrada em massa para abastecimento entre estoques/lojas.',
            'Acompanhe Painel em Tempo Real para status da operacao diaria.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-expedicao.svg',
            'alt': 'Fluxograma da operacao de expedicao e entrega.',
            'legenda': 'Ordem recomendada: separar, roteirizar, despachar, entregar e abastecer quando necessario.',
        },
        'alertas': [
            'Pedido sem separacao concluida nao deve seguir para roteirizacao.',
            'Nao misture ajuste de estoque com despacho sem registrar a movimentacao correspondente.',
            'Etiquetas e nota fiscal dependem de configuracoes da empresa e da elegibilidade do pedido.',
        ],
        'duvidas': [
            {
                'pergunta': 'Como cadastrar frota propria e terceiros?',
                'resposta': 'Use Expedicao > Frota Propria e Terceiros para salvar veiculos, Correios, freelancers e transportadoras.',
            },
            {
                'pergunta': 'Quando usar transferencia entre estoques?',
                'resposta': 'Quando houver necessidade de reposicao de loja/CD ou ajuste fisico de enderecamento.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Pedido nao libera para roteirizacao.',
                'acao': 'Confirme se a separacao foi concluida e se o pedido nao esta cancelado/fechado.',
            },
            {
                'situacao': 'Etiqueta nao imprime.',
                'acao': 'Valide se emissao de etiqueta esta ativa nas configuracoes da empresa e se o pedido esta elegivel.',
            },
        ],
    },
    'financeiro-lancamentos': {
        'slug': 'financeiro-lancamentos',
        'titulo': 'Financeiro, lancamentos e fundos',
        'resumo': 'Padrao de registro monetario para conciliacao e exportacao contabil.',
        'paginas': ['financeiro'],
        'objetivo': 'Organizar entradas, saidas e fundos para facilitar conciliacao interna e envio contabil.',
        'checklist': [
            'Definir periodo de analise antes de lancar ou exportar dados.',
            'Conferir natureza, competencia e centro responsavel do registro.',
            'Separar o que e operacao diaria, fundo interno e item contabil.',
        ],
        'passos': [
            'Use Visao Financeira para acompanhar indicadores principais.',
            'Cadastre lancamentos com natureza, competencia e centro responsavel.',
            'Marque itens para contabilidade quando houver envio para contador.',
            'Registre solicitacoes e liberacoes de fundos conforme politica interna.',
            'Exporte arquivos para conferencia antes do fechamento mensal.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-financeiro.svg',
            'alt': 'Fluxograma do processo financeiro com lancamentos e conciliacao.',
            'legenda': 'Fluxo basico do financeiro: registrar, classificar, conciliar, aprovar fundos e exportar.',
        },
        'alertas': [
            'Lancamento sem competencia ou classificacao correta gera distorcao de relatorio.',
            'Antes de exportar, revise filtros de data e registros marcados para contabilidade.',
            'Consumo proprio, taxas e ajustes precisam ficar separados para manter rastreabilidade.',
        ],
        'duvidas': [
            {
                'pergunta': 'O que entra como lancamento contabil?',
                'resposta': 'Entradas e saidas monetarias que precisam compor conciliacao, declaracao e envio ao contador.',
            },
            {
                'pergunta': 'Posso separar consumo proprio da empresa?',
                'resposta': 'Sim. Classifique lancamentos de consumo proprio para manter rastreio e separacao fiscal.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Valores nao batem com o periodo.',
                'acao': 'Revise filtros de data, status do caixa e se houve pedidos cancelados no intervalo.',
            },
            {
                'situacao': 'Exportacao incompleta.',
                'acao': 'Confirme se lancamentos obrigatorios foram marcados para contabilidade e salvos.',
            },
        ],
    },
    'rh-seguranca': {
        'slug': 'rh-seguranca',
        'titulo': 'RH e seguranca de acesso',
        'resumo': 'Define cargos, permissoes e auditoria para controle operacional.',
        'paginas': ['funcionarios', 'rh_funcoes', 'rh_indicadores', 'rh_organograma', 'auditoria'],
        'objetivo': 'Controlar quem acessa cada modulo, como a estrutura da equipe e montada e quem alterou dados sensiveis.',
        'checklist': [
            'Criar ou revisar cargos e perfis antes de cadastrar muitos usuarios.',
            'Definir superior, departamento, time e nivel organizacional quando aplicavel.',
            'Validar se o controle de acesso por pagina ja pode ser ativado.',
        ],
        'passos': [
            'Crie funcoes/perfis com permissoes por pagina.',
            'Associe cada colaborador ao cargo correto.',
            'Monte a arvore de hierarquia no RH > Organograma.',
            'Ative o controle de acesso quando terminar a configuracao.',
            'Revise os acessos por funcionario periodicamente.',
            'Use a auditoria para rastrear alteracoes e acoes sensiveis.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-rh.svg',
            'alt': 'Fluxograma do processo de RH e seguranca de acesso.',
            'legenda': 'Fluxo recomendado para criar cargos, vincular colaboradores, ativar acessos e auditar alteracoes.',
        },
        'alertas': [
            'Ativar controle de acesso cedo demais pode esconder telas essenciais da equipe.',
            'Usuario ativo sem superior, nivel ou cargo correto pode distorcer organograma e permissoes.',
            'A auditoria deve ser revisada em mudancas sensiveis de cadastro, financeiro e operacao.',
        ],
        'duvidas': [
            {
                'pergunta': 'Como liberar upload de foto de perfil para colaborador?',
                'resposta': 'No cadastro do funcionario, habilite a opcao de permitir edicao de imagem de perfil.',
            },
            {
                'pergunta': 'Matricula pode ser usada como login?',
                'resposta': 'Sim. A matricula pode ser usada como identificador de acesso junto com senha.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Colaborador nao aparece no organograma.',
                'acao': 'Verifique superior, nivel, departamento/time e se o cadastro esta ativo.',
            },
            {
                'situacao': 'Permissoes nao foram aplicadas.',
                'acao': 'Confira cargo/perfil e, apos salvar, faça novo login para atualizar sessao.',
            },
        ],
    },
    'ecommerce-config': {
        'slug': 'ecommerce-config',
        'titulo': 'Configuracao do e-commerce',
        'resumo': 'Separa ativacao da loja online da configuracao visual, comercial e tecnica da vitrine.',
        'paginas': ['ecommerce_config', 'ecommerce_marketing'],
        'objetivo': 'Manter a loja online ativa, ajustar a vitrine e organizar campanhas, cupons e sazonalidades sem misturar configuracao visual com calendario comercial.',
        'checklist': [
            'Validar se a ativacao da loja online esta ligada antes de divulgar o link para clientes.',
            'Separar banners, logos e textos oficiais antes da configuracao.',
            'Definir periodo de vigencia para campanhas, cupons e temas sazonais.',
            'Verificar cores, rodape e imagem padrao de produto em mobile e desktop.',
        ],
        'passos': [
            'Acesse E-commerce > Ativacao da Loja e confirme se o canal publico esta liberado.',
            'Acesse E-commerce > Tema e Loja Online.',
            'Defina cores da vitrine, banners e checkout.',
            'Acesse E-commerce > Promocoes e Campanhas.',
            'Cadastre a mensagem promocional principal, campanhas e cupons com vigencia.',
            'Associe cada campanha a um tema sazonal quando houver data especial.',
            'Ajuste rodape, favicon e imagem padrao de produto.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-ecommerce.svg',
            'alt': 'Fluxograma da configuracao do e-commerce.',
            'legenda': 'Sequencia sugerida para ajustar a vitrine, separar marketing sazonal e revisar a experiencia final da loja.',
        },
        'alertas': [
            'Imagens fora do tamanho recomendado podem prejudicar a leitura da vitrine no celular.',
            'Campanhas sem vigencia clara podem continuar aparecendo fora do periodo esperado.',
            'Sempre revise a loja publica depois de alterar tema, banner, promocao ou cupom.',
        ],
        'duvidas': [
            {
                'pergunta': 'Como deixar a loja com visual padrao de marketplace?',
                'resposta': 'Ajuste paleta, banners e cards de produto no modulo de tema para manter consistencia mobile.',
            },
            {
                'pergunta': 'Posso agendar campanhas automaticamente?',
                'resposta': 'Sim. Configure inicio/fim de vigencia para promocoes, cupons e campanhas programadas na tela de marketing.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Banner nao aparece na loja.',
                'acao': 'Verifique vigencia, status ativo e dimensoes da imagem.',
            },
            {
                'situacao': 'Produto sem foto na vitrine.',
                'acao': 'Confirme imagem do produto ou fallback padrao configurado no e-commerce.',
            },
        ],
    },
    'servicos-tecnicos': {
        'slug': 'servicos-tecnicos',
        'titulo': 'Chamados internos e servicos tecnicos',
        'resumo': 'Explica como abrir, encaminhar e executar chamados e ordens de servico.',
        'paginas': ['servicos_tecnicos', 'chamados_internos'],
        'objetivo': 'Padronizar o registro da demanda interna ate a execucao tecnica e o retorno final.',
        'checklist': [
            'Definir se o caso deve nascer como chamado interno ou ordem tecnica.',
            'Informar produto, local, prioridade e descricao objetiva do problema.',
            'Validar se o servico de montagem/instalacao esta habilitado quando aplicavel.',
        ],
        'passos': [
            'Abra o chamado interno com contexto claro do problema ou necessidade.',
            'Classifique prioridade, local, responsavel e observacoes relevantes.',
            'Converta ou vincule a demanda a uma ordem de servico tecnica quando necessario.',
            'Acompanhe status, execucao em campo e retorno do tecnico.',
            'Feche a demanda somente apos validar a solucao ou o motivo do encerramento.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-servicos.svg',
            'alt': 'Fluxograma de chamados internos e ordens de servico.',
            'legenda': 'Fluxo recomendado para abertura, triagem, execucao e fechamento de servicos tecnicos.',
        },
        'alertas': [
            'Chamado sem descricao minima dificulta triagem e aumenta retrabalho.',
            'Nao feche a ordem antes de registrar o retorno tecnico e o resultado da visita.',
            'Servicos de montagem e instalacao dependem de empresa e produto habilitados.',
        ],
        'duvidas': [
            {
                'pergunta': 'Quando abrir chamado e quando abrir ordem de servico?',
                'resposta': 'O chamado registra a necessidade interna; a ordem organiza a execucao tecnica da atividade.',
            },
            {
                'pergunta': 'Posso acompanhar apenas minhas ordens?',
                'resposta': 'Sim. O modulo possui visao de ordens proprias e visao geral conforme o acesso liberado.',
            },
        ],
        'problemas': [
            {
                'situacao': 'Nao consigo criar ordem de instalacao.',
                'acao': 'Verifique se a empresa e o produto estao com servico de montagem/instalacao liberados.',
            },
            {
                'situacao': 'Chamado parado sem responsavel.',
                'acao': 'Revise prioridade, encaminhamento e se existe tecnico/perfil apto para receber a demanda.',
            },
        ],
    },
    'app-mobile': {
        'slug': 'app-mobile',
        'titulo': 'App mobile e tela inicial',
        'resumo': 'Mostra como instalar o sistema na tela inicial do celular e usar o acesso rapido.',
        'paginas': [],
        'objetivo': 'Facilitar o acesso diario no celular com comportamento de app, especialmente para operacao em campo.',
        'checklist': [
            'Abrir o sistema no celular usando navegador atualizado.',
            'Estar conectado ao sistema com a tela carregada normalmente.',
            'Verificar se o atalho ainda nao foi instalado no aparelho.',
        ],
        'passos': [
            'Ao entrar no sistema pelo celular, observe o aviso para adicionar o app a tela inicial.',
            'No Android, toque em Adicionar agora quando o navegador oferecer a instalacao.',
            'No iPhone, use Compartilhar > Adicionar a Tela de Inicio quando o aviso orientar esse passo.',
            'Abra o atalho criado para usar o sistema em tela cheia e com acesso mais rapido.',
            'Se o aviso sumir, abra novamente o sistema no navegador principal do aparelho e repita o processo.',
        ],
        'fluxograma': {
            'imagem': 'img/ajuda/fluxo-app-mobile.svg',
            'alt': 'Fluxograma para instalacao do sistema como app na tela inicial do celular.',
            'legenda': 'Fluxo simples para instalar o atalho do sistema no Android ou no iPhone.',
        },
        'alertas': [
            'O aviso de instalacao nao aparece para quem ja instalou ou dispensou recentemente.',
            'No iPhone a instalacao e manual pelo menu Compartilhar do Safari.',
            'Alguns navegadores de terceiros limitam o comportamento de app; prefira Chrome, Edge ou Safari.',
        ],
        'duvidas': [
            {
                'pergunta': 'Instalar na tela inicial muda meu acesso?',
                'resposta': 'Nao. O atalho apenas facilita abrir o sistema como app; seu perfil e permissao continuam os mesmos.',
            },
            {
                'pergunta': 'Posso remover e instalar de novo depois?',
                'resposta': 'Sim. O atalho pode ser removido do aparelho e reinstalado quando necessario.',
            },
        ],
        'problemas': [
            {
                'situacao': 'O botao de instalar nao apareceu no Android.',
                'acao': 'Atualize o navegador, recarregue a pagina e abra o sistema no navegador principal do aparelho.',
            },
            {
                'situacao': 'No iPhone nao apareceu botao de instalar.',
                'acao': 'Use o Safari e siga o caminho Compartilhar > Adicionar a Tela de Inicio.',
            },
        ],
    },
}

AJUDA_ETAPAS = {
    'primeiros-passos': [
        {
            'titulo': 'Preparar responsaveis e acessos',
            'descricao': 'Organize quem vai configurar o sistema e garanta que os perfis iniciais consigam navegar.',
            'passos': [
                'Defina quem sera o responsavel por empresa, acessos e operacao do dia a dia.',
                'Revise cadastros de funcionarios, cargo e paginas liberadas antes de iniciar a implantacao.',
                'Confirme quais modulos vao entrar em uso imediato para evitar liberar menu sem necessidade.',
            ],
        },
        {
            'titulo': 'Montar a base minima de operacao',
            'descricao': 'Cadastre o essencial para o sistema conseguir simular a rotina real.',
            'passos': [
                'Preencha os dados da empresa e escolha o tipo de negocio e o canal principal de operacao.',
                'Cadastre categorias e produtos de teste com preco, estoque e situacao corretos.',
                'Revise meios de pagamento, estoques e recursos extras como expedicao ou e-commerce, se estiverem no escopo.',
            ],
        },
        {
            'titulo': 'Validar a operacao ponta a ponta',
            'descricao': 'Execute um teste controlado antes de abrir a operacao para a equipe.',
            'passos': [
                'Abra um caixa de teste para simular o inicio do turno.',
                'Registre uma venda no PDV ou um pedido interno com itens reais de exemplo.',
                'Confira se pedido, caixa e estoque refletiram corretamente antes de liberar vendas reais.',
            ],
        },
    ],
    'duvidas-acesso': [
        {
            'titulo': 'Confirmar cadastro e sessao',
            'descricao': 'Elimine primeiro os problemas basicos de usuario, senha ou sessao expirada.',
            'passos': [
                'Confirme se o colaborador esta ativo e com a matricula correta.',
                'Valide se a senha usada corresponde ao cadastro atual ou se precisa redefinicao.',
                'Peca um novo login sempre que houver troca de permissao ou suspeita de sessao desatualizada.',
            ],
        },
        {
            'titulo': 'Revisar perfil e paginas liberadas',
            'descricao': 'Cheque se o bloqueio vem do cargo, do perfil ou de liberacao individual de pagina.',
            'passos': [
                'Abra Funcionarios e confira cargo, perfil e configuracao de acesso do colaborador.',
                'Se o controle de acesso por pagina estiver ativo, revise pagina por pagina que o usuario precisa enxergar.',
                'Confirme tambem se o modulo esta ativo nas configuracoes da empresa.',
            ],
        },
        {
            'titulo': 'Validar o comportamento na tela',
            'descricao': 'Depois do ajuste, teste exatamente a tela ou a acao que estava bloqueada.',
            'passos': [
                'Compare o menu do usuario com um perfil que ja possua o acesso esperado.',
                'Peca para o colaborador abrir novamente a tela ou executar a acao que estava indisponivel.',
                'Se o bloqueio continuar, anote a mensagem exibida para identificar se o caso e permissao, configuracao ou rota.',
            ],
        },
    ],
    'estoque-operacao': [
        {
            'titulo': 'Preparar cadastros base',
            'descricao': 'Comece pelos registros que sustentam entradas, saidas e reposicao.',
            'passos': [
                'Cadastre fornecedores, estoques e categorias antes de criar um volume grande de produtos.',
                'Defina um padrao de descricao, unidade e organizacao para o catalogo.',
                'Revise se os estoques necessarios para loja, retaguarda ou CD ja estao criados.',
            ],
        },
        {
            'titulo': 'Padronizar cadastro e saldo dos produtos',
            'descricao': 'Deixe cada item pronto para compra, venda e controle interno.',
            'passos': [
                'Cadastre produtos com preco de custo, preco de venda e estoque minimo.',
                'Associe o produto ao estoque correto e, quando aplicavel, ao endereco de armazenagem.',
                'Confirme se os itens de maior giro ficaram com parametros suficientes para reposicao e picking.',
            ],
        },
        {
            'titulo': 'Operar entradas, saidas e acompanhamento',
            'descricao': 'Registre as movimentacoes no fluxo certo e acompanhe os indicadores.',
            'passos': [
                'Use recebimentos por fornecedor para entrada formal e movimentacoes para ajustes internos.',
                'Evite alterar saldo manualmente sem uma movimentacao correspondente.',
                'Acompanhe relatorios, alertas de ruptura e necessidade de reposicao para corrigir desvios cedo.',
            ],
        },
    ],
    'vendas-pdv': [
        {
            'titulo': 'Abrir o turno de venda',
            'descricao': 'Garanta que a operacao comece com caixa, regras e parametros corretos.',
            'passos': [
                'Abra o caixa com o valor inicial previsto para o turno.',
                'Confira se os meios de pagamento e itens de venda estao liberados.',
                'Se houver mesa, comanda ou garcom, valide a configuracao da empresa antes da primeira venda.',
            ],
        },
        {
            'titulo': 'Registrar e finalizar a venda',
            'descricao': 'Conduza a venda do lancamento dos itens ate a confirmacao do pagamento.',
            'passos': [
                'Lance os itens no PDV e revise quantidades, observacoes e totais.',
                'Escolha o metodo de pagamento e confirme o valor recebido ou a forma de fechamento.',
                'Finalize a venda apenas depois de revisar o resumo do pedido.',
            ],
        },
        {
            'titulo': 'Acompanhar pedidos e encerrar o turno',
            'descricao': 'Monitore o que ficou em aberto e feche o caixa com seguranca.',
            'passos': [
                'Use a tela de Pedidos para acompanhar status, preparo e entregas vinculadas.',
                'Resolva pagamentos pendentes, pedidos travados ou divergencias antes do fechamento.',
                'Feche o caixa ao final do expediente e confira historico, sangrias e saldo final.',
            ],
        },
    ],
    'expedicao-entregas': [
        {
            'titulo': 'Preparar a fila de entrega',
            'descricao': 'Somente pedidos prontos e elegiveis devem entrar no fluxo de saida.',
            'passos': [
                'Confirme se a separacao de entrega esta ativa na empresa e se a fila do dia esta organizada.',
                'Valide se os pedidos ja foram separados e estao aptos para roteirizacao.',
                'Cheque frota, etiquetas e regras de despacho antes de iniciar a montagem das rotas.',
            ],
        },
        {
            'titulo': 'Roteirizar e despachar',
            'descricao': 'Monte a saida com informacoes suficientes para rastrear a operacao.',
            'passos': [
                'Finalize separacao e embalagem conforme a fila operacional.',
                'Preencha rota, ordem, local de saida, motorista e veiculo responsavel.',
                'Emita etiqueta e nota fiscal quando o pedido estiver elegivel para despacho.',
            ],
        },
        {
            'titulo': 'Acompanhar entrega e abastecimento',
            'descricao': 'Feche o ciclo controlando status de entrega e reposicao entre estoques.',
            'passos': [
                'Use o Painel de Expedicao para acompanhar saida, andamento e retorno da operacao.',
                'Registre transferencias entre lojas ou CDs quando houver abastecimento fisico.',
                'Revise pedidos que nao liberaram, etiquetas falhadas ou pendencias de conferencia ao fim do ciclo.',
            ],
        },
    ],
    'financeiro-lancamentos': [
        {
            'titulo': 'Definir o contexto do registro',
            'descricao': 'Comece sabendo em qual periodo, natureza e centro o lancamento deve entrar.',
            'passos': [
                'Escolha o periodo de analise antes de criar ou revisar registros.',
                'Confirme natureza, competencia e centro responsavel do movimento.',
                'Separe o que e despesa operacional, consumo proprio, fundo interno ou item contabil.',
            ],
        },
        {
            'titulo': 'Registrar e classificar os movimentos',
            'descricao': 'Padronize os cadastros para a conciliacao funcionar sem retrabalho.',
            'passos': [
                'Cadastre lancamentos com os dados monetarios e classificacoes obrigatorias.',
                'Marque para contabilidade apenas o que precisa entrar no fluxo contabil.',
                'Registre fundos, liberacoes e solicitacoes conforme a politica da empresa.',
            ],
        },
        {
            'titulo': 'Conferir e exportar',
            'descricao': 'Feche o periodo revisando filtros, pendencias e consistencia dos valores.',
            'passos': [
                'Revise o periodo, os filtros e a consistencia dos totais antes de exportar.',
                'Compare os registros com caixas, pedidos e demais origens operacionais quando houver divergencia.',
                'Exporte os arquivos para conferencia final e envio ao contador.',
            ],
        },
    ],
    'rh-seguranca': [
        {
            'titulo': 'Desenhar cargos e regras de acesso',
            'descricao': 'Defina a estrutura antes de liberar muitos usuarios no sistema.',
            'passos': [
                'Crie cargos e perfis com as paginas que cada grupo realmente precisa acessar.',
                'Defina uma regra clara para acessos administrativos, gerenciais e operacionais.',
                'Decida quando o controle de acesso por pagina sera ativado.',
            ],
        },
        {
            'titulo': 'Vincular a equipe a estrutura correta',
            'descricao': 'Associe cada colaborador ao cargo, superior e nivel adequados.',
            'passos': [
                'Cadastre ou revise colaboradores com cargo, matricula e status ativo.',
                'Informe superior, departamento, time e nivel organizacional quando fizer parte da rotina.',
                'Confirme se cada colaborador ficou ligado ao perfil de acesso esperado.',
            ],
        },
        {
            'titulo': 'Ativar controles e auditar',
            'descricao': 'Depois de configurar, mantenha a revisao periodica dos acessos.',
            'passos': [
                'Ative o controle de acesso somente depois da revisao inicial das permissoes.',
                'Peca novo login para atualizar a sessao apos mudancas relevantes.',
                'Use a auditoria para acompanhar alteracoes sensiveis em cadastros, financeiro e operacao.',
            ],
        },
    ],
    'ecommerce-config': [
        {
            'titulo': 'Ativar o canal e separar materiais',
            'descricao': 'Antes de mexer na vitrine, confirme se a loja pode ser publicada.',
            'passos': [
                'Verifique na ativacao da loja se o canal publico esta liberado.',
                'Separe logos, banners, textos e imagens oficiais antes da configuracao.',
                'Defina a vigencia esperada para campanhas, cupons e datas sazonais que vao entrar no ar.',
            ],
        },
        {
            'titulo': 'Separar vitrine e marketing',
            'descricao': 'Ajuste a identidade visual em uma tela e concentre o calendario promocional em outra.',
            'passos': [
                'Configure tema, paleta, banners e rodape da loja online.',
                'Cadastre a mensagem promocional principal, campanhas e cupons na tela de marketing.',
                'Use temas sazonais para Natal, Pascoa, Black Friday e outras datas relevantes.',
                'Revise favicon, imagem padrao de produto e demais elementos visuais de apoio.',
            ],
        },
        {
            'titulo': 'Validar a loja publica',
            'descricao': 'Sempre revise a experiencia final antes de divulgar o link aos clientes.',
            'passos': [
                'Abra a loja em mobile e desktop para validar leitura, contraste e encaixe dos banners.',
                'Confirme se os produtos aparecem com foto, preco e fallback correto.',
                'Divulgue o link somente apos a revisao final da vitrine publicada.',
            ],
        },
    ],
    'servicos-tecnicos': [
        {
            'titulo': 'Abrir a demanda correta',
            'descricao': 'Escolha o tipo de registro certo para a necessidade interna.',
            'passos': [
                'Decida se o caso deve nascer como chamado interno ou ordem tecnica.',
                'Registre descricao objetiva, local, prioridade e contexto do problema.',
                'Quando houver servico de montagem ou instalacao, confirme se empresa e produto estao habilitados.',
            ],
        },
        {
            'titulo': 'Encaminhar e executar',
            'descricao': 'Direcione a demanda para quem vai tratar o caso e acompanhe a execucao.',
            'passos': [
                'Classifique responsavel, observacoes e urgencia da solicitacao.',
                'Converta ou vincule o chamado a uma ordem tecnica quando a atividade exigir execucao operacional.',
                'Acompanhe status, execucao em campo e retorno do tecnico durante o atendimento.',
            ],
        },
        {
            'titulo': 'Fechar com rastreabilidade',
            'descricao': 'O encerramento deve deixar claro o que foi feito e qual foi o resultado.',
            'passos': [
                'Registre retorno tecnico, resultado da visita e observacoes finais.',
                'Valide com a area solicitante se a demanda foi resolvida ou precisa nova tratativa.',
                'Encerre o chamado ou a ordem somente depois de documentar o motivo do fechamento.',
            ],
        },
    ],
    'app-mobile': [
        {
            'titulo': 'Confirmar o ambiente do aparelho',
            'descricao': 'Prepare o celular para instalar o atalho sem falhas de compatibilidade.',
            'passos': [
                'Abra o sistema em um navegador atualizado no celular.',
                'Confirme que a pagina carregou normalmente e que o login foi concluido.',
                'Verifique se o atalho ainda nao foi instalado anteriormente no aparelho.',
            ],
        },
        {
            'titulo': 'Instalar o atalho na tela inicial',
            'descricao': 'Siga o caminho correto conforme o sistema operacional do aparelho.',
            'passos': [
                'No Android, aceite a sugestao de adicionar o app quando o navegador exibir o aviso.',
                'No iPhone, use Compartilhar > Adicionar a Tela de Inicio para criar o atalho.',
                'Aguarde a criacao do icone e confirme se o nome do atalho ficou identificavel para a equipe.',
            ],
        },
        {
            'titulo': 'Validar o uso diario',
            'descricao': 'Depois da instalacao, teste o comportamento do atalho como app.',
            'passos': [
                'Abra o atalho criado e confira se o sistema abriu em tela cheia.',
                'Teste o acesso rapido nas telas mais usadas pela operacao.',
                'Se o aviso de instalacao nao voltar, remova o atalho e repita o processo pelo navegador principal do aparelho.',
            ],
        },
    ],
}


def _textos_unicos_ajuda(itens):
    vistos = set()
    resultado = []
    for item in itens or []:
        texto = (item or '').strip()
        chave = texto.lower()
        if not texto or chave in vistos:
            continue
        vistos.add(chave)
        resultado.append(texto)
    return resultado


def _etapas_topico_ajuda(topico):
    slug = topico.get('slug')
    etapas_brutas = list(topico.get('etapas') or AJUDA_ETAPAS.get(slug) or [])
    etapas = []
    for indice, item in enumerate(etapas_brutas, start=1):
        if not isinstance(item, dict):
            continue
        titulo = (item.get('titulo') or f'Etapa {indice}').strip() or f'Etapa {indice}'
        descricao = (item.get('descricao') or '').strip()
        passos = _textos_unicos_ajuda(item.get('passos') or [])
        if not descricao and not passos:
            continue
        etapas.append({
            'ordem': indice,
            'titulo': titulo,
            'descricao': descricao,
            'passos': passos,
        })

    if etapas:
        return etapas

    checklist = _textos_unicos_ajuda(topico.get('checklist') or [])
    passos = _textos_unicos_ajuda(topico.get('passos') or [])
    etapas = []
    if checklist:
        etapas.append({
            'ordem': 1,
            'titulo': 'Preparacao',
            'descricao': 'Confirme os pre-requisitos antes de executar o fluxo.',
            'passos': checklist[:3],
        })
    if passos:
        etapas.append({
            'ordem': len(etapas) + 1,
            'titulo': 'Execucao',
            'descricao': 'Siga a sequencia principal da operacao.',
            'passos': passos[:4],
        })
        if len(passos) > 4:
            etapas.append({
                'ordem': len(etapas) + 1,
                'titulo': 'Conferencia',
                'descricao': 'Finalize revisando as ultimas validacoes do processo.',
                'passos': passos[4:],
            })
    return etapas


def _passos_topico_ajuda(topico):
    passos = list(topico.get('passos') or [])
    for etapa in _etapas_topico_ajuda(topico):
        passos.extend(etapa.get('passos') or [])
    return _textos_unicos_ajuda(passos)

AJUDA_MENU_TOPICOS = {
    'Dashboard': ('primeiros-passos', 'duvidas-acesso', 'app-mobile'),
    'Gestao': (),
    'Financeiro': ('financeiro-lancamentos',),
    'Vendas': ('vendas-pdv',),
    'Estoque': ('estoque-operacao',),
    'Recebimento': ('estoque-operacao',),
    'Expedicao': ('expedicao-entregas',),
    'Meu RH': ('rh-seguranca',),
    'E-commerce': ('ecommerce-config',),
    'Servicos': ('servicos-tecnicos',),
    'Ajuda': (),
}

AJUDA_MENU_DESCRICOES = {
    'Dashboard': 'Orientacoes iniciais, acesso ao sistema e uso no celular.',
    'Gestao': 'Fluxos administrativos e configuracoes gerais do negocio.',
    'Financeiro': 'Lancamentos, fundos e conferencia financeira da operacao.',
    'Vendas': 'Atendimento, PDV, mesas, caixas e andamento dos pedidos.',
    'Estoque': 'Cadastros, enderecamento, ajustes internos e consultas operacionais.',
    'Recebimento': 'Entradas de fornecedores, conferencia e armazenagem por tipo de recebimento.',
    'Expedicao': 'Separacao, transferencias entre lojas/CDs e entrega de pedidos.',
    'Meu RH': 'Equipe, acessos, perfis, auditoria e estrutura organizacional.',
    'E-commerce': 'Ativacao da loja online, identidade visual, campanhas, cupons e integracoes.',
    'Servicos': 'Chamados internos, suporte tecnico e atendimentos de manutencao.',
}


def _rotulo_paginas_fluxo(paginas):
    rotulos = []
    for pagina in paginas:
        titulo = PAGINAS_SISTEMA.get(pagina)
        if titulo and titulo not in rotulos:
            rotulos.append(titulo)
    return rotulos


def _paginas_da_secao_menu(secao_nome):
    for nome, paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        if nome == secao_nome:
            return paginas
    return ()


def _enriquecer_topico_ajuda(topico, secao_nome=None, secao_paginas=()):
    topico_formatado = dict(topico)
    paginas_topico = list(topico.get('paginas') or [])
    fluxo_paginas = []
    if secao_paginas:
        fluxo_paginas = [pagina for pagina in secao_paginas if pagina in paginas_topico]
    if not fluxo_paginas:
        fluxo_paginas = paginas_topico
    etapas = _etapas_topico_ajuda(topico)
    topico_formatado['menu_secao'] = secao_nome
    topico_formatado['fluxo_paginas'] = _rotulo_paginas_fluxo(fluxo_paginas)
    topico_formatado['etapas'] = etapas
    topico_formatado['passos_tutorial'] = _passos_topico_ajuda(topico)
    topico_formatado['etapas_total'] = len(etapas)
    return topico_formatado


def _organizar_topicos_ajuda_por_menu(topicos):
    topicos_map = {item['slug']: item for item in topicos}
    grupos = []
    topicos_ordenados = []
    slugs_usados = set()

    for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        slugs_secao = AJUDA_MENU_TOPICOS.get(secao_nome, ())
        topicos_secao = []
        for slug in slugs_secao:
            topico = topicos_map.get(slug)
            if not topico:
                continue
            topico_formatado = _enriquecer_topico_ajuda(topico, secao_nome, secao_paginas)
            topicos_secao.append(topico_formatado)
            topicos_ordenados.append(topico_formatado)
            slugs_usados.add(slug)
        if topicos_secao:
            grupos.append({
                'secao': secao_nome,
                'descricao': AJUDA_MENU_DESCRICOES.get(secao_nome),
                'topicos': topicos_secao,
            })

    topicos_restantes = []
    for topico in sorted(topicos_map.values(), key=lambda item: item.get('titulo') or ''):
        if topico['slug'] in slugs_usados:
            continue
        topico_formatado = _enriquecer_topico_ajuda(topico)
        topicos_restantes.append(topico_formatado)
        topicos_ordenados.append(topico_formatado)

    if topicos_restantes:
        grupos.append({
            'secao': 'Outros fluxos',
            'descricao': 'Guias complementares liberados para o seu perfil.',
            'topicos': topicos_restantes,
        })

    return grupos, topicos_ordenados


def _topicos_relacionados_ajuda(topico_atual, paginas_permitidas):
    grupo_atual = None
    for secao_nome, slugs_secao in AJUDA_MENU_TOPICOS.items():
        if topico_atual.get('slug') in slugs_secao:
            grupo_atual = secao_nome
            break

    relacionados_mesmo_grupo = []
    relacionados_outros = []

    for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        slugs_secao = AJUDA_MENU_TOPICOS.get(secao_nome, ())
        for slug in slugs_secao:
            if slug == topico_atual.get('slug'):
                continue
            item = AJUDA_TOPICOS.get(slug)
            if not item:
                continue
            paginas_item = set(item.get('paginas') or [])
            if paginas_item and not paginas_item.intersection(paginas_permitidas):
                continue
            topico_formatado = _enriquecer_topico_ajuda(item, secao_nome, secao_paginas)
            if secao_nome == grupo_atual:
                relacionados_mesmo_grupo.append(topico_formatado)
            else:
                relacionados_outros.append(topico_formatado)

    for slug, item in sorted(AJUDA_TOPICOS.items(), key=lambda entry: (entry[1].get('titulo') or '')):
        if slug == topico_atual.get('slug'):
            continue
        if slug in {rel['slug'] for rel in relacionados_mesmo_grupo + relacionados_outros}:
            continue
        paginas_item = set(item.get('paginas') or [])
        if paginas_item and not paginas_item.intersection(paginas_permitidas):
            continue
        relacionados_outros.append(_enriquecer_topico_ajuda(item))

    return grupo_atual, relacionados_mesmo_grupo + relacionados_outros


ASSISTENTE_PAGINA_ENDPOINT_PRINCIPAL = {
    'inicio': 'boas_vindas',
    'gestao_negocio': 'gestao_negocio',
    'financeiro': 'financeiro_lancamentos',
    'pdv': 'pdv',
    'estoques': 'listar_estoques',
    'produtos': 'listar_produtos',
    'categorias': 'listar_categorias',
    'fornecedores': 'listar_fornecedores',
    'enderecos_estoque': 'listar_enderecos_estoque',
    'movimentacoes': 'listar_movimentacoes',
    'almoxarifado': 'listar_almoxarifado',
    'recebimentos': 'listar_recebimentos_fornecedor',
    'relatorios': 'relatorios',
    'equipamentos_estoque': 'listar_equipamentos_movimentacao',
    'enderecos_inteligentes': 'enderecos_inteligentes',
    'caixas': 'listar_caixas',
    'mesas': 'listar_mesas',
    'pedidos': 'listar_pedidos',
    'expedicao': 'central_expedicao',
    'transferencias_estoque': 'listar_transferencias_estoque',
    'funcionarios': 'listar_funcionarios',
    'rh_funcoes': 'listar_funcoes_rh',
    'rh_indicadores': 'indicadores_rh',
    'rh_organograma': 'organograma_rh',
    'auditoria': 'auditoria_sistema',
    'empresa': 'editar_empresa',
    'ecommerce_config': 'configurar_ecommerce',
    'servicos_tecnicos': 'listar_ordens_servico',
    'chamados_internos': 'listar_chamados_internos',
    'garcons': 'listar_garcons',
    'ajuda': 'central_ajuda',
}


def _endpoint_principal_assistente_pagina(pagina):
    endpoint = ASSISTENTE_PAGINA_ENDPOINT_PRINCIPAL.get(pagina)
    if endpoint:
        return endpoint
    endpoints = sorted(PAGINA_ENDPOINTS.get(pagina) or [])
    return endpoints[0] if endpoints else None


def _url_assistente(endpoint, **kwargs):
    if not endpoint:
        return None
    try:
        return url_for(endpoint, **kwargs)
    except Exception:
        return None


def _secao_menu_assistente_topico(topico):
    slug = topico.get('slug')
    for secao_nome, slugs in AJUDA_MENU_TOPICOS.items():
        if slug in slugs:
            return secao_nome
    for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        if set(topico.get('paginas') or []).intersection(secao_paginas):
            return secao_nome
    return 'Ajuda'


def _resumo_topico_assistente(topico):
    partes = [
        topico.get('resumo'),
        topico.get('objetivo'),
    ]
    etapas = [item.get('titulo') for item in _etapas_topico_ajuda(topico) if item.get('titulo')]
    if etapas:
        partes.append('Etapas: ' + ' '.join(etapas[:3]))
    passos = _passos_topico_ajuda(topico)
    if passos:
        partes.append('Fluxo: ' + ' '.join(passos[:3]))
    duvidas = [item.get('pergunta') for item in (topico.get('duvidas') or []) if item.get('pergunta')]
    if duvidas:
        partes.append('Duvidas comuns: ' + ' '.join(duvidas[:2]))
    alertas = [item for item in (topico.get('alertas') or []) if item]
    if alertas:
        partes.append('Atencao: ' + ' '.join(alertas[:2]))
    return ' '.join(item.strip() for item in partes if item).strip()


def _acoes_assistente_para_paginas(paginas):
    acoes = []
    for pagina in paginas or []:
        endpoint = _endpoint_principal_assistente_pagina(pagina)
        url = _url_assistente(endpoint)
        if not url:
            continue
        acoes.append({
            'label': f'Abrir {PAGINAS_SISTEMA.get(pagina, pagina)}',
            'url': url,
            'page': pagina,
            'reason': PAGINAS_SISTEMA.get(pagina, pagina),
            'kind': 'navigate',
        })
    return acoes


def _construir_documentos_assistente_local():
    documentos = []
    topicos_por_pagina = {}

    def _deduplicar_textos(itens):
        vistos = set()
        resultado = []
        for item in itens or []:
            texto = (item or '').strip()
            chave = texto.lower()
            if not texto or chave in vistos:
                continue
            vistos.add(chave)
            resultado.append(texto)
        return resultado

    def _deduplicar_problemas(itens):
        vistos = set()
        resultado = []
        for item in itens or []:
            situacao = (item.get('situacao') or '').strip()
            acao = (item.get('acao') or '').strip()
            chave = (situacao.lower(), acao.lower())
            if not situacao or not acao or chave in vistos:
                continue
            vistos.add(chave)
            resultado.append({
                'situation': situacao,
                'action': acao,
            })
        return resultado

    def _faq_pairs(topico):
        pares = []
        for duvida in topico.get('duvidas') or []:
            pergunta = (duvida.get('pergunta') or '').strip()
            resposta = (duvida.get('resposta') or '').strip()
            if pergunta and resposta:
                pares.append({
                    'question': pergunta,
                    'answer': resposta,
                })
        return pares

    for topico in AJUDA_TOPICOS.values():
        for pagina in topico.get('paginas') or []:
            topicos_por_pagina.setdefault(pagina, []).append(topico)

    for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        for pagina in secao_paginas:
            titulo = PAGINAS_SISTEMA.get(pagina)
            if not titulo:
                continue
            topicos_relacionados = topicos_por_pagina.get(pagina, [])
            resumo_relacionado = ' '.join(
                item.get('resumo')
                for item in topicos_relacionados[:2]
                if item.get('resumo')
            ).strip()
            snippet = resumo_relacionado or f'Tela principal de {titulo} na secao {secao_nome}.'
            faq_pairs = []
            passos_relacionados = []
            checklist_relacionado = []
            alertas_relacionados = []
            problemas_relacionados = []
            palavras_chave = {
                pagina,
                titulo,
                secao_nome,
                'navegacao',
                'como usar',
                'onde fica',
            }
            for item in topicos_relacionados:
                if item.get('titulo'):
                    palavras_chave.add(item['titulo'])
                if item.get('slug'):
                    palavras_chave.add(item['slug'].replace('-', ' '))
                for etapa in _etapas_topico_ajuda(item):
                    if etapa.get('titulo'):
                        palavras_chave.add(etapa['titulo'])
                faq_pairs.extend(_faq_pairs(item))
                passos_relacionados.extend(_passos_topico_ajuda(item))
                checklist_relacionado.extend(item.get('checklist') or [])
                alertas_relacionados.extend(item.get('alertas') or [])
                problemas_relacionados.extend(item.get('problemas') or [])

            endpoint = _endpoint_principal_assistente_pagina(pagina)
            documentos.append({
                'id': f'page:{pagina}',
                'kind': 'page',
                'page': pagina,
                'pages': [pagina],
                'page_labels': [titulo],
                'section': secao_nome,
                'title': titulo,
                'summary': snippet,
                'snippet': snippet,
                'keywords': sorted(palavras_chave),
                'faq_pairs': faq_pairs,
                'steps': _deduplicar_textos(passos_relacionados),
                'checklist': _deduplicar_textos(checklist_relacionado),
                'alerts': _deduplicar_textos(alertas_relacionados),
                'problems': _deduplicar_problemas(problemas_relacionados),
                'url': _url_assistente(endpoint),
                'actions': _acoes_assistente_para_paginas([pagina]) + (
                    [{
                        'label': 'Abrir guia',
                        'url': _url_assistente('central_ajuda'),
                        'page': 'ajuda',
                        'reason': 'Ajuda e treinamento',
                        'kind': 'guide',
                    }] if _url_assistente('central_ajuda') else []
                ),
            })

    for topico in AJUDA_TOPICOS.values():
        secao_nome = _secao_menu_assistente_topico(topico)
        paginas = list(topico.get('paginas') or [])
        resumo = _resumo_topico_assistente(topico)
        faq_pairs = _faq_pairs(topico)
        problems = _deduplicar_problemas(topico.get('problemas') or [])
        palavras_chave = {
            topico.get('slug', '').replace('-', ' '),
            topico.get('titulo'),
            secao_nome,
            'fluxo',
            'passo a passo',
            'etapas',
            'tutorial',
            'ajuda',
            'como fazer',
        }
        for etapa in _etapas_topico_ajuda(topico):
            if etapa.get('titulo'):
                palavras_chave.add(etapa['titulo'])
        for pergunta in topico.get('duvidas') or []:
            if pergunta.get('pergunta'):
                palavras_chave.add(pergunta['pergunta'])

        documentos.append({
            'id': f"topic:{topico.get('slug')}",
            'kind': 'topic',
            'page': paginas[0] if paginas else 'ajuda',
            'pages': paginas or ['ajuda'],
            'page_labels': [PAGINAS_SISTEMA.get(item, item) for item in (paginas or ['ajuda'])],
            'section': secao_nome,
            'title': topico.get('titulo'),
            'summary': topico.get('resumo') or topico.get('objetivo'),
            'snippet': resumo,
            'keywords': sorted(item for item in palavras_chave if item),
            'faq_pairs': faq_pairs,
            'steps': _deduplicar_textos(_passos_topico_ajuda(topico)),
            'checklist': _deduplicar_textos(topico.get('checklist') or []),
            'alerts': _deduplicar_textos(topico.get('alertas') or []),
            'problems': problems,
            'source_topic': topico.get('titulo'),
            'url': _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')),
            'actions': (
                [{
                    'label': 'Abrir guia detalhado',
                    'url': _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')),
                    'page': 'ajuda',
                    'reason': topico.get('titulo'),
                    'kind': 'guide',
                }] if _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')) else []
            ) + _acoes_assistente_para_paginas(paginas[:3]),
        })

        for indice, item in enumerate(faq_pairs, start=1):
            documentos.append({
                'id': f"faq:{topico.get('slug')}:{indice}",
                'kind': 'faq',
                'page': paginas[0] if paginas else 'ajuda',
                'pages': paginas or ['ajuda'],
                'page_labels': [PAGINAS_SISTEMA.get(pagina, pagina) for pagina in (paginas or ['ajuda'])],
                'section': secao_nome,
                'title': item.get('question'),
                'summary': item.get('answer'),
                'snippet': f"Duvida comum em {topico.get('titulo')}: {item.get('answer')}",
                'keywords': sorted({
                    topico.get('titulo'),
                    secao_nome,
                    topico.get('slug', '').replace('-', ' '),
                    item.get('question'),
                    'duvida comum',
                    'resposta',
                } - {None, ''}),
                'faq_pairs': [item],
                'steps': _deduplicar_textos(_passos_topico_ajuda(topico)),
                'checklist': _deduplicar_textos(topico.get('checklist') or []),
                'alerts': _deduplicar_textos(topico.get('alertas') or []),
                'problems': problems,
                'source_topic': topico.get('titulo'),
                'url': _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')),
                'actions': (
                    [{
                        'label': 'Abrir guia detalhado',
                        'url': _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')),
                        'page': 'ajuda',
                        'reason': topico.get('titulo'),
                        'kind': 'guide',
                    }] if _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')) else []
                ) + _acoes_assistente_para_paginas(paginas[:3]),
            })

        for indice, item in enumerate(problems, start=1):
            documentos.append({
                'id': f"issue:{topico.get('slug')}:{indice}",
                'kind': 'issue',
                'page': paginas[0] if paginas else 'ajuda',
                'pages': paginas or ['ajuda'],
                'page_labels': [PAGINAS_SISTEMA.get(pagina, pagina) for pagina in (paginas or ['ajuda'])],
                'section': secao_nome,
                'title': item.get('situation'),
                'summary': item.get('action'),
                'snippet': f"Tratativa sugerida em {topico.get('titulo')}: {item.get('action')}",
                'keywords': sorted({
                    topico.get('titulo'),
                    secao_nome,
                    topico.get('slug', '').replace('-', ' '),
                    item.get('situation'),
                    'problema',
                    'nao consigo',
                    'erro',
                } - {None, ''}),
                'faq_pairs': faq_pairs,
                'steps': _deduplicar_textos(_passos_topico_ajuda(topico)),
                'checklist': _deduplicar_textos(topico.get('checklist') or []),
                'alerts': _deduplicar_textos(topico.get('alertas') or []),
                'problems': [item],
                'source_topic': topico.get('titulo'),
                'url': _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')),
                'actions': (
                    [{
                        'label': 'Abrir guia detalhado',
                        'url': _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')),
                        'page': 'ajuda',
                        'reason': topico.get('titulo'),
                        'kind': 'guide',
                    }] if _url_assistente('detalhe_ajuda', topico_slug=topico.get('slug')) else []
                ) + _acoes_assistente_para_paginas(paginas[:3]),
            })

    documentos.append({
        'id': 'topic:home-operacional',
        'kind': 'topic',
        'page': 'inicio',
        'pages': ['inicio', 'pedidos', 'recebimentos', 'expedicao', 'financeiro'],
        'page_labels': [PAGINAS_SISTEMA.get(item, item) for item in ['inicio', 'pedidos', 'recebimentos', 'expedicao', 'financeiro']],
        'section': 'Dashboard',
        'title': 'Home Operacional',
        'summary': 'Painel inicial com prioridades do dia, atalhos e fluxo recomendado.',
        'snippet': 'Use a Home Operacional para ver o que esta pendente, entrar nas filas de trabalho e abrir os atalhos mais usados do seu perfil.',
        'keywords': ['home operacional', 'prioridades', 'pendencias', 'inicio', 'o que fazer hoje'],
        'faq_pairs': [],
        'steps': [],
        'checklist': [],
        'alerts': [],
        'problems': [],
        'url': _url_assistente('boas_vindas'),
        'actions': _acoes_assistente_para_paginas(['inicio', 'ajuda']),
    })

    return documentos


@app.route('/ajuda')
@login_required
def central_ajuda():
    funcionario = get_funcionario_logado()
    perfil_usuario = (funcionario.role if funcionario else 'operador')
    paginas_permitidas = _paginas_permitidas_para_funcionario(funcionario)
    topicos_ajuda = []
    for topico in AJUDA_TOPICOS.values():
        paginas_topico = set(topico.get('paginas') or [])
        if not paginas_topico or paginas_topico.intersection(paginas_permitidas):
            topicos_ajuda.append(topico)

    grupos_ajuda_menu, topicos_ordenados = _organizar_topicos_ajuda_por_menu(topicos_ajuda)

    faq_rapido = []
    for topico in topicos_ordenados:
        for duvida in (topico.get('duvidas') or []):
            faq_rapido.append({
                'topico_slug': topico.get('slug'),
                'topico_titulo': topico.get('titulo'),
                'pergunta': duvida.get('pergunta'),
            })
            if len(faq_rapido) >= 10:
                break
        if len(faq_rapido) >= 10:
            break

    resumo_ajuda = {
        'topicos': len(topicos_ordenados),
        'passos': sum(len(topico.get('passos_tutorial') or []) for topico in topicos_ordenados),
        'etapas': sum(len(topico.get('etapas') or []) for topico in topicos_ordenados),
        'fluxogramas': sum(1 for topico in topicos_ordenados if topico.get('fluxograma')),
    }

    return render_template(
        'sistema/ajuda.html',
        topicos_ajuda=topicos_ordenados,
        grupos_ajuda_menu=grupos_ajuda_menu,
        faq_rapido=faq_rapido,
        resumo_ajuda=resumo_ajuda,
        perfil_usuario=perfil_usuario,
        cargo_usuario=(funcionario.cargo if funcionario else None),
    )


@app.route('/ajuda/<string:topico_slug>')
@login_required
def detalhe_ajuda(topico_slug):
    topico = AJUDA_TOPICOS.get(topico_slug)
    if not topico:
        flash('Topico de ajuda nao encontrado.', 'warning')
        return redirect(url_for('central_ajuda'))

    funcionario = get_funcionario_logado()
    paginas_permitidas = _paginas_permitidas_para_funcionario(funcionario)
    paginas_topico = set(topico.get('paginas') or [])
    if paginas_topico and not paginas_topico.intersection(paginas_permitidas):
        flash('Este guia nao esta disponivel para o seu perfil de acesso.', 'warning')
        return redirect(url_for('central_ajuda'))

    grupo_menu_atual, relacionados = _topicos_relacionados_ajuda(topico, paginas_permitidas)
    return render_template(
        'sistema/ajuda_detalhe.html',
        topico=_enriquecer_topico_ajuda(topico, grupo_menu_atual, _paginas_da_secao_menu(grupo_menu_atual)),
        topicos_relacionados=relacionados,
        grupo_menu_atual=grupo_menu_atual,
    )


# ============ ROTAS - FUNCIONARIOS ============

@app.route('/funcionarios')
@require_role('admin', 'gerente')
def listar_funcionarios():
    page = max(request.args.get('page', 1, type=int), 1)
    per_page = min(max(request.args.get('per_page', 25, type=int), 10), 100)
    busca = (request.args.get('busca') or '').strip()
    role = (request.args.get('role') or '').strip().lower()
    status = (request.args.get('status') or '').strip().lower()
    departamento = (request.args.get('departamento') or '').strip()
    ordenar = (request.args.get('ordenar') or 'nome_asc').strip().lower()

    query = Funcionario.query.options(
        selectinload(Funcionario.superior),
        selectinload(Funcionario.perfil_acesso),
    )
    if busca:
        termo = f'%{busca}%'
        query = query.filter(
            db.or_(
                Funcionario.nome.ilike(termo),
                Funcionario.email.ilike(termo),
                Funcionario.matricula.ilike(termo),
                db.cast(Funcionario.numero_cadastro, db.String).ilike(termo),
                Funcionario.cargo.ilike(termo),
                Funcionario.departamento.ilike(termo),
                Funcionario.time_nome.ilike(termo),
                Funcionario.cpf.ilike(termo),
            )
        )
    if role in ROLES_PERMITIDOS:
        query = query.filter(Funcionario.role == role)
    if status == 'ativos':
        query = query.filter(Funcionario.ativo.is_(True))
    elif status == 'inativos':
        query = query.filter(Funcionario.ativo.is_(False))
    if departamento:
        query = query.filter(Funcionario.departamento == departamento)

    ordenacoes = {
        'nome_asc': (Funcionario.nome.asc(), Funcionario.id.asc()),
        'nome_desc': (Funcionario.nome.desc(), Funcionario.id.desc()),
        'recentes': (Funcionario.criado_em.desc(), Funcionario.id.desc()),
        'cargo': (Funcionario.cargo.asc(), Funcionario.nome.asc()),
        'departamento': (Funcionario.departamento.asc(), Funcionario.nome.asc()),
        'perfil': (Funcionario.role.asc(), Funcionario.nome.asc()),
    }
    if ordenar not in ordenacoes:
        ordenar = 'nome_asc'

    paginacao = query.order_by(*ordenacoes[ordenar]).paginate(page=page, per_page=per_page, error_out=False)
    funcionarios = paginacao.items
    departamentos_disponiveis = sorted({
        (item[0] or '').strip()
        for item in db.session.query(Funcionario.departamento).distinct().all()
        if (item[0] or '').strip()
    }, key=str.lower)
    resumo_funcionarios = {
        'total': Funcionario.query.count(),
        'ativos': Funcionario.query.filter(Funcionario.ativo.is_(True)).count(),
        'inativos': Funcionario.query.filter(Funcionario.ativo.is_(False)).count(),
        'filtrados': paginacao.total,
    }
    return render_template(
        'funcionarios/listar.html',
        funcionarios=funcionarios,
        paginacao=paginacao,
        resumo_funcionarios=resumo_funcionarios,
        departamentos_disponiveis=departamentos_disponiveis,
        filtros={
            'busca': busca,
            'role': role,
            'status': status,
            'departamento': departamento,
            'ordenar': ordenar,
            'per_page': per_page,
        },
        primeiro_funcionario_id=_primeiro_funcionario_id(),
    )


@app.route('/funcionarios/novo', methods=['GET', 'POST'])
@require_role('admin')
def criar_funcionario():
    if request.method == 'POST':
        nome = request.form.get('nome', '').strip()
        email = request.form.get('email', '').strip()
        senha = request.form.get('senha', '')
        confirmacao_senha = request.form.get('confirmacao_senha', '')
        role = _normalizar_texto(request.form.get('role', 'operador'))
        cargo = (request.form.get('cargo') or '').strip()
        perfil_acesso_id = request.form.get('perfil_acesso_id', type=int)
        cpf = _normalizar_cpf(request.form.get('cpf'))
        rg = _normalizar_campo_organograma(request.form.get('rg'))
        data_nascimento = _parse_date_iso(request.form.get('data_nascimento'))
        celular = _normalizar_campo_organograma(request.form.get('celular'))
        cep = _normalizar_campo_organograma(request.form.get('cep'))
        endereco = _normalizar_campo_organograma(request.form.get('endereco'))
        bairro = _normalizar_campo_organograma(request.form.get('bairro'))
        cidade = _normalizar_campo_organograma(request.form.get('cidade'))
        estado = _normalizar_estado(request.form.get('estado'))
        superior_id = request.form.get('superior_id', type=int)
        departamento = _normalizar_campo_organograma(request.form.get('departamento'))
        time_nome = _normalizar_campo_organograma(request.form.get('time_nome'))
        nivel_organograma = _normalizar_campo_organograma(request.form.get('nivel_organograma'))
        permitir_editar_imagem_perfil = (request.form.get('permitir_editar_imagem_perfil') == 'on')
        vinculos_estoque, erro_vinculos_estoque = _resolver_vinculos_estoque_funcionario(request)
        perfil_acesso = None

        if not nome or not email or not senha:
            flash('Nome, email e senha são obrigatórios.', 'danger')
            return redirect(url_for('criar_funcionario'))

        if senha != confirmacao_senha:
            flash('As senhas não conferem.', 'danger')
            return redirect(url_for('criar_funcionario'))

        if len(senha) < 6:
            flash('A senha deve ter no minimo 6 caracteres.', 'danger')
            return redirect(url_for('criar_funcionario'))

        if Funcionario.query.filter_by(email=email).first():
            flash('Email ja cadastrado.', 'danger')
            return redirect(url_for('criar_funcionario'))


        if cpf == '__invalid__':
            flash('CPF invalido. Informe 11 digitos.', 'danger')
            return redirect(url_for('criar_funcionario'))
        if cpf and Funcionario.query.filter_by(cpf=cpf).first():
            flash('CPF ja cadastrado para outro funcionario.', 'danger')
            return redirect(url_for('criar_funcionario'))

        if role not in ROLES_PERMITIDOS:
            flash('Tipo de usuario invalido.', 'danger')
            return redirect(url_for('criar_funcionario'))

        if perfil_acesso_id:
            perfil_acesso = PerfilAcesso.query.get(perfil_acesso_id)
            if not perfil_acesso:
                flash('Perfil de acesso padrao invalido.', 'danger')
                return redirect(url_for('criar_funcionario'))

        if nivel_organograma and nivel_organograma not in NIVEIS_ORGANOGRAMA:
            flash('Nivel de organograma invalido.', 'danger')
            return redirect(url_for('criar_funcionario'))

        if erro_vinculos_estoque:
            flash(erro_vinculos_estoque, 'danger')
            return redirect(url_for('criar_funcionario'))

        try:
            validate_employee_payload(
                nome=nome,
                email=email,
                role=role,
                cargo=(cargo or _role_para_cargo_padrao(role)),
                departamento=departamento,
                ativo=True,
                controle_acesso_ativo=bool(perfil_acesso),
                perfil_acesso_id=(perfil_acesso.id if perfil_acesso else None),
                restricao_estoques_ativa=bool(vinculos_estoque and vinculos_estoque['restricao_estoques_ativa']),
                estoque_principal_id=(
                    vinculos_estoque['estoque_principal'].id
                    if vinculos_estoque and vinculos_estoque['estoque_principal']
                    else None
                ),
            )
        except ValidationError as exc:
            flash(str(exc), 'danger')
            return redirect(url_for('criar_funcionario'))

        nivel_hierarquico = nivel_organograma or _role_para_nivel_organograma(role)
        if not superior_id and _deve_exigir_superior_hierarquico(role, nivel_hierarquico):
            flash('Informe o responsavel direto do colaborador para manter o organograma organizado.', 'danger')
            return redirect(url_for('criar_funcionario'))

        superior = None
        if superior_id:
            superior = Funcionario.query.get(superior_id)
            if not superior:
                flash('Superior hierarquico invalido.', 'danger')
                return redirect(url_for('criar_funcionario'))

        novo_funcionario = Funcionario(
            nome=nome,
            email=email,
            matricula=None,
            cpf=(cpf if cpf != '__invalid__' else None),
            rg=rg,
            data_nascimento=data_nascimento,
            celular=celular,
            cep=cep,
            endereco=endereco,
            bairro=bairro,
            cidade=cidade,
            estado=estado,
            permitir_editar_imagem_perfil=permitir_editar_imagem_perfil,
            role=role,
            cargo=cargo or _role_para_cargo_padrao(role),
            perfil_acesso_id=(perfil_acesso.id if perfil_acesso else None),
            departamento=departamento,
            time_nome=time_nome,
            nivel_organograma=nivel_hierarquico,
            restricao_estoques_ativa=bool(vinculos_estoque and vinculos_estoque['restricao_estoques_ativa']),
            estoque_principal_id=(
                vinculos_estoque['estoque_principal'].id
                if vinculos_estoque and vinculos_estoque['estoque_principal']
                else None
            ),
            superior_id=(superior.id if superior else None),
            controle_acesso_ativo=bool(perfil_acesso),
        )
        novo_funcionario.set_password(senha)
        db.session.add(novo_funcionario)
        db.session.flush()
        novo_funcionario.estoques_permitidos = (
            vinculos_estoque['estoques_permitidos']
            if vinculos_estoque else []
        )
        novo_funcionario.numero_cadastro = _gerar_numero_cadastro_unico(novo_funcionario)
        novo_funcionario.matricula = _gerar_matricula_unica(novo_funcionario)
        sincronizar_garcom_funcionario(novo_funcionario)
        db.session.commit()

        flash(f'Funcionario {nome} criado com sucesso!', 'success')
        return redirect(url_for('listar_funcionarios'))

    funcoes_rh = FuncaoRH.query.filter_by(ativo=True).order_by(FuncaoRH.nome.asc()).all()
    perfis_acesso = PerfilAcesso.query.filter_by(ativo=True).order_by(PerfilAcesso.nome.asc()).all()
    superiores = Funcionario.query.filter_by(ativo=True).order_by(Funcionario.nome.asc()).all()
    departamentos_existentes, times_existentes = _listar_cadastros_organograma()
    return render_template(
        'funcionarios/criar.html',
        funcoes_rh=funcoes_rh,
        perfis_acesso=perfis_acesso,
        superiores=superiores,
        estoques_cadastrados=_listar_estoques_para_vinculo_funcionario(),
        estoques_permitidos_ids=[],
        niveis_organograma=NIVEIS_ORGANOGRAMA,
        departamentos_existentes=departamentos_existentes,
        times_existentes=times_existentes,
    )


@app.route('/funcionarios/<int:funcionario_id>/editar', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def editar_funcionario(funcionario_id):
    funcionario = Funcionario.query.get_or_404(funcionario_id)

    funcionario_logado = get_funcionario_logado()
    if funcionario_logado.role == 'gerente' and funcionario.role in ['admin', 'gerente'] and funcionario.id != funcionario_logado.id:
        flash('Você não tem permissão para editar este funcionário.', 'danger')
        return redirect(url_for('listar_funcionarios'))

    if request.method == 'POST':
        nome = request.form.get('nome', '').strip()
        email = request.form.get('email', '').strip()
        cpf = _normalizar_cpf(request.form.get('cpf'))
        rg = _normalizar_campo_organograma(request.form.get('rg'))
        data_nascimento = _parse_date_iso(request.form.get('data_nascimento'))
        celular = _normalizar_campo_organograma(request.form.get('celular'))
        cep = _normalizar_campo_organograma(request.form.get('cep'))
        endereco = _normalizar_campo_organograma(request.form.get('endereco'))
        bairro = _normalizar_campo_organograma(request.form.get('bairro'))
        cidade = _normalizar_campo_organograma(request.form.get('cidade'))
        estado = _normalizar_estado(request.form.get('estado'))
        role = _normalizar_texto(request.form.get('role', funcionario.role))
        cargo = (request.form.get('cargo') or '').strip()
        perfil_acesso_id = request.form.get('perfil_acesso_id', type=int)
        superior_id = request.form.get('superior_id', type=int)
        departamento = _normalizar_campo_organograma(request.form.get('departamento'))
        time_nome = _normalizar_campo_organograma(request.form.get('time_nome'))
        nivel_organograma = _normalizar_campo_organograma(request.form.get('nivel_organograma'))
        permitir_editar_imagem_perfil = (request.form.get('permitir_editar_imagem_perfil') == 'on')
        ativo = request.form.get('ativo') == 'on'
        nova_senha = request.form.get('nova_senha', '')
        vinculos_estoque, erro_vinculos_estoque = _resolver_vinculos_estoque_funcionario(request)
        perfil_acesso = None

        if not nome or not email:
            flash('Nome e email são obrigatórios.', 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        outro_func = Funcionario.query.filter_by(email=email).first()
        if outro_func and outro_func.id != funcionario.id:
            flash('Email ja cadastrado por outro funcionario.', 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))


        if cpf == '__invalid__':
            flash('CPF invalido. Informe 11 digitos.', 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))
        if cpf:
            outro_cpf = Funcionario.query.filter(
                Funcionario.cpf == cpf,
                Funcionario.id != funcionario.id
            ).first()
            if outro_cpf:
                flash('CPF ja cadastrado para outro funcionario.', 'danger')
                return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        if erro_vinculos_estoque:
            flash(erro_vinculos_estoque, 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        if perfil_acesso_id:
            perfil_acesso = PerfilAcesso.query.get(perfil_acesso_id)
            if not perfil_acesso:
                flash('Perfil de acesso padrao invalido.', 'danger')
                return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        try:
            validate_employee_payload(
                nome=nome,
                email=email,
                role=(role if funcionario_logado.role == 'admin' else funcionario.role),
                cargo=(cargo or funcionario.cargo or _role_para_cargo_padrao(funcionario.role)),
                departamento=departamento,
                ativo=ativo,
                controle_acesso_ativo=bool(perfil_acesso or funcionario.controle_acesso_ativo),
                perfil_acesso_id=(perfil_acesso.id if perfil_acesso else funcionario.perfil_acesso_id),
                restricao_estoques_ativa=bool(vinculos_estoque and vinculos_estoque['restricao_estoques_ativa']),
                estoque_principal_id=(
                    vinculos_estoque['estoque_principal'].id
                    if vinculos_estoque and vinculos_estoque['estoque_principal']
                    else None
                ),
            )
        except ValidationError as exc:
            flash(str(exc), 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        funcionario.nome = nome
        funcionario.email = email
        funcionario.cpf = cpf if cpf != '__invalid__' else None
        funcionario.rg = rg
        funcionario.data_nascimento = data_nascimento
        funcionario.celular = celular
        funcionario.cep = cep
        funcionario.endereco = endereco
        funcionario.bairro = bairro
        funcionario.cidade = cidade
        funcionario.estado = estado
        funcionario.permitir_editar_imagem_perfil = permitir_editar_imagem_perfil
        funcionario.cargo = cargo or funcionario.cargo or _role_para_cargo_padrao(funcionario.role)
        funcionario.perfil_acesso_id = perfil_acesso.id if perfil_acesso else None
        funcionario.departamento = departamento
        funcionario.time_nome = time_nome
        funcionario.restricao_estoques_ativa = bool(vinculos_estoque and vinculos_estoque['restricao_estoques_ativa'])
        funcionario.estoque_principal_id = (
            vinculos_estoque['estoque_principal'].id
            if vinculos_estoque and vinculos_estoque['estoque_principal']
            else None
        )
        funcionario.estoques_permitidos = (
            vinculos_estoque['estoques_permitidos']
            if vinculos_estoque else []
        )
        funcionario.ativo = ativo
        funcionario.numero_cadastro = _gerar_numero_cadastro_unico(funcionario)
        funcionario.matricula = _gerar_matricula_unica(funcionario)

        if nivel_organograma and nivel_organograma not in NIVEIS_ORGANOGRAMA:
            flash('Nivel de organograma invalido.', 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        nivel_hierarquico = nivel_organograma or _role_para_nivel_organograma(funcionario.role)
        if not superior_id and _deve_exigir_superior_hierarquico(role if funcionario_logado.role == 'admin' else funcionario.role, nivel_hierarquico):
            flash('Informe o responsavel direto do colaborador para manter o organograma organizado.', 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        if superior_id == funcionario.id:
            flash('Um funcionario nao pode ser superior de si mesmo.', 'danger')
            return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))

        superior = None
        if superior_id:
            superior = Funcionario.query.get(superior_id)
            if not superior:
                flash('Superior hierarquico invalido.', 'danger')
                return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))
            cursor = superior
            visitados = set()
            while cursor:
                if cursor.id in visitados:
                    flash('Hierarquia invalida detectada. Revise o organograma da empresa.', 'danger')
                    return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))
                visitados.add(cursor.id)
                if cursor.id == funcionario.id:
                    flash('Hierarquia invalida: ciclo detectado na estrutura da empresa.', 'danger')
                    return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))
                cursor = cursor.superior
        funcionario.superior_id = superior.id if superior else None

        if funcionario_logado.role == 'admin':
            if role not in ROLES_PERMITIDOS:
                flash('Tipo de usuario invalido.', 'danger')
                return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))
            funcionario.role = role

        if funcionario.role != 'admin' and perfil_acesso:
            funcionario.controle_acesso_ativo = True

        funcionario.nivel_organograma = nivel_hierarquico

        if nova_senha:
            if len(nova_senha) < 6:
                flash('A nova senha deve ter no minimo 6 caracteres.', 'danger')
                return redirect(url_for('editar_funcionario', funcionario_id=funcionario_id))
            funcionario.set_password(nova_senha)

        sincronizar_garcom_funcionario(funcionario)
        db.session.commit()
        flash('Funcionario atualizado com sucesso!', 'success')
        return redirect(url_for('listar_funcionarios'))

    funcoes_rh = FuncaoRH.query.filter_by(ativo=True).order_by(FuncaoRH.nome.asc()).all()
    perfis_acesso = PerfilAcesso.query.order_by(PerfilAcesso.nome.asc()).all()
    superiores = Funcionario.query.filter(
        Funcionario.id != funcionario.id,
        Funcionario.ativo.is_(True)
    ).order_by(Funcionario.nome.asc()).all()
    departamentos_existentes, times_existentes = _listar_cadastros_organograma()
    return render_template(
        'funcionarios/editar.html',
        funcionario=funcionario,
        funcoes_rh=funcoes_rh,
        funcoes_rh_nomes=[f.nome for f in funcoes_rh],
        perfis_acesso=perfis_acesso,
        superiores=superiores,
        estoques_cadastrados=_listar_estoques_para_vinculo_funcionario(),
        estoques_permitidos_ids=sorted(_estoques_permitidos_ids_funcionario(funcionario) or []),
        niveis_organograma=NIVEIS_ORGANOGRAMA,
        departamentos_existentes=departamentos_existentes,
        times_existentes=times_existentes,
    )


@app.route('/funcionarios/<int:funcionario_id>/deletar', methods=['POST'])
@require_role('admin')
def deletar_funcionario(funcionario_id):
    if funcionario_id == session.get('funcionario_id'):
        flash('Você não pode deletar sua própria conta.', 'danger')
        return redirect(url_for('listar_funcionarios'))

    funcionario = Funcionario.query.get_or_404(funcionario_id)
    try:
        Funcionario.query.filter_by(superior_id=funcionario.id).update(
            {Funcionario.superior_id: None},
            synchronize_session=False
        )
        db.session.delete(funcionario)
        db.session.commit()
        flash(f'Funcionario {funcionario.nome} deletado com sucesso!', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao deletar funcionario: {str(e)}', 'danger')

    return redirect(url_for('listar_funcionarios'))


@app.route('/funcionarios/<int:funcionario_id>/acessos', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def editar_acessos_funcionario(funcionario_id):
    funcionario = Funcionario.query.get_or_404(funcionario_id)
    funcionario_logado = get_funcionario_logado()
    primeiro_id = _primeiro_funcionario_id()
    estoques_cadastrados = _listar_estoques_para_vinculo_funcionario()
    restricao_estoques_ativa_form = bool(funcionario.restricao_estoques_ativa)
    estoque_principal_id_form = funcionario.estoque_principal_id
    estoques_permitidos_ids = sorted(_estoques_permitidos_ids_funcionario(funcionario) or [])

    if funcionario.id == primeiro_id:
        flash('O primeiro registro do sistema possui acesso protegido e nao pode ser alterado nesta tela.', 'warning')
        return redirect(url_for('listar_funcionarios'))

    if funcionario_logado and funcionario_logado.role == 'gerente':
        if funcionario.role in ['admin', 'gerente'] and funcionario.id != funcionario_logado.id:
            flash('Gerente nao pode alterar acessos de admin/gerente.', 'danger')
            return redirect(url_for('listar_funcionarios'))

    permissoes_atuais = _paginas_efetivas_funcionario(funcionario)
    if request.method == 'POST':
        paginas_enviadas = set(request.form.getlist('paginas'))
        permissoes_atuais = paginas_enviadas
        restricao_estoques_ativa_form = (request.form.get('restricao_estoques_ativa') == 'on')
        estoque_principal_id_form = request.form.get('estoque_principal_id', type=int)
        estoques_permitidos_ids = sorted({
            int(valor)
            for valor in request.form.getlist('estoques_permitidos_ids')
            if str(valor).isdigit()
        } | ({estoque_principal_id_form} if estoque_principal_id_form else set()))
        vinculos_estoque, erro_vinculo = _resolver_vinculos_estoque_funcionario(request)

        if erro_vinculo:
            flash(erro_vinculo, 'danger')
        else:
            try:
                funcionario.restricao_estoques_ativa = vinculos_estoque['restricao_estoques_ativa']
                funcionario.estoque_principal = vinculos_estoque['estoque_principal']
                funcionario.estoques_permitidos = vinculos_estoque['estoques_permitidos']
                funcionario.numero_cadastro = _gerar_numero_cadastro_unico(funcionario)
                funcionario.matricula = _gerar_matricula_unica(funcionario)
                _salvar_permissoes_funcionario(funcionario, paginas_enviadas)
                db.session.commit()
                flash(f'Acessos de {funcionario.nome} atualizados com sucesso.', 'success')
                return redirect(url_for('listar_funcionarios'))
            except Exception as e:
                db.session.rollback()
                flash(f'Erro ao salvar acessos: {str(e)}', 'danger')
    
    permissoes_personalizadas = _mapa_permissoes_personalizadas_funcionario(funcionario)
    permissoes_base = _paginas_perfil_acesso(funcionario.perfil_acesso)
    paginas_ordenadas_menu = []
    for secao_nome, secao_paginas in PAGINAS_SISTEMA_MENU_ORDEM:
        itens_secao = [
            (pagina_key, PAGINAS_SISTEMA[pagina_key])
            for pagina_key in secao_paginas
            if pagina_key in PAGINAS_SISTEMA
        ]
        if itens_secao:
            paginas_ordenadas_menu.append((secao_nome, itens_secao))
    
    return render_template(
        'funcionarios/acessos.html',
        funcionario=funcionario,
        paginas_sistema=PAGINAS_SISTEMA,
        paginas_ordenadas_menu=paginas_ordenadas_menu,
        permissoes_atuais=permissoes_atuais,
        permissoes_base=permissoes_base,
        permissoes_personalizadas=permissoes_personalizadas,
        estoques_cadastrados=estoques_cadastrados,
        restricao_estoques_ativa_form=restricao_estoques_ativa_form,
        estoque_principal_id_form=estoque_principal_id_form,
        estoques_permitidos_ids=estoques_permitidos_ids,
    )


@app.route('/rh/funcoes')
@require_role('admin', 'gerente')
def listar_funcoes_rh():
    funcoes = FuncaoRH.query.order_by(FuncaoRH.nome.asc()).all()
    funcionarios_por_cargo = {
        (funcao.nome or '').strip().lower(): Funcionario.query.filter(
            db.func.lower(Funcionario.cargo) == (funcao.nome or '').strip().lower()
        ).count()
        for funcao in funcoes
    }
    cargos_em_uso = Funcionario.query.filter(
        Funcionario.cargo.isnot(None),
        Funcionario.cargo != ''
    ).count()
    return render_template(
        'rh/funcoes.html',
        funcoes=funcoes,
        resumo_cargos={
            'total': len(funcoes),
            'ativas': sum(1 for funcao in funcoes if funcao.ativo),
            'em_uso': cargos_em_uso,
        },
        funcionarios_por_cargo=funcionarios_por_cargo,
    )


@app.route('/rh/perfis')
@require_role('admin', 'gerente')
def listar_perfis_rh():
    perfis_acesso = PerfilAcesso.query.order_by(PerfilAcesso.nome.asc()).all()
    perfis_acesso_permissoes = []
    perfis_sem_paginas = 0
    for perfil in perfis_acesso:
        permissoes = _paginas_perfil_acesso(perfil)
        if not permissoes:
            perfis_sem_paginas += 1
        perfis_acesso_permissoes.append((perfil, permissoes))

    paginas_ordenadas_menu = _paginas_ordenadas_menu()
    funcionarios_por_perfil = {
        perfil.id: Funcionario.query.filter_by(perfil_acesso_id=perfil.id).count()
        for perfil in perfis_acesso
    }
    return render_template(
        'rh/perfis.html',
        perfis_acesso_permissoes=perfis_acesso_permissoes,
        resumo_perfis={
            'total': len(perfis_acesso),
            'ativos': sum(1 for perfil in perfis_acesso if perfil.ativo),
            'com_paginas': sum(1 for _, acessos in perfis_acesso_permissoes if acessos),
            'sem_paginas': perfis_sem_paginas,
        },
        funcionarios_por_perfil=funcionarios_por_perfil,
        paginas_ordenadas_menu=paginas_ordenadas_menu,
        paginas_sistema=PAGINAS_SISTEMA,
    )


@app.route('/rh/indicadores')
@require_role('admin', 'gerente')
def indicadores_rh():
    total_funcionarios = Funcionario.query.count()
    funcionarios_ativos = Funcionario.query.filter_by(ativo=True).count()
    funcionarios_inativos = Funcionario.query.filter_by(ativo=False).count()
    acessos_controlados = Funcionario.query.filter_by(controle_acesso_ativo=True).count()

    funcoes_total = FuncaoRH.query.count()
    funcoes_ativas = FuncaoRH.query.filter_by(ativo=True).count()
    perfis_acesso_total = PerfilAcesso.query.count()
    perfis_acesso_ativos = PerfilAcesso.query.filter_by(ativo=True).count()

    data_limite = datetime.utcnow() - timedelta(days=30)
    admissoes_30_dias = Funcionario.query.filter(Funcionario.criado_em >= data_limite).count()
    inicio_hoje = datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)

    distribuicao_roles = db.session.query(
        Funcionario.role.label('role'),
        db.func.count(Funcionario.id).label('quantidade')
    ).group_by(Funcionario.role).order_by(db.desc('quantidade')).all()

    distribuicao_perfis_acesso = [
        {
            'perfil': item.perfil,
            'quantidade': item.quantidade,
        }
        for item in db.session.query(
            PerfilAcesso.nome.label('perfil'),
            db.func.count(Funcionario.id).label('quantidade')
        ).outerjoin(
            Funcionario,
            Funcionario.perfil_acesso_id == PerfilAcesso.id
        ).group_by(PerfilAcesso.id, PerfilAcesso.nome).order_by(db.desc('quantidade'), PerfilAcesso.nome.asc()).all()
    ]
    sem_perfil_padrao = Funcionario.query.filter(Funcionario.perfil_acesso_id.is_(None)).count()
    if sem_perfil_padrao:
        distribuicao_perfis_acesso.append({
            'perfil': 'Sem perfil padrao',
            'quantidade': sem_perfil_padrao,
        })

    distribuicao_cargos = db.session.query(
        Funcionario.cargo.label('cargo'),
        db.func.count(Funcionario.id).label('quantidade'),
        db.func.sum(db.case((Funcionario.ativo.is_(True), 1), else_=0)).label('ativos')
    ).group_by(Funcionario.cargo).order_by(db.desc('quantidade')).all()

    try:
        pedidos_pendentes = Pedido.query.filter(
            Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO, Pedido.STATUS_ENTREGUE])
        ).count()
        pedidos_hoje = Pedido.query.filter(Pedido.criado_em >= inicio_hoje).count()
        pedidos_30_dias = Pedido.query.filter(Pedido.criado_em >= data_limite).count()
    except (OperationalError, ProgrammingError):
        app.logger.warning('indicadores_rh: contadores de pedidos indisponiveis por schema/banco inconsistente', exc_info=True)
        db.session.rollback()
        pedidos_pendentes = 0
        pedidos_hoje = 0
        pedidos_30_dias = 0
    media_pedidos_dia_30 = (pedidos_30_dias / 30.0) if pedidos_30_dias else 0.0

    equipe_operacao_ativa = Funcionario.query.filter(
        Funcionario.ativo.is_(True),
        Funcionario.role.in_(['caixa', 'operador', 'garcom', 'gerente'])
    ).count()
    pendencias_por_colaborador = (
        pedidos_pendentes / equipe_operacao_ativa
        if equipe_operacao_ativa > 0 else float(pedidos_pendentes)
    )
    equipe_minima_recomendada = max(1, int((pedidos_pendentes + 7) // 8)) if pedidos_pendentes > 0 else 1
    deficit_equipe = max(0, equipe_minima_recomendada - equipe_operacao_ativa)

    funcoes_ativas_lista = FuncaoRH.query.filter_by(ativo=True).all()
    cargos_sem_cobertura = []
    for funcao in funcoes_ativas_lista:
        ativos_no_cargo = Funcionario.query.filter(
            Funcionario.ativo.is_(True),
            db.func.lower(Funcionario.cargo) == (funcao.nome or '').lower()
        ).count()
        if ativos_no_cargo == 0:
            cargos_sem_cobertura.append(funcao.nome)

    acoes_rh = []
    if deficit_equipe > 0:
        acoes_rh.append(f'Ajustar escala imediata: deficit atual de {deficit_equipe} colaborador(es).')
    if pendencias_por_colaborador > 8:
        acoes_rh.append('Revisar capacidade operacional: fila acima de 8 pendencias por colaborador.')
    if funcionarios_inativos > funcionarios_ativos and total_funcionarios > 0:
        acoes_rh.append('Revisar base de colaboradores inativos e normalizar quadro ativo.')
    if cargos_sem_cobertura:
        acoes_rh.append('Priorizar cobertura dos cargos sem pessoas ativas.')
    if perfis_acesso_ativos == 0:
        acoes_rh.append('Cadastrar perfis de acesso padrao para reduzir configuracoes manuais por colaborador.')
    if not acoes_rh:
        acoes_rh.append('Quadro dentro da capacidade atual. Manter monitoramento semanal.')

    funcionarios_recentes = Funcionario.query.options(
        selectinload(Funcionario.perfil_acesso)
    ).order_by(Funcionario.criado_em.desc()).limit(10).all()

    produtividade_map = {}
    pedidos_fechados_periodo = Pedido.query.filter(
        Pedido.status == Pedido.STATUS_FECHADO,
        Pedido.fechado_em.isnot(None),
        Pedido.fechado_em >= data_limite
    ).all()

    for pedido in pedidos_fechados_periodo:
        funcionario_produtividade = None
        if pedido.garcom and pedido.garcom.funcionario:
            funcionario_produtividade = pedido.garcom.funcionario
        elif pedido.caixa and pedido.caixa.funcionario:
            funcionario_produtividade = pedido.caixa.funcionario

        if not funcionario_produtividade:
            continue

        registro = produtividade_map.setdefault(
            funcionario_produtividade.id,
            {
                'nome': funcionario_produtividade.nome,
                'cargo': funcionario_produtividade.cargo or funcionario_produtividade.role.upper(),
                'role': funcionario_produtividade.role,
                'pedidos': 0,
                'itens': 0,
                'faturamento': 0.0,
            }
        )
        registro['pedidos'] += 1
        registro['faturamento'] += float(pedido.total or 0.0)
        registro['itens'] += sum((item.quantidade or 0) for item in (pedido.itens or []))

    produtividade_colaboradores = []
    for registro in produtividade_map.values():
        pedidos = float(registro['pedidos'] or 0)
        itens = float(registro['itens'] or 0)
        faturamento = float(registro['faturamento'] or 0.0)
        registro['ticket_medio'] = (faturamento / pedidos) if pedidos > 0 else 0.0
        registro['pedidos_dia'] = pedidos / 30.0
        registro['indice_produtividade'] = pedidos + (itens / 10.0)
        produtividade_colaboradores.append(registro)

    produtividade_colaboradores = sorted(
        produtividade_colaboradores,
        key=lambda item: (item['indice_produtividade'], item['faturamento']),
        reverse=True
    )
    top_produtividade = produtividade_colaboradores[0] if produtividade_colaboradores else None
    total_pedidos_produtividade = sum(item['pedidos'] for item in produtividade_colaboradores)
    media_pedidos_produtividade = (
        total_pedidos_produtividade / len(produtividade_colaboradores)
        if produtividade_colaboradores else 0.0
    )

    return render_template(
        'rh/indicadores.html',
        total_funcionarios=total_funcionarios,
        funcionarios_ativos=funcionarios_ativos,
        funcionarios_inativos=funcionarios_inativos,
        acessos_controlados=acessos_controlados,
        funcoes_total=funcoes_total,
        funcoes_ativas=funcoes_ativas,
        perfis_acesso_total=perfis_acesso_total,
        perfis_acesso_ativos=perfis_acesso_ativos,
        admissoes_30_dias=admissoes_30_dias,
        pedidos_pendentes=pedidos_pendentes,
        pedidos_hoje=pedidos_hoje,
        pedidos_30_dias=pedidos_30_dias,
        media_pedidos_dia_30=media_pedidos_dia_30,
        equipe_operacao_ativa=equipe_operacao_ativa,
        pendencias_por_colaborador=pendencias_por_colaborador,
        equipe_minima_recomendada=equipe_minima_recomendada,
        deficit_equipe=deficit_equipe,
        acoes_rh=acoes_rh,
        distribuicao_roles=distribuicao_roles,
        distribuicao_perfis_acesso=distribuicao_perfis_acesso,
        distribuicao_cargos=distribuicao_cargos,
        cargos_sem_cobertura=cargos_sem_cobertura,
        funcionarios_recentes=funcionarios_recentes,
        produtividade_colaboradores=produtividade_colaboradores,
        top_produtividade=top_produtividade,
        media_pedidos_produtividade=media_pedidos_produtividade,
    )


@app.route('/rh/organograma')
@login_required
def organograma_rh():
    funcionario_logado = get_funcionario_logado()
    todos_funcionarios = Funcionario.query.filter_by(ativo=True).order_by(Funcionario.nome.asc()).all()
    todos_map = {f.id: f for f in todos_funcionarios}
    filhos_map_total = {}
    for f in todos_funcionarios:
        if f.superior_id and f.superior_id in todos_map:
            filhos_map_total.setdefault(f.superior_id, []).append(f)

    def _ids_visiveis_para_funcionario(funcionario):
        if not funcionario:
            return set()
        role_norm = (funcionario.role or '').strip().lower()
        if role_norm in {'admin', 'gerente'}:
            return {f.id for f in todos_funcionarios}

        ids_visiveis = {funcionario.id}

        cursor = funcionario.superior
        visitados_superiores = set()
        while cursor and cursor.id not in visitados_superiores:
            visitados_superiores.add(cursor.id)
            ids_visiveis.add(cursor.id)
            cursor = cursor.superior

        fila = deque([funcionario.id])
        visitados_desc = set()
        while fila:
            atual_id = fila.popleft()
            if atual_id in visitados_desc:
                continue
            visitados_desc.add(atual_id)
            for filho in filhos_map_total.get(atual_id, []):
                ids_visiveis.add(filho.id)
                fila.append(filho.id)

        return ids_visiveis

    ids_visiveis = _ids_visiveis_para_funcionario(funcionario_logado)
    funcionarios_base = [f for f in todos_funcionarios if f.id in ids_visiveis] if ids_visiveis else todos_funcionarios

    departamentos_disponiveis = sorted({
        (f.departamento or '').strip()
        for f in funcionarios_base
        if (f.departamento or '').strip()
    }, key=str.lower)
    cargos_disponiveis = sorted({
        (f.cargo or '').strip()
        for f in funcionarios_base
        if (f.cargo or '').strip()
    }, key=str.lower)
    times_disponiveis = sorted({
        (f.time_nome or '').strip()
        for f in funcionarios_base
        if (f.time_nome or '').strip()
    }, key=str.lower)
    niveis_disponiveis = sorted({
        (f.nivel_organograma or '').strip()
        for f in funcionarios_base
        if (f.nivel_organograma or '').strip()
    }, key=str.lower)
    roles_disponiveis = sorted({
        (f.role or '').strip().lower()
        for f in funcionarios_base
        if (f.role or '').strip()
    })
    perfis_disponiveis = sorted({
        (f.perfil_acesso.nome or '').strip()
        for f in funcionarios_base
        if getattr(f, 'perfil_acesso', None) and (f.perfil_acesso.nome or '').strip()
    }, key=str.lower)

    departamento_filtro = (request.args.get('departamento') or '').strip()
    cargo_filtro = (request.args.get('cargo') or '').strip()
    time_filtro = (request.args.get('time') or '').strip()
    nivel_filtro = (request.args.get('nivel') or '').strip()
    role_filtro = (request.args.get('role') or '').strip().lower()
    perfil_filtro = (request.args.get('perfil') or '').strip()
    lideranca_filtro = (request.args.get('lideranca') or '').strip().lower()
    vinculo_filtro = (request.args.get('vinculo') or '').strip().lower()
    acesso_filtro = (request.args.get('acesso') or '').strip().lower()
    busca_filtro = (request.args.get('busca') or '').strip()

    funcionarios = funcionarios_base
    if departamento_filtro:
        funcionarios = [
            f for f in funcionarios
            if (f.departamento or '').strip() == departamento_filtro
        ]
    if cargo_filtro:
        funcionarios = [
            f for f in funcionarios
            if (f.cargo or '').strip() == cargo_filtro
        ]
    if time_filtro:
        funcionarios = [
            f for f in funcionarios
            if (f.time_nome or '').strip() == time_filtro
        ]
    if nivel_filtro:
        funcionarios = [
            f for f in funcionarios
            if (f.nivel_organograma or '').strip() == nivel_filtro
        ]
    if role_filtro:
        funcionarios = [
            f for f in funcionarios
            if (f.role or '').strip().lower() == role_filtro
        ]
    if perfil_filtro:
        funcionarios = [
            f for f in funcionarios
            if getattr(f, 'perfil_acesso', None) and (f.perfil_acesso.nome or '').strip() == perfil_filtro
        ]
    if lideranca_filtro == 'lideres':
        funcionarios = [f for f in funcionarios if filhos_map_total.get(f.id)]
    elif lideranca_filtro == 'sem_lideranca':
        funcionarios = [f for f in funcionarios if not filhos_map_total.get(f.id)]
    if vinculo_filtro == 'com_superior':
        funcionarios = [
            f for f in funcionarios
            if f.superior_id and f.superior_id in ids_visiveis
        ]
    elif vinculo_filtro == 'sem_superior':
        funcionarios = [
            f for f in funcionarios
            if not f.superior_id or f.superior_id not in ids_visiveis
        ]
    if acesso_filtro == 'controlado':
        funcionarios = [f for f in funcionarios if f.controle_acesso_ativo]
    elif acesso_filtro == 'livre':
        funcionarios = [f for f in funcionarios if not f.controle_acesso_ativo]
    if busca_filtro:
        busca_normalizada = busca_filtro.lower()
        funcionarios = [
            f for f in funcionarios
            if busca_normalizada in (f.nome or '').lower()
            or busca_normalizada in (f.cargo or '').lower()
            or busca_normalizada in (f.departamento or '').lower()
            or busca_normalizada in (f.time_nome or '').lower()
            or busca_normalizada in (f.role or '').lower()
            or busca_normalizada in ((f.perfil_acesso.nome if getattr(f, 'perfil_acesso', None) else '') or '').lower()
        ]

    funcionarios_map = {f.id: f for f in funcionarios}

    role_ordem = {
        'admin': 0,
        'gerente': 1,
        'caixa': 2,
        'operador': 3,
        'garcom': 4,
    }

    def _sort_key(func):
        return (
            role_ordem.get((func.role or '').lower(), 99),
            (func.nivel_organograma or '').lower(),
            (func.cargo or '').lower(),
            (func.nome or '').lower(),
        )

    arvore = {}
    raizes = []
    for funcionario in funcionarios:
        if funcionario.superior_id and funcionario.superior_id in funcionarios_map:
            arvore.setdefault(funcionario.superior_id, []).append(funcionario)
        else:
            raizes.append(funcionario)

    for chave in list(arvore.keys()):
        arvore[chave] = sorted(arvore[chave], key=_sort_key)
    raizes = sorted(raizes, key=_sort_key)

    subordinados_totais = {}

    def _contar_subordinados(funcionario_id, visitados=None):
        if visitados is None:
            visitados = set()
        if funcionario_id in visitados:
            return 0
        visitados.add(funcionario_id)
        total = 0
        for filho in arvore.get(funcionario_id, []):
            total += 1 + _contar_subordinados(filho.id, visitados.copy())
        subordinados_totais[funcionario_id] = total
        return total

    for raiz in raizes:
        _contar_subordinados(raiz.id)
    for funcionario in funcionarios:
        subordinados_totais.setdefault(funcionario.id, len(arvore.get(funcionario.id, [])))

    camadas = {}
    fila = deque((raiz, 0) for raiz in raizes)
    visitados = set()
    while fila:
        funcionario, camada = fila.popleft()
        if funcionario.id in visitados:
            continue
        visitados.add(funcionario.id)
        camadas.setdefault(camada, []).append(funcionario)
        for filho in arvore.get(funcionario.id, []):
            fila.append((filho, camada + 1))

    for funcionario in funcionarios:
        if funcionario.id not in visitados:
            camadas.setdefault(0, []).append(funcionario)

    camadas_organograma = []
    for indice, colaboradores in sorted(camadas.items(), key=lambda item: item[0]):
        colaboradores = sorted(colaboradores, key=_sort_key)
        contagem_niveis = {}
        for colaborador in colaboradores:
            nivel = (colaborador.nivel_organograma or '').strip()
            if not nivel:
                continue
            contagem_niveis[nivel] = contagem_niveis.get(nivel, 0) + 1
        rotulo_camada = max(contagem_niveis, key=contagem_niveis.get) if contagem_niveis else f'Camada {indice + 1}'
        camadas_organograma.append({
            'indice': indice + 1,
            'rotulo': rotulo_camada,
            'colaboradores': colaboradores,
        })

    lideres = [f for f in funcionarios if f.id in arvore]
    sem_superior = [f for f in funcionarios if not f.superior_id or f.superior_id not in funcionarios_map]

    resumo_departamentos_map = {}
    for funcionario in funcionarios:
        departamento_nome = (funcionario.departamento or '').strip() or 'Sem departamento'
        registro = resumo_departamentos_map.setdefault(departamento_nome, {'nome': departamento_nome, 'total': 0, 'lideres': 0})
        registro['total'] += 1
        if funcionario.id in arvore:
            registro['lideres'] += 1

    resumo_departamentos = sorted(
        resumo_departamentos_map.values(),
        key=lambda item: (-item['total'], item['nome'].lower())
    )

    return render_template(
        'rh/organograma.html',
        raizes=raizes,
        arvore=arvore,
        camadas_organograma=camadas_organograma,
        subordinados_totais=subordinados_totais,
        departamentos_disponiveis=departamentos_disponiveis,
        cargos_disponiveis=cargos_disponiveis,
        times_disponiveis=times_disponiveis,
        niveis_disponiveis=niveis_disponiveis,
        roles_disponiveis=roles_disponiveis,
        perfis_disponiveis=perfis_disponiveis,
        departamento_filtro=departamento_filtro,
        cargo_filtro=cargo_filtro,
        time_filtro=time_filtro,
        nivel_filtro=nivel_filtro,
        role_filtro=role_filtro,
        perfil_filtro=perfil_filtro,
        lideranca_filtro=lideranca_filtro,
        vinculo_filtro=vinculo_filtro,
        acesso_filtro=acesso_filtro,
        busca_filtro=busca_filtro,
        profundidade_maxima=(max(camadas.keys()) + 1 if camadas else 0),
        resumo_departamentos=resumo_departamentos,
        total_colaboradores=len(funcionarios),
        total_lideres=len(lideres),
        total_sem_superior=len(sem_superior),
        modo_visao=('completo' if funcionario_logado and funcionario_logado.role in {'admin', 'gerente'} else 'restrito'),
        funcionario_logado=funcionario_logado,
    )


@app.route('/api/rh/analytics')
@require_role('admin', 'gerente')
def analytics_rh_api():
    periodo = request.args.get('periodo', type=int) or 30
    if periodo not in {30, 90, 365}:
        periodo = 30
    cache = extensions.cache
    cache_key = f'analytics:rh:{periodo}'
    if cache is not None:
        cached_payload = cache.get(cache_key)
        if cached_payload is not None:
            return jsonify(cached_payload)

    agora = datetime.utcnow()
    data_limite = agora - timedelta(days=periodo)
    admissoes_periodo = Funcionario.query.filter(Funcionario.criado_em >= data_limite).count()
    total_funcionarios = Funcionario.query.count()
    funcionarios_ativos = Funcionario.query.filter(Funcionario.ativo.is_(True)).count()
    funcionarios_inativos = Funcionario.query.filter(Funcionario.ativo.is_(False)).count()
    acessos_controlados = Funcionario.query.filter(
        Funcionario.controle_acesso_ativo.is_(True)
    ).count()

    distribuicao_roles = db.session.query(
        Funcionario.role.label('role'),
        db.func.count(Funcionario.id).label('quantidade')
    ).group_by(Funcionario.role).order_by(db.desc('quantidade')).all()
    distribuicao_cargos = db.session.query(
        Funcionario.cargo.label('cargo'),
        db.func.count(Funcionario.id).label('quantidade'),
        db.func.sum(db.case((Funcionario.ativo.is_(True), 1), else_=0)).label('ativos')
    ).group_by(Funcionario.cargo).order_by(db.desc('quantidade')).all()
    distribuicao_perfis = db.session.query(
        PerfilAcesso.nome.label('perfil'),
        db.func.count(Funcionario.id).label('quantidade')
    ).outerjoin(
        Funcionario, Funcionario.perfil_acesso_id == PerfilAcesso.id
    ).group_by(
        PerfilAcesso.id, PerfilAcesso.nome
    ).order_by(
        db.desc('quantidade'), PerfilAcesso.nome.asc()
    ).all()

    pedidos_pendentes = Pedido.query.filter(
        Pedido.status.in_([Pedido.STATUS_ABERTO, Pedido.STATUS_EM_PREPARO, Pedido.STATUS_ENTREGUE])
    ).count()
    equipe_operacao_ativa = Funcionario.query.filter(
        Funcionario.ativo.is_(True),
        Funcionario.role.in_(['caixa', 'operador', 'garcom', 'gerente'])
    ).count()
    pendencias_por_colaborador = (
        pedidos_pendentes / equipe_operacao_ativa
        if equipe_operacao_ativa > 0 else float(pedidos_pendentes)
    )
    deficit_equipe = max(0, max(1, int((pedidos_pendentes + 7) // 8)) - equipe_operacao_ativa) if pedidos_pendentes else 0

    admissoes_por_mes_raw = db.session.query(
        db.func.strftime('%Y-%m', Funcionario.criado_em).label('mes'),
        db.func.count(Funcionario.id).label('quantidade')
    ).filter(
        Funcionario.criado_em >= agora - timedelta(days=180)
    ).group_by(
        db.func.strftime('%Y-%m', Funcionario.criado_em)
    ).order_by(
        db.func.strftime('%Y-%m', Funcionario.criado_em).asc()
    ).all()

    admissoes_diarias = [
        {'dia': item.mes, 'quantidade': int(item.quantidade or 0)}
        for item in admissoes_por_mes_raw
    ]

    funcoes_ativas_lista = FuncaoRH.query.filter_by(ativo=True).all()
    cargos_sem_cobertura = []
    for funcao in funcoes_ativas_lista:
        ativos_no_cargo = Funcionario.query.filter(
            Funcionario.ativo.is_(True),
            db.func.lower(Funcionario.cargo) == (funcao.nome or '').lower()
        ).count()
        if ativos_no_cargo == 0:
            cargos_sem_cobertura.append({
                'cargo': funcao.nome,
                'ativos': 0,
            })

    funcionarios_recentes = Funcionario.query.order_by(Funcionario.criado_em.desc()).limit(5).all()

    payload = {
        'success': True,
        'message': 'Analytics RH carregado com sucesso.',
        'data': {
            'periodo_dias': periodo,
            'kpis': {
                'total_funcionarios': int(total_funcionarios or 0),
                'funcionarios_ativos': int(funcionarios_ativos or 0),
                'funcionarios_inativos': int(funcionarios_inativos or 0),
                'acessos_controlados': int(acessos_controlados or 0),
                'produtividade_media': 0.0,
                'equipe_operacional_ativa': int(equipe_operacao_ativa or 0),
                'deficit_equipe': int(deficit_equipe or 0),
            },
            'admissoes_periodo': admissoes_periodo,
            'distribuicao_roles': [
                {'role': item.role, 'quantidade': int(item.quantidade or 0)}
                for item in distribuicao_roles
            ],
            'distribuicao_cargos': [
                {'cargo': item.cargo or '-', 'quantidade': int(item.quantidade or 0), 'ativos': int(item.ativos or 0)}
                for item in distribuicao_cargos
            ],
            'distribuicao_perfis_acesso': [
                {'perfil': item.perfil or 'Sem perfil padrao', 'quantidade': int(item.quantidade or 0)}
                for item in distribuicao_perfis
            ],
            'admissoes_diarias': admissoes_diarias,
            'pedidos_pendentes': pedidos_pendentes,
            'equipe_operacao_ativa': equipe_operacao_ativa,
            'pendencias_por_colaborador': round(pendencias_por_colaborador, 2),
            'cargos_sem_cobertura': cargos_sem_cobertura,
            'funcionarios_recentes': [
                {
                    'nome': item.nome,
                    'cargo': item.cargo or item.role.upper(),
                    'data_admissao': item.criado_em.strftime('%d/%m/%Y') if item.criado_em else '-',
                }
                for item in funcionarios_recentes
            ],
            'alertas': [],
            'produtividade_vs_faturamento': [],
        }
    }

    produtividade_map = {}
    pedidos_fechados_periodo = Pedido.query.options(
        selectinload(Pedido.garcom).selectinload(Garcom.funcionario),
        selectinload(Pedido.caixa).selectinload(Caixa.funcionario),
    ).filter(
        Pedido.status == Pedido.STATUS_FECHADO,
        Pedido.fechado_em.isnot(None),
        Pedido.fechado_em >= data_limite
    ).all()
    for pedido in pedidos_fechados_periodo:
        funcionario_produtividade = None
        if pedido.garcom and pedido.garcom.funcionario:
            funcionario_produtividade = pedido.garcom.funcionario
        elif pedido.caixa and pedido.caixa.funcionario:
            funcionario_produtividade = pedido.caixa.funcionario
        if not funcionario_produtividade:
            continue
        registro = produtividade_map.setdefault(
            funcionario_produtividade.id,
            {
                'nome': funcionario_produtividade.nome,
                'cargo': funcionario_produtividade.cargo or funcionario_produtividade.role.upper(),
                'pedidos': 0,
                'faturamento': 0.0,
            }
        )
        registro['pedidos'] += 1
        registro['faturamento'] += float(pedido.total or 0.0)

    produtividade_items = [
        {
            'nome': item['nome'],
            'cargo': item['cargo'],
            'pedidos': int(item['pedidos'] or 0),
            'faturamento': float(item['faturamento'] or 0.0),
            'ticket_medio': float((item['faturamento'] or 0.0) / item['pedidos']) if item['pedidos'] else 0.0,
        }
        for item in produtividade_map.values()
    ]
    produtividade_items.sort(key=lambda item: (item['pedidos'], item['faturamento']), reverse=True)
    payload['data']['produtividade_vs_faturamento'] = produtividade_items
    payload['data']['top_produtividade'] = produtividade_items[:5]
    payload['data']['kpis']['produtividade_media'] = round(
        (
            sum(item['pedidos'] for item in produtividade_items) / len(produtividade_items)
            if produtividade_items else 0.0
        ),
        2,
    )
    if cargos_sem_cobertura:
        payload['data']['alertas'].append({
            'nivel': 'warning',
            'titulo': 'Cargos sem cobertura',
            'descricao': f'{len(cargos_sem_cobertura)} cargo(s) ativo(s) sem colaborador ativo.',
        })
    if deficit_equipe > 0:
        payload['data']['alertas'].append({
            'nivel': 'danger',
            'titulo': 'Deficit operacional',
            'descricao': f'Deficit estimado de {deficit_equipe} colaborador(es) para a fila atual.',
        })
    perfis_sem_usuario = [item for item in payload['data']['distribuicao_perfis_acesso'] if item['quantidade'] == 0]
    if perfis_sem_usuario:
        payload['data']['alertas'].append({
            'nivel': 'info',
            'titulo': 'Perfis sem aplicacao',
            'descricao': f'{len(perfis_sem_usuario)} perfil(is) ativo(s) sem usuarios vinculados.',
        })
    if cache is not None:
        cache.set(cache_key, payload, timeout=60)
    return jsonify(payload)


@app.route('/auditoria')
@require_role('admin', 'gerente')
def auditoria_sistema():
    funcionario_id = request.args.get('funcionario_id', type=int)
    acao = (request.args.get('acao') or '').strip()
    entidade = (request.args.get('entidade') or '').strip()
    metodo = (request.args.get('metodo') or '').strip().upper()
    page = max(request.args.get('page', 1, type=int), 1)
    per_page = min(max(request.args.get('per_page', 30, type=int), 10), 100)

    query = AuditoriaEvento.query
    if funcionario_id:
        query = query.filter(AuditoriaEvento.funcionario_id == funcionario_id)
    if acao:
        query = query.filter(AuditoriaEvento.acao.ilike(f'%{acao}%'))
    if entidade:
        query = query.filter(AuditoriaEvento.entidade.ilike(f'%{entidade}%'))
    if metodo in {'GET', 'POST', 'PUT', 'PATCH', 'DELETE'}:
        query = query.filter(AuditoriaEvento.metodo == metodo)

    paginacao = query.order_by(AuditoriaEvento.criado_em.desc()).paginate(page=page, per_page=per_page, error_out=False)
    eventos = paginacao.items
    funcionarios = Funcionario.query.order_by(Funcionario.nome.asc()).all()
    return render_template(
        'sistema/auditoria.html',
        eventos=eventos,
        paginacao=paginacao,
        funcionarios=funcionarios,
        filtros={
            'funcionario_id': funcionario_id,
            'acao': acao,
            'entidade': entidade,
            'metodo': metodo,
            'per_page': per_page,
        }
    )


@app.route('/rh/funcoes/nova', methods=['POST'])
@require_role('admin', 'gerente')
def nova_funcao_rh():
    nome = (request.form.get('nome') or '').strip()
    descricao = (request.form.get('descricao') or '').strip() or None
    ativo = request.form.get('ativo') == 'on'

    if not nome:
        flash('Nome do cargo e obrigatorio.', 'danger')
        return redirect(url_for('listar_funcoes_rh'))

    existente = FuncaoRH.query.filter(db.func.lower(FuncaoRH.nome) == nome.lower()).first()
    if existente:
        flash('Ja existe um cargo com esse nome.', 'warning')
        return redirect(url_for('listar_funcoes_rh'))

    try:
        funcao = FuncaoRH(nome=nome, descricao=descricao, ativo=ativo)
        db.session.add(funcao)
        db.session.commit()
        flash(f'Cargo "{nome}" criado com sucesso!', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao criar cargo: {str(e)}', 'danger')

    return redirect(url_for('listar_funcoes_rh'))


@app.route('/rh/funcoes/<int:funcao_id>/editar', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def editar_funcao_rh(funcao_id):
    funcao = FuncaoRH.query.get_or_404(funcao_id)
    nome_anterior = funcao.nome

    if request.method == 'POST':
        nome = (request.form.get('nome') or '').strip()
        descricao = (request.form.get('descricao') or '').strip() or None
        ativo = request.form.get('ativo') == 'on'

        if not nome:
            flash('Nome do cargo e obrigatorio.', 'danger')
            return redirect(url_for('editar_funcao_rh', funcao_id=funcao_id))

        existente = FuncaoRH.query.filter(
            db.func.lower(FuncaoRH.nome) == nome.lower(),
            FuncaoRH.id != funcao.id
        ).first()
        if existente:
            flash('Ja existe outro cargo com esse nome.', 'warning')
            return redirect(url_for('editar_funcao_rh', funcao_id=funcao_id))

        try:
            funcao.nome = nome
            funcao.descricao = descricao
            funcao.ativo = ativo
            if (nome_anterior or '').strip().lower() != nome.lower():
                Funcionario.query.filter(
                    db.func.lower(Funcionario.cargo) == (nome_anterior or '').strip().lower()
                ).update({Funcionario.cargo: nome}, synchronize_session=False)
            db.session.commit()
            flash('Cargo atualizado com sucesso!', 'success')
            return redirect(url_for('listar_funcoes_rh'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar cargo: {str(e)}', 'danger')

    return render_template(
        'rh/editar_funcao.html',
        funcao=funcao,
    )


@app.route('/rh/funcoes/<int:funcao_id>/deletar', methods=['POST'])
@require_role('admin', 'gerente')
def deletar_funcao_rh(funcao_id):
    funcao = FuncaoRH.query.get_or_404(funcao_id)
    vinculado = Funcionario.query.filter(
        db.func.lower(Funcionario.cargo) == (funcao.nome or '').strip().lower()
    ).first()
    if vinculado:
        flash('Este cargo ainda está vinculado a colaboradores. Realoque ou edite os cadastros antes de excluir.', 'warning')
        return redirect(url_for('listar_funcoes_rh'))
    try:
        db.session.delete(funcao)
        db.session.commit()
        flash('Cargo removido com sucesso.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao remover cargo: {str(e)}', 'danger')
    return redirect(url_for('listar_funcoes_rh'))


@app.route('/rh/perfis-acesso/novo', methods=['POST'])
@require_role('admin', 'gerente')
def novo_perfil_acesso_rh():
    nome = (request.form.get('nome') or '').strip()
    descricao = (request.form.get('descricao') or '').strip() or None
    ativo = request.form.get('ativo') == 'on'
    permissoes_padrao = _extrair_permissoes_padrao_form()

    if not nome:
        flash('Nome do perfil de acesso é obrigatório.', 'danger')
        return redirect(url_for('listar_perfis_rh'))

    existente = PerfilAcesso.query.filter(db.func.lower(PerfilAcesso.nome) == nome.lower()).first()
    if existente:
        flash('Já existe um perfil de acesso com esse nome.', 'warning')
        return redirect(url_for('listar_perfis_rh'))

    try:
        perfil = PerfilAcesso(
            nome=nome,
            descricao=descricao,
            ativo=ativo,
            permissoes_padrao=_serializar_paginas_json(permissoes_padrao),
        )
        db.session.add(perfil)
        db.session.commit()
        flash(f'Perfil de acesso "{nome}" criado com sucesso!', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao criar perfil de acesso: {str(e)}', 'danger')

    return redirect(url_for('listar_perfis_rh'))


@app.route('/rh/perfis-acesso/<int:perfil_id>/editar', methods=['GET', 'POST'])
@require_role('admin', 'gerente')
def editar_perfil_acesso_rh(perfil_id):
    perfil = PerfilAcesso.query.get_or_404(perfil_id)
    permissoes_atuais = _paginas_perfil_acesso(perfil)

    if request.method == 'POST':
        nome = (request.form.get('nome') or '').strip()
        descricao = (request.form.get('descricao') or '').strip() or None
        ativo = request.form.get('ativo') == 'on'
        permissoes_padrao = _extrair_permissoes_padrao_form()

        if not nome:
            flash('Nome do perfil de acesso é obrigatório.', 'danger')
            return redirect(url_for('editar_perfil_acesso_rh', perfil_id=perfil_id))

        existente = PerfilAcesso.query.filter(
            db.func.lower(PerfilAcesso.nome) == nome.lower(),
            PerfilAcesso.id != perfil.id
        ).first()
        if existente:
            flash('Já existe outro perfil de acesso com esse nome.', 'warning')
            return redirect(url_for('editar_perfil_acesso_rh', perfil_id=perfil_id))

        try:
            perfil.nome = nome
            perfil.descricao = descricao
            perfil.ativo = ativo
            perfil.permissoes_padrao = _serializar_paginas_json(permissoes_padrao)
            db.session.commit()
            flash('Perfil de acesso atualizado com sucesso!', 'success')
            return redirect(url_for('listar_perfis_rh'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar perfil de acesso: {str(e)}', 'danger')

    return render_template(
        'rh/editar_perfil_acesso.html',
        perfil=perfil,
        permissoes_atuais=permissoes_atuais,
        paginas_ordenadas_menu=_paginas_ordenadas_menu(),
        paginas_sistema=PAGINAS_SISTEMA,
    )


@app.route('/rh/perfis-acesso/<int:perfil_id>/deletar', methods=['POST'])
@require_role('admin', 'gerente')
def deletar_perfil_acesso_rh(perfil_id):
    perfil = PerfilAcesso.query.get_or_404(perfil_id)
    colaborador_vinculado = Funcionario.query.filter_by(perfil_acesso_id=perfil.id).first()
    if colaborador_vinculado:
        flash('Este perfil de acesso está vinculado a colaboradores. Remova os vínculos antes de excluir.', 'warning')
        return redirect(url_for('listar_perfis_rh'))
    try:
        db.session.delete(perfil)
        db.session.commit()
        flash('Perfil de acesso removido com sucesso.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao remover perfil de acesso: {str(e)}', 'danger')
    return redirect(url_for('listar_perfis_rh'))


# ============ ROTAS - SERVICOS TECNICOS ============

def _servicos_tecnicos_empresa_ativos():
    empresa = EmpresaConfig.query.first()
    return bool(empresa and empresa.servicos_tecnicos_ativos)


def _funcionarios_tecnicos_montadores():
    funcionarios = Funcionario.query.filter_by(ativo=True).order_by(Funcionario.nome.asc()).all()
    elegiveis = []
    for func in funcionarios:
        cargo_norm = (func.cargo or '').strip().lower()
        role_norm = (func.role or '').strip().lower()
        if role_norm in {'admin', 'gerente'}:
            continue
        if 'tecn' in cargo_norm or 'montador' in cargo_norm or 'instal' in cargo_norm:
            elegiveis.append(func)
    return elegiveis or funcionarios


def _pode_gerir_chamados(funcionario):
    role_norm = ((funcionario.role if funcionario else '') or '').strip().lower()
    return role_norm in {'admin', 'gerente'}


@app.route('/chamados')
@login_required
def listar_chamados_internos():
    funcionario = get_funcionario_logado()
    if not funcionario:
        flash('Sessao invalida. Faca login novamente.', 'danger')
        return redirect(url_for('login'))

    status = (request.args.get('status') or '').strip().lower()
    categoria = (request.args.get('categoria') or '').strip().lower()
    busca = (request.args.get('busca') or '').strip()

    query = ChamadoInterno.query

    if not _pode_gerir_chamados(funcionario):
        query = query.filter(
            db.or_(
                ChamadoInterno.solicitante_id == funcionario.id,
                ChamadoInterno.responsavel_id == funcionario.id,
            )
        )

    if status in ChamadoInterno.STATUS_VALIDOS:
        query = query.filter(ChamadoInterno.status == status)
    if categoria in ChamadoInterno.CATEGORIAS_VALIDAS:
        query = query.filter(ChamadoInterno.categoria == categoria)
    if busca:
        termo = f'%{busca}%'
        query = query.filter(
            db.or_(
                ChamadoInterno.titulo.ilike(termo),
                ChamadoInterno.descricao.ilike(termo),
                db.cast(ChamadoInterno.id, db.String).ilike(termo),
                ChamadoInterno.solicitante.has(Funcionario.nome.ilike(termo)),
                ChamadoInterno.responsavel.has(Funcionario.nome.ilike(termo)),
            )
        )

    chamados = query.order_by(
        db.case((ChamadoInterno.status.in_([ChamadoInterno.STATUS_ABERTO, ChamadoInterno.STATUS_TRIAGEM]), 0), else_=1),
        db.case((ChamadoInterno.prioridade == 'critica', 0), (ChamadoInterno.prioridade == 'alta', 1), (ChamadoInterno.prioridade == 'media', 2), else_=3),
        ChamadoInterno.aberto_em.desc(),
    ).all()

    return render_template(
        'servicos/chamados/listar.html',
        chamados=chamados,
        filtros={'status': status, 'categoria': categoria, 'busca': busca},
        categorias_validas=ChamadoInterno.CATEGORIAS_VALIDAS,
        status_validos=ChamadoInterno.STATUS_VALIDOS,
        prioridades_validas=ChamadoInterno.PRIORIDADES_VALIDAS,
        pode_gerir=_pode_gerir_chamados(funcionario),
    )


@app.route('/chamados/novo', methods=['GET', 'POST'])
@login_required
def criar_chamado_interno():
    funcionario = get_funcionario_logado()
    if not funcionario:
        flash('Sessao invalida. Faca login novamente.', 'danger')
        return redirect(url_for('login'))

    responsaveis = Funcionario.query.filter_by(ativo=True).order_by(Funcionario.nome.asc()).all() if _pode_gerir_chamados(funcionario) else []

    if request.method == 'POST':
        try:
            titulo = (request.form.get('titulo') or '').strip()
            categoria = (request.form.get('categoria') or ChamadoInterno.CATEGORIA_SISTEMA).strip().lower()
            prioridade = (request.form.get('prioridade') or 'media').strip().lower()
            descricao = (request.form.get('descricao') or '').strip()
            setor_origem = (request.form.get('setor_origem') or funcionario.departamento or '').strip() or None
            responsavel_id = request.form.get('responsavel_id', type=int) or None
            status = (request.form.get('status') or ChamadoInterno.STATUS_ABERTO).strip().lower()
            resolucao = (request.form.get('resolucao') or '').strip() or None

            if not titulo or not descricao:
                flash('Titulo e descricao do chamado sao obrigatorios.', 'danger')
                return redirect(url_for('criar_chamado_interno'))
            if categoria not in ChamadoInterno.CATEGORIAS_VALIDAS:
                categoria = ChamadoInterno.CATEGORIA_SISTEMA
            if prioridade not in ChamadoInterno.PRIORIDADES_VALIDAS:
                prioridade = 'media'
            if status not in ChamadoInterno.STATUS_VALIDOS or not _pode_gerir_chamados(funcionario):
                status = ChamadoInterno.STATUS_ABERTO

            responsavel = Funcionario.query.get(responsavel_id) if responsavel_id else None
            if responsavel_id and not responsavel:
                flash('Responsavel invalido.', 'danger')
                return redirect(url_for('criar_chamado_interno'))

            chamado = ChamadoInterno(
                titulo=titulo,
                categoria=categoria,
                prioridade=prioridade,
                status=status,
                descricao=descricao,
                setor_origem=setor_origem,
                solicitante_id=funcionario.id,
                responsavel_id=(responsavel.id if responsavel else None),
                resolucao=resolucao,
            )
            if chamado.status == ChamadoInterno.STATUS_CONCLUIDO:
                chamado.concluido_em = datetime.utcnow()
            db.session.add(chamado)
            db.session.commit()
            flash(f'Chamado #{chamado.id} aberto com sucesso.', 'success')
            return redirect(url_for('listar_chamados_internos'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao abrir chamado: {str(e)}', 'danger')

    return render_template(
        'servicos/chamados/form.html',
        chamado=None,
        categorias_validas=ChamadoInterno.CATEGORIAS_VALIDAS,
        status_validos=ChamadoInterno.STATUS_VALIDOS,
        prioridades_validas=ChamadoInterno.PRIORIDADES_VALIDAS,
        responsaveis=responsaveis,
        pode_gerir=_pode_gerir_chamados(funcionario),
        funcionario_logado=funcionario,
    )


@app.route('/chamados/<int:chamado_id>/editar', methods=['GET', 'POST'])
@login_required
def editar_chamado_interno(chamado_id):
    funcionario = get_funcionario_logado()
    chamado = ChamadoInterno.query.get_or_404(chamado_id)
    pode_gerir = _pode_gerir_chamados(funcionario)

    if not pode_gerir and chamado.solicitante_id != funcionario.id and chamado.responsavel_id != funcionario.id:
        flash('Voce nao tem permissao para acessar este chamado.', 'danger')
        return redirect(url_for('listar_chamados_internos'))

    responsaveis = Funcionario.query.filter_by(ativo=True).order_by(Funcionario.nome.asc()).all() if pode_gerir else []

    if request.method == 'POST':
        try:
            chamado.titulo = (request.form.get('titulo') or chamado.titulo).strip() or chamado.titulo
            categoria = (request.form.get('categoria') or chamado.categoria).strip().lower()
            prioridade = (request.form.get('prioridade') or chamado.prioridade).strip().lower()
            descricao = (request.form.get('descricao') or chamado.descricao).strip() or chamado.descricao
            chamado.setor_origem = (request.form.get('setor_origem') or chamado.setor_origem or '').strip() or None

            if categoria in ChamadoInterno.CATEGORIAS_VALIDAS:
                chamado.categoria = categoria
            if prioridade in ChamadoInterno.PRIORIDADES_VALIDAS:
                chamado.prioridade = prioridade
            chamado.descricao = descricao

            if pode_gerir:
                status = (request.form.get('status') or chamado.status).strip().lower()
                if status in ChamadoInterno.STATUS_VALIDOS:
                    chamado.status = status
                chamado.responsavel_id = request.form.get('responsavel_id', type=int) or None
                chamado.resolucao = (request.form.get('resolucao') or '').strip() or None
                if chamado.status == ChamadoInterno.STATUS_CONCLUIDO:
                    chamado.concluido_em = chamado.concluido_em or datetime.utcnow()
                elif chamado.status != ChamadoInterno.STATUS_CONCLUIDO:
                    chamado.concluido_em = None

            db.session.commit()
            flash(f'Chamado #{chamado.id} atualizado com sucesso.', 'success')
            return redirect(url_for('listar_chamados_internos'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar chamado: {str(e)}', 'danger')

    return render_template(
        'servicos/chamados/form.html',
        chamado=chamado,
        categorias_validas=ChamadoInterno.CATEGORIAS_VALIDAS,
        status_validos=ChamadoInterno.STATUS_VALIDOS,
        prioridades_validas=ChamadoInterno.PRIORIDADES_VALIDAS,
        responsaveis=responsaveis,
        pode_gerir=pode_gerir,
        funcionario_logado=funcionario,
    )


@app.route('/chamados/<int:chamado_id>/status', methods=['POST'])
@login_required
def atualizar_status_chamado_interno(chamado_id):
    funcionario = get_funcionario_logado()
    chamado = ChamadoInterno.query.get_or_404(chamado_id)
    pode_gerir = _pode_gerir_chamados(funcionario)

    if not pode_gerir and chamado.responsavel_id != funcionario.id:
        flash('Voce nao tem permissao para alterar o status deste chamado.', 'danger')
        return redirect(url_for('listar_chamados_internos'))

    novo_status = (request.form.get('status') or '').strip().lower()
    if novo_status not in ChamadoInterno.STATUS_VALIDOS:
        flash('Status invalido.', 'warning')
        return redirect(url_for('listar_chamados_internos'))

    try:
        chamado.status = novo_status
        if not chamado.responsavel_id and funcionario:
            chamado.responsavel_id = funcionario.id
        if novo_status == ChamadoInterno.STATUS_CONCLUIDO:
            chamado.concluido_em = datetime.utcnow()
            chamado.resolucao = (request.form.get('resolucao') or chamado.resolucao or '').strip() or chamado.resolucao
        elif novo_status != ChamadoInterno.STATUS_CONCLUIDO:
            chamado.concluido_em = None
        db.session.commit()
        flash(f'Status do chamado #{chamado.id} atualizado para {novo_status}.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao atualizar status do chamado: {str(e)}', 'danger')
    return redirect(url_for('listar_chamados_internos'))


@app.route('/servicos')
@require_role('admin', 'gerente', 'operador', 'caixa')
def listar_ordens_servico():
    if not _servicos_tecnicos_empresa_ativos():
        flash('Modulo de servicos tecnicos desativado na empresa.', 'warning')
        return redirect(url_for('dashboard'))

    status = (request.args.get('status') or '').strip().lower()
    tipo = (request.args.get('tipo') or '').strip().lower()
    busca = (request.args.get('busca') or '').strip()

    query = OrdemServico.query
    if status in OrdemServico.STATUS_VALIDOS:
        query = query.filter(OrdemServico.status == status)
    if tipo in OrdemServico.TIPOS_VALIDOS:
        query = query.filter(OrdemServico.tipo == tipo)
    if busca:
        termo = f'%{busca}%'
        query = query.filter(
            db.or_(
                OrdemServico.titulo.ilike(termo),
                OrdemServico.descricao.ilike(termo),
                db.cast(OrdemServico.id, db.String).ilike(termo),
                db.cast(OrdemServico.pedido_id, db.String).ilike(termo),
                OrdemServico.funcionario_destino.has(Funcionario.nome.ilike(termo)),
                OrdemServico.produto.has(Produto.nome.ilike(termo)),
            )
        )

    ordens = query.order_by(
        db.case((OrdemServico.status == OrdemServico.STATUS_ABERTA, 0), else_=1),
        OrdemServico.criado_em.desc()
    ).all()

    return render_template(
        'servicos/ordens/listar.html',
        ordens=ordens,
        filtros={'status': status, 'tipo': tipo, 'busca': busca},
        tipos_validos=OrdemServico.TIPOS_VALIDOS,
        status_validos=OrdemServico.STATUS_VALIDOS,
    )


@app.route('/servicos/minhas-ordens')
@login_required
def minhas_ordens_servico():
    funcionario = get_funcionario_logado()
    if not funcionario:
        flash('Sessao invalida. Faca login novamente.', 'danger')
        return redirect(url_for('login'))

    ordens = OrdemServico.query.filter_by(funcionario_destino_id=funcionario.id).order_by(
        db.case((OrdemServico.status.in_([OrdemServico.STATUS_ENVIADA, OrdemServico.STATUS_EM_EXECUCAO]), 0), else_=1),
        OrdemServico.data_agendada.asc().nullsfirst(),
        OrdemServico.criado_em.desc(),
    ).all()

    return render_template(
        'servicos/ordens/minhas.html',
        ordens=ordens,
        funcionario=funcionario,
    )


@app.route('/servicos/nova', methods=['GET', 'POST'])
@require_role('admin', 'gerente', 'operador', 'caixa')
def criar_ordem_servico():
    if not _servicos_tecnicos_empresa_ativos():
        flash('Modulo de servicos tecnicos desativado na empresa.', 'warning')
        return redirect(url_for('dashboard'))

    empresa = EmpresaConfig.query.first()
    produtos = Produto.query.filter_by(ativo=True).order_by(Produto.nome.asc()).all()
    pedidos = Pedido.query.order_by(Pedido.criado_em.desc()).limit(300).all()
    funcionarios_destino = _funcionarios_tecnicos_montadores()
    funcionario_logado = get_funcionario_logado()

    if request.method == 'POST':
        try:
            titulo = (request.form.get('titulo') or '').strip()
            tipo = (request.form.get('tipo') or '').strip().lower()
            status = (request.form.get('status') or OrdemServico.STATUS_ABERTA).strip().lower()
            servico_tipo = (request.form.get('servico_tipo') or OrdemServico.SERVICO_NENHUM).strip().lower()
            prioridade = (request.form.get('prioridade') or '').strip() or None
            descricao = (request.form.get('descricao') or '').strip() or None
            observacoes = (request.form.get('observacoes') or '').strip() or None
            avaria_detalhes = (request.form.get('avaria_detalhes') or '').strip() or None
            inspecao_detalhes = (request.form.get('inspecao_detalhes') or '').strip() or None
            resultado_inspecao = (request.form.get('resultado_inspecao') or '').strip() or None
            produto_id = request.form.get('produto_id', type=int) or None
            pedido_id = request.form.get('pedido_id', type=int) or None
            funcionario_destino_id = request.form.get('funcionario_destino_id', type=int) or None
            data_agendada = _parse_date_iso(request.form.get('data_agendada'))

            if not titulo:
                flash('Titulo da ordem e obrigatorio.', 'danger')
                return redirect(url_for('criar_ordem_servico'))
            if tipo not in OrdemServico.TIPOS_VALIDOS:
                flash('Tipo de ordem invalido.', 'danger')
                return redirect(url_for('criar_ordem_servico'))
            if status not in OrdemServico.STATUS_VALIDOS:
                status = OrdemServico.STATUS_ABERTA
            if servico_tipo not in OrdemServico.SERVICOS_VALIDOS:
                servico_tipo = OrdemServico.SERVICO_NENHUM

            produto = Produto.query.get(produto_id) if produto_id else None
            if produto_id and not produto:
                flash('Produto informado nao existe.', 'danger')
                return redirect(url_for('criar_ordem_servico'))
            pedido = Pedido.query.get(pedido_id) if pedido_id else None
            if pedido_id and not pedido:
                flash('Venda/pedido informado nao existe.', 'danger')
                return redirect(url_for('criar_ordem_servico'))

            if servico_tipo in {OrdemServico.SERVICO_MONTAGEM, OrdemServico.SERVICO_INSTALACAO}:
                if not empresa or not empresa.servico_montagem_instalacao_ativo:
                    flash('Servico de montagem/instalacao nao liberado para a empresa.', 'danger')
                    return redirect(url_for('criar_ordem_servico'))
                if not produto:
                    flash('Selecione um produto para servico de montagem/instalacao.', 'danger')
                    return redirect(url_for('criar_ordem_servico'))
                if servico_tipo == OrdemServico.SERVICO_MONTAGEM and not produto.servico_montagem_disponivel:
                    flash('Produto sem liberacao para servico de montagem.', 'danger')
                    return redirect(url_for('criar_ordem_servico'))
                if servico_tipo == OrdemServico.SERVICO_INSTALACAO and not produto.servico_instalacao_disponivel:
                    flash('Produto sem liberacao para servico de instalacao.', 'danger')
                    return redirect(url_for('criar_ordem_servico'))

            destino = Funcionario.query.get(funcionario_destino_id) if funcionario_destino_id else None
            if funcionario_destino_id and not destino:
                flash('Funcionario de destino invalido.', 'danger')
                return redirect(url_for('criar_ordem_servico'))

            ordem = OrdemServico(
                titulo=titulo,
                tipo=tipo,
                status=status,
                servico_tipo=servico_tipo,
                prioridade=prioridade,
                descricao=descricao,
                observacoes=observacoes,
                avaria_detalhes=avaria_detalhes,
                inspecao_detalhes=inspecao_detalhes,
                resultado_inspecao=resultado_inspecao,
                produto_id=(produto.id if produto else None),
                pedido_id=(pedido.id if pedido else None),
                funcionario_destino_id=(destino.id if destino else None),
                criado_por_id=(funcionario_logado.id if funcionario_logado else None),
                data_agendada=data_agendada,
            )
            if ordem.status in {OrdemServico.STATUS_ENVIADA, OrdemServico.STATUS_EM_EXECUCAO, OrdemServico.STATUS_CONCLUIDA}:
                ordem.enviada_em = datetime.utcnow()
            if ordem.status == OrdemServico.STATUS_CONCLUIDA:
                ordem.concluida_em = datetime.utcnow()
            db.session.add(ordem)
            db.session.commit()
            flash(f'Ordem #{ordem.id} criada com sucesso.', 'success')
            return redirect(url_for('listar_ordens_servico'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao criar ordem: {str(e)}', 'danger')

    return render_template(
        'servicos/ordens/form.html',
        ordem=None,
        produtos=produtos,
        pedidos=pedidos,
        funcionarios_destino=funcionarios_destino,
        tipos_validos=OrdemServico.TIPOS_VALIDOS,
        status_validos=OrdemServico.STATUS_VALIDOS,
        servicos_validos=OrdemServico.SERVICOS_VALIDOS,
        empresa=empresa,
    )


@app.route('/servicos/<int:ordem_id>/editar', methods=['GET', 'POST'])
@require_role('admin', 'gerente', 'operador', 'caixa')
def editar_ordem_servico(ordem_id):
    if not _servicos_tecnicos_empresa_ativos():
        flash('Modulo de servicos tecnicos desativado na empresa.', 'warning')
        return redirect(url_for('dashboard'))

    ordem = OrdemServico.query.get_or_404(ordem_id)
    empresa = EmpresaConfig.query.first()
    produtos = Produto.query.filter_by(ativo=True).order_by(Produto.nome.asc()).all()
    pedidos = Pedido.query.order_by(Pedido.criado_em.desc()).limit(300).all()
    funcionarios_destino = _funcionarios_tecnicos_montadores()

    if request.method == 'POST':
        try:
            ordem.titulo = (request.form.get('titulo') or '').strip() or ordem.titulo
            tipo = (request.form.get('tipo') or ordem.tipo).strip().lower()
            status = (request.form.get('status') or ordem.status).strip().lower()
            servico_tipo = (request.form.get('servico_tipo') or ordem.servico_tipo).strip().lower()
            if tipo in OrdemServico.TIPOS_VALIDOS:
                ordem.tipo = tipo
            if status in OrdemServico.STATUS_VALIDOS:
                ordem.status = status
            if servico_tipo in OrdemServico.SERVICOS_VALIDOS:
                ordem.servico_tipo = servico_tipo
            ordem.prioridade = (request.form.get('prioridade') or '').strip() or None
            ordem.descricao = (request.form.get('descricao') or '').strip() or None
            ordem.observacoes = (request.form.get('observacoes') or '').strip() or None
            ordem.avaria_detalhes = (request.form.get('avaria_detalhes') or '').strip() or None
            ordem.inspecao_detalhes = (request.form.get('inspecao_detalhes') or '').strip() or None
            ordem.resultado_inspecao = (request.form.get('resultado_inspecao') or '').strip() or None
            ordem.data_agendada = _parse_date_iso(request.form.get('data_agendada'))
            ordem.produto_id = request.form.get('produto_id', type=int) or None
            ordem.pedido_id = request.form.get('pedido_id', type=int) or None
            ordem.funcionario_destino_id = request.form.get('funcionario_destino_id', type=int) or None

            produto = Produto.query.get(ordem.produto_id) if ordem.produto_id else None
            if ordem.pedido_id and not Pedido.query.get(ordem.pedido_id):
                flash('Venda/pedido informado nao existe.', 'danger')
                return redirect(url_for('editar_ordem_servico', ordem_id=ordem.id))
            if ordem.servico_tipo in {OrdemServico.SERVICO_MONTAGEM, OrdemServico.SERVICO_INSTALACAO}:
                if not empresa or not empresa.servico_montagem_instalacao_ativo:
                    flash('Servico de montagem/instalacao nao liberado para a empresa.', 'danger')
                    return redirect(url_for('editar_ordem_servico', ordem_id=ordem.id))
                if not produto:
                    flash('Selecione um produto para servico de montagem/instalacao.', 'danger')
                    return redirect(url_for('editar_ordem_servico', ordem_id=ordem.id))
                if ordem.servico_tipo == OrdemServico.SERVICO_MONTAGEM and not produto.servico_montagem_disponivel:
                    flash('Produto sem liberacao para servico de montagem.', 'danger')
                    return redirect(url_for('editar_ordem_servico', ordem_id=ordem.id))
                if ordem.servico_tipo == OrdemServico.SERVICO_INSTALACAO and not produto.servico_instalacao_disponivel:
                    flash('Produto sem liberacao para servico de instalacao.', 'danger')
                    return redirect(url_for('editar_ordem_servico', ordem_id=ordem.id))

            if ordem.status in {OrdemServico.STATUS_ENVIADA, OrdemServico.STATUS_EM_EXECUCAO, OrdemServico.STATUS_CONCLUIDA} and not ordem.enviada_em:
                ordem.enviada_em = datetime.utcnow()
            if ordem.status == OrdemServico.STATUS_CONCLUIDA:
                ordem.concluida_em = datetime.utcnow()
            elif ordem.status != OrdemServico.STATUS_CONCLUIDA:
                ordem.concluida_em = None

            db.session.commit()
            flash(f'Ordem #{ordem.id} atualizada com sucesso.', 'success')
            return redirect(url_for('listar_ordens_servico'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar ordem: {str(e)}', 'danger')

    return render_template(
        'servicos/ordens/form.html',
        ordem=ordem,
        produtos=produtos,
        pedidos=pedidos,
        funcionarios_destino=funcionarios_destino,
        tipos_validos=OrdemServico.TIPOS_VALIDOS,
        status_validos=OrdemServico.STATUS_VALIDOS,
        servicos_validos=OrdemServico.SERVICOS_VALIDOS,
        empresa=empresa,
    )


@app.route('/servicos/<int:ordem_id>/enviar', methods=['POST'])
@require_role('admin', 'gerente', 'operador', 'caixa')
def enviar_ordem_servico(ordem_id):
    if not _servicos_tecnicos_empresa_ativos():
        flash('Modulo de servicos tecnicos desativado na empresa.', 'warning')
        return redirect(url_for('dashboard'))

    ordem = OrdemServico.query.get_or_404(ordem_id)
    if not ordem.funcionario_destino_id:
        flash('Defina um funcionario de destino antes de enviar.', 'warning')
        return redirect(url_for('editar_ordem_servico', ordem_id=ordem.id))
    try:
        ordem.status = OrdemServico.STATUS_ENVIADA
        ordem.enviada_em = datetime.utcnow()
        db.session.commit()
        flash(f'Ordem #{ordem.id} enviada para o colaborador.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao enviar ordem: {str(e)}', 'danger')
    return redirect(url_for('listar_ordens_servico'))


@app.route('/servicos/<int:ordem_id>/tecnico', methods=['GET', 'POST'])
@login_required
def executar_ordem_servico_tecnico(ordem_id):
    funcionario = get_funcionario_logado()
    ordem = OrdemServico.query.get_or_404(ordem_id)

    if not funcionario:
        flash('Sessao invalida. Faca login novamente.', 'danger')
        return redirect(url_for('login'))

    pode_gerir = (funcionario.role or '').strip().lower() in {'admin', 'gerente'}
    if not pode_gerir and ordem.funcionario_destino_id != funcionario.id:
        flash('Voce nao tem permissao para executar esta ordem.', 'danger')
        return redirect(url_for('minhas_ordens_servico'))

    if request.method == 'POST':
        try:
            acao = (request.form.get('acao') or 'salvar').strip().lower()
            retorno_tecnico = (request.form.get('retorno_tecnico') or '').strip() or None
            observacoes_extra = (request.form.get('observacoes') or '').strip() or None
            resultado_inspecao = (request.form.get('resultado_inspecao') or '').strip() or None

            if acao == 'iniciar':
                ordem.status = OrdemServico.STATUS_EM_EXECUCAO
                ordem.iniciado_em = ordem.iniciado_em or datetime.utcnow()
                ordem.enviada_em = ordem.enviada_em or datetime.utcnow()
            elif acao == 'concluir':
                ordem.status = OrdemServico.STATUS_CONCLUIDA
                ordem.iniciado_em = ordem.iniciado_em or datetime.utcnow()
                ordem.concluida_em = datetime.utcnow()
            elif acao == 'aguardar':
                ordem.status = OrdemServico.STATUS_ENVIADA

            ordem.retorno_tecnico = retorno_tecnico
            ordem.observacoes = observacoes_extra
            ordem.resultado_inspecao = resultado_inspecao
            db.session.commit()
            flash(f'Ordem #{ordem.id} atualizada pelo tecnico.', 'success')
            return redirect(url_for('executar_ordem_servico_tecnico', ordem_id=ordem.id))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao registrar execucao da ordem: {str(e)}', 'danger')

    return render_template(
        'servicos/ordens/tecnico.html',
        ordem=ordem,
        funcionario=funcionario,
        pode_gerir=pode_gerir,
    )


# ============ REGISTRO DE MODULOS DE DOMINIO ============
register_domain_modules(
    app,
    login_required=login_required,
    require_role=require_role,
    aplicar_movimentacao_estoque=aplicar_movimentacao_estoque,
    sincronizar_matriculas_funcionarios=_garantir_matriculas_funcionarios,
)


def _normalizar_historico_assistente(payload_historico):
    historico = []
    for item in payload_historico or []:
        if not isinstance(item, dict):
            continue
        role = (item.get('role') or '').strip().lower()
        text = (item.get('text') or '').strip()
        if role not in {'user', 'assistant'} or not text:
            continue
        historico.append({
            'role': role,
            'text': text[:1000],
        })
        if len(historico) >= 8:
            break
    return historico


def _normalizar_voto_assistente(valor):
    voto = (valor or '').strip().lower()
    if voto in AssistenteLocalFeedback.VOTOS_VALIDOS:
        return voto
    return ''


def _carregar_feedbacks_assistente_local(pagina_atual=None, limite=180):
    query = AssistenteLocalFeedback.query.order_by(AssistenteLocalFeedback.criado_em.desc())
    registros = query.limit(limite).all()
    feedbacks = []
    for item in registros:
        doc_ids = []
        try:
            dados = json.loads(item.matched_doc_ids_json or '[]')
            if isinstance(dados, list):
                doc_ids = [str(valor).strip() for valor in dados if str(valor).strip()]
        except Exception:
            doc_ids = []

        feedbacks.append({
            'vote': item.vote,
            'question': item.question or '',
            'reason': item.reason or '',
            'pagina_atual': item.pagina_atual or '',
            'doc_ids': doc_ids,
        })
    return feedbacks


local_ai_assistant = LocalAIAssistant(app, _construir_documentos_assistente_local)
if app.config.get('LOCAL_AI_ENABLED'):
    local_ai_assistant.start_background_prepare()

register_api_module(app, {
    'login_required': login_required,
    '_limit': _limit,
    '_parse_date_range': _parse_date_range,
    '_coletar_dashboard_analytics': _coletar_dashboard_analytics,
    'get_funcionario_logado': get_funcionario_logado,
    'ENDPOINT_TO_PAGINA': ENDPOINT_TO_PAGINA,
    'local_ai_assistant_getter': lambda: local_ai_assistant,
    '_normalizar_historico_assistente': _normalizar_historico_assistente,
    '_normalizar_voto_assistente': _normalizar_voto_assistente,
    '_carregar_feedbacks_assistente_local': _carregar_feedbacks_assistente_local,
    '_paginas_permitidas_para_funcionario': _paginas_permitidas_para_funcionario,
    'json_dumps': json.dumps,
})


@app.before_request
def validar_acesso_por_pagina():
    ensure_csrf_token()
    csrf_result = csrf_protect_request()
    if csrf_result is not None:
        return csrf_result

    endpoint = request.endpoint
    if not endpoint:
        return None
    if endpoint == 'static' or endpoint.startswith('static'):
        return None
    if endpoint in {'login', 'logout', 'registro', 'index', 'central_ajuda', 'detalhe_ajuda', 'public.cardapio_mesa', 'public.enviar_pedido_qr'}:
        return None
    if 'funcionario_id' not in session:
        return None

    funcionario = get_funcionario_logado()
    if not funcionario or not funcionario.ativo:
        session.clear()
        flash('Sua sessao expirou. Faca login novamente.', 'warning')
        return redirect(url_for('login'))

    if getattr(funcionario, 'senha_provisoria', False):
        endpoints_liberados = {
            'dashboard',
            'logout',
            'pwa_manifest',
            'pwa_service_worker',
        }
        if endpoint not in endpoints_liberados:
            flash('Altere a senha temporaria para continuar usando o sistema.', 'warning')
            return redirect(url_for('dashboard'))

    if not funcionario_tem_acesso(funcionario, endpoint):
        if is_json_request():
            return json_response(False, 'Sem permissao para acessar este recurso.', status=403, code='forbidden')
        flash('Acesso negado para esta pagina.', 'danger')
        return redirect(url_for('boas_vindas'))
    return None


@app.after_request
def registrar_auditoria_pos_resposta(response):
    try:
        if request.method in {'POST', 'PUT', 'PATCH', 'DELETE'}:
            endpoint = request.endpoint or ''
            if not endpoint.startswith('static'):
                ignorar = {
                    'login',
                    'logout',
                }
                if endpoint not in ignorar:
                    funcionario = get_funcionario_logado()
                    detalhes = _resumir_payload_requisicao()
                    registrar_evento_auditoria(
                        funcionario=funcionario,
                        acao=f'{request.method.lower()}_{endpoint or "sem_endpoint"}',
                        entidade='operacao_sistema',
                        detalhes=detalhes,
                        status_code=response.status_code
                    )
    except Exception:
        pass

    try:
        if request.method == 'GET':
            endpoint = request.endpoint or ''
            caminho = request.path or ''
            if caminho == '/sw.js' or endpoint in {'pwa_manifest', 'store_pwa_manifest'}:
                response.headers['Cache-Control'] = 'no-cache, no-store, max-age=0, must-revalidate'
                response.headers['Pragma'] = 'no-cache'
                response.headers['Expires'] = '0'
            elif caminho.startswith('/static/'):
                if request.args.get('v'):
                    response.headers['Cache-Control'] = 'public, max-age=31536000, immutable'
                else:
                    response.headers['Cache-Control'] = 'no-cache, max-age=0, must-revalidate'
                    response.headers['Pragma'] = 'no-cache'
                    response.headers['Expires'] = '0'
            elif response.mimetype == 'text/html':
                response.headers['Cache-Control'] = 'no-cache, no-store, max-age=0, must-revalidate'
                response.headers['Pragma'] = 'no-cache'
                response.headers['Expires'] = '0'
    except Exception:
        pass

    return response


@app.route('/manifest.webmanifest')
def pwa_manifest():
    empresa = EmpresaConfig.query.first()
    app_name = (empresa.nome_fantasia if empresa and empresa.nome_fantasia else APP_NAME)
    icon_path = (
        (getattr(empresa, 'app_icon_path', None) if empresa else None)
        or (getattr(empresa, 'logo_path', None) if empresa else None)
        or (getattr(empresa, 'favicon_path', None) if empresa else None)
        or (empresa.ecom_favicon_path if empresa else None)
        or 'uploads/empresa/Loja_do_Lucas_logo.png'
    )
    extensao_icone = os.path.splitext(icon_path)[1].lower()
    tipo_icone = {
        '.png': 'image/png',
        '.jpg': 'image/jpeg',
        '.jpeg': 'image/jpeg',
        '.webp': 'image/webp',
        '.ico': 'image/x-icon',
        '.svg': 'image/svg+xml',
    }.get(extensao_icone, 'image/png')
    manifest_data = {
        'name': app_name,
        'short_name': app_name[:12] if len(app_name) > 12 else app_name,
        'description': 'Sistema multiplataforma para operacao, vendas, estoque e expedicao.',
        'start_url': url_for('boas_vindas'),
        'scope': '/',
        'display': 'standalone',
        'background_color': '#f4f7fb',
        'theme_color': '#0f766e',
        'lang': 'pt-BR',
        'icons': [
            {
                'src': url_for('static', filename=icon_path),
                'sizes': '192x192',
                'type': tipo_icone,
                'purpose': 'any maskable',
            },
            {
                'src': url_for('static', filename=icon_path),
                'sizes': '512x512',
                'type': tipo_icone,
                'purpose': 'any maskable',
            },
        ],
    }
    return Response(
        json.dumps(manifest_data, ensure_ascii=False),
        mimetype='application/manifest+json',
    )


@app.route('/manifest-loja.webmanifest')
def store_pwa_manifest():
    empresa = EmpresaConfig.query.first()
    store_name = (empresa.nome_fantasia if empresa and empresa.nome_fantasia else APP_NAME)
    icon_path = (
        (getattr(empresa, 'app_icon_path', None) if empresa else None)
        or (getattr(empresa, 'logo_path', None) if empresa else None)
        or (empresa.ecom_favicon_path if empresa else None)
        or (getattr(empresa, 'favicon_path', None) if empresa else None)
        or 'uploads/empresa/Loja_do_Lucas_logo.png'
    )
    extensao_icone = os.path.splitext(icon_path)[1].lower()
    tipo_icone = {
        '.png': 'image/png',
        '.jpg': 'image/jpeg',
        '.jpeg': 'image/jpeg',
        '.webp': 'image/webp',
        '.ico': 'image/x-icon',
        '.svg': 'image/svg+xml',
    }.get(extensao_icone, 'image/png')
    theme_color = (
        (empresa.ecom_cor_primaria if empresa and empresa.ecom_cor_primaria else None)
        or '#ff7848'
    )
    manifest_data = {
        'id': url_for('index'),
        'name': f'{store_name} Loja',
        'short_name': store_name[:12] if len(store_name) > 12 else store_name,
        'description': 'Loja online com pedido rapido e atalho para abrir como app no celular.',
        'start_url': url_for('index'),
        'scope': '/',
        'display': 'standalone',
        'background_color': '#fffdfa',
        'theme_color': theme_color,
        'lang': 'pt-BR',
        'icons': [
            {
                'src': url_for('static', filename=icon_path),
                'sizes': '192x192',
                'type': tipo_icone,
                'purpose': 'any maskable',
            },
            {
                'src': url_for('static', filename=icon_path),
                'sizes': '512x512',
                'type': tipo_icone,
                'purpose': 'any maskable',
            },
        ],
    }
    return Response(
        json.dumps(manifest_data, ensure_ascii=False),
        mimetype='application/manifest+json',
    )


@app.route('/sw.js')
def pwa_service_worker():
    response = send_from_directory(app.static_folder, 'sw.js', mimetype='application/javascript')
    response.headers['Cache-Control'] = 'no-cache, no-store, max-age=0, must-revalidate'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    return response


# ============ CONTEXT PROCESSORS ============
register_context_processors(
    app,
    dependencies={
        'get_funcionario_logado': get_funcionario_logado,
        'empresa_config_model': EmpresaConfig,
        'produto_model': Produto,
        'endpoint_to_pagina': ENDPOINT_TO_PAGINA,
        'titulo_tela_atual': _titulo_tela_atual,
        'estoques_contexto_disponiveis_fn': _estoques_contexto_disponiveis,
        'estoque_contexto_selecionado_id_fn': _estoque_contexto_selecionado_id,
        'ensure_csrf_token_fn': ensure_csrf_token,
        'csrf_input_tag_fn': csrf_input_tag,
        'paginas_permitidas_para_funcionario_fn': _paginas_permitidas_para_funcionario,
        'menu_agrupado_para_paginas_fn': _menu_agrupado_para_paginas,
        'menu_navegacao_principal_fn': _menu_navegacao_principal,
        'montar_indicadores_contexto_usuario_fn': _montar_indicadores_contexto_usuario,
        'registrar_debug_menu_fn': _registrar_debug_menu,
        'local_ai_assistant_getter': lambda: local_ai_assistant,
    },
)


# ============ ERROR HANDLERS ============

@app.errorhandler(400)
def bad_request(error):
    mensagem = getattr(error, 'description', None) or 'Requisicao invalida.'
    acao = resolve_action(code='bad_request', status_code=400)
    if is_json_request():
        return json_response(False, mensagem, status=400, code='bad_request', action=acao)
    return render_template('errors/400.html', error_message=mensagem), 400


@app.errorhandler(403)
def forbidden(error):
    mensagem = getattr(error, 'description', None) or 'Acesso negado.'
    acao = resolve_action(code='forbidden', status_code=403)
    if is_json_request():
        return json_response(False, mensagem, status=403, code='forbidden', action=acao)
    return render_template('errors/403.html', error_message=mensagem), 403


@app.errorhandler(404)
def page_not_found(error):
    return render_template('errors/404.html'), 404


@app.errorhandler(AppError)
def handle_app_error(error):
    mensagem = str(error)
    status_code = getattr(error, 'status_code', 500)
    code = getattr(error, 'code', 'app_error')
    fields = getattr(error, 'fields', {}) or {}
    acao = resolve_action(code=code, status_code=status_code, action=getattr(error, 'action', None))
    if is_json_request():
        return json_response(False, mensagem, status=status_code, code=code, fields=fields, action=acao)
    flash(
        build_flash_message('Aviso' if status_code == 403 else 'Erro', mensagem, acao),
        flash_category_for_status(status_code),
    )
    destino = request.referrer
    if request.endpoint == 'login':
        destino = url_for('login')
    return redirect(destino or url_for('dashboard'))


@app.errorhandler(BusinessRuleError)
def handle_business_rule_error(error):
    return handle_app_error(error)


@app.errorhandler(ValidationError)
def handle_validation_error(error):
    return handle_app_error(error)


@app.errorhandler(PermissionDenied)
def handle_permission_denied(error):
    return handle_app_error(error)


@app.errorhandler(NotFound)
def handle_not_found_error(error):
    return handle_app_error(error)


@app.errorhandler(500)
def internal_error(error):
    db.session.rollback()
    return render_template('errors/500.html'), 500


def create_app():
    """Factory compatível com Flask CLI e scripts do projeto."""
    return app


__all__ = [
    'app',
    'create_app',
    'db',
    'sincronizar_garcom_funcionario',
]


